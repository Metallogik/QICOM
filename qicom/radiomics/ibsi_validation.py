"""Official-resource IBSI reference checks shared by the GUI and CI.

QICOM does not claim IBSI certification.  A successful result means only
that this installation reproduced the mandatory values selected by one
exact, hash-pinned manifest within its declared tolerances.
"""

from __future__ import annotations

import hashlib
import html
import importlib.metadata
import json
import math
import os
import platform
import shutil
import tempfile
import urllib.request
import zipfile
from collections.abc import Callable, Mapping
from dataclasses import dataclass, replace
from datetime import UTC, datetime
from importlib import resources
from pathlib import Path
from typing import Any, Literal

import numpy as np
import pandas as pd
import platformdirs
import SimpleITK as sitk
from openpyxl import load_workbook

from qicom import __version__ as qicom_version
from qicom.atomic_files import commit_temporary, temporary_sibling
from qicom.radiomics.config import MIRP_ADDITIONAL_CLASSES, RadiomicsConfig
from qicom.radiomics.extraction import extract_image_features
from qicom.radiomics.io import load_radiomics_image, read_spatial_image
from qicom.radiomics.models import RadiomicsImage
from qicom.reporting.excel_safety import safe_excel_dataframe
from qicom.schema_version import is_supported_version
from qicom.session.archive_safety import WORKBOOK_LIMITS, validate_zip_inventory
from qicom.session.secure_fs import ensure_private_directory, ensure_private_file

IBSI_MANIFEST_SCHEMA_VERSION = 1
# The "-v2" rule stops truncating a near-zero comparison to exactly zero; see
# _official_comparison_error. A new version changes PASS/FAIL verdicts.
IBSI_COMPARISON_RULE_ID = "official-rounded-difference-3-significant-digits+absolute-near-zero-v2"
IBSI_VERIFICATION_PROJECT_VERSION = 1
IBSI_VERIFICATION_COLUMNS = [
    "feature_key",
    "official_tag",
    "image_type",
    "feature_class",
    "feature_name",
    "official_family",
    "official_feature_name",
    "unit",
    "qicom_value",
    "ibsi_reference_value",
    "absolute_error",
    "relative_error",
    "comparison_error",
    "absolute_tolerance",
    "tolerance_applied",
    "status",
    "reason",
    "backend_version",
    "profile_hash",
    "validation_status",
]

DownloadFunction = Callable[[str, Path], None]
ProgressFunction = Callable[[str, int, int], None]


class ReferenceResourcesMissing(FileNotFoundError):
    """Raised when official files are absent and download was not consented."""


@dataclass(frozen=True, slots=True)
class ReferenceResource:
    resource_id: str
    filename: str
    url: str
    sha256: str
    #: Exact size of the pinned file. Pinned alongside the digest so the
    #: download has a real ceiling: a SHA-256 only tells you the file was
    #: wrong *after* the whole of it has been written, which is no help at
    #: all when "the whole of it" is what the other end chose to send.
    #: 0 means a manifest written before this field existed - the digest
    #: still applies, and the shared cap below still bounds the transfer.
    size_bytes: int = 0


@dataclass(frozen=True, slots=True)
class ReferenceFeature:
    feature_key: str
    official_tag: str
    unit: str
    official_family: str
    official_feature_name: str
    reference_value: float
    absolute_tolerance: float
    mandatory: bool = True


@dataclass(frozen=True, slots=True)
class IbsiReferenceSet:
    manifest: dict[str, Any]
    manifest_hash: str
    resource_paths: dict[str, Path]
    resource_hashes: dict[str, str]
    image: RadiomicsImage
    mask: np.ndarray
    config: RadiomicsConfig
    features: tuple[ReferenceFeature, ...]

    @property
    def profile_hash(self) -> str:
        return self.manifest_hash


@dataclass(frozen=True, slots=True)
class IbsiVerificationResult:
    table: pd.DataFrame
    summary: dict[str, Any]
    manifest: dict[str, Any]
    receipt: dict[str, Any]

    @property
    def overall_status(self) -> str:
        return str(self.summary["overall_status"])


@dataclass(frozen=True, slots=True)
class IbsiVerificationStudy:
    reference_set_id: str
    manifest_hash: str
    mode: Literal["reference-verification", "derived-experiment"] = "reference-verification"
    verifiable: bool = True
    config: RadiomicsConfig | None = None
    result: IbsiVerificationResult | None = None
    derived_reason: str | None = None

    def duplicate_as_experiment(
        self, config: RadiomicsConfig | None = None
    ) -> IbsiVerificationStudy:
        resolved = self.config if config is None else config
        return replace(
            self,
            mode="derived-experiment",
            verifiable=False,
            config=replace(resolved, ibsi_preset="none") if resolved is not None else None,
            result=None,
            derived_reason=(
                "Configuration or mask may be modified; this copy no longer corresponds to "
                "the hash-pinned IBSI reference configuration and cannot produce a verdict."
            ),
        )


def _canonical_json(payload: Mapping[str, Any]) -> str:
    return json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def default_reference_cache() -> Path:
    return Path(platformdirs.user_data_dir("qicom", appauthor=False)) / "ibsi-reference-sets"


def available_reference_set_ids() -> tuple[str, ...]:
    root = resources.files("qicom.radiomics").joinpath("reference_sets")
    return tuple(sorted(item.name.removesuffix(".json") for item in root.iterdir()))


def load_reference_manifest(
    reference_set_id: str = "ibsi1_digital_phantom",
) -> dict[str, Any]:
    if not reference_set_id or Path(reference_set_id).name != reference_set_id:
        raise ValueError("reference_set_id must be a simple packaged identifier.")
    item = resources.files("qicom.radiomics").joinpath("reference_sets", f"{reference_set_id}.json")
    try:
        payload = json.loads(item.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"Unknown IBSI reference set {reference_set_id!r}.") from exc
    _validate_manifest(payload)
    return payload


def _validate_manifest(payload: Any) -> None:
    if not isinstance(payload, dict):
        raise ValueError("IBSI manifest must be a JSON object.")
    if not is_supported_version(payload.get("schema_version"), (IBSI_MANIFEST_SCHEMA_VERSION,)):
        raise ValueError("Unsupported IBSI manifest schema version.")
    required = {
        "reference_set_id",
        "reference_set_version",
        "verification_type",
        "scope",
        "pipeline_components",
        "resources",
        "expected_geometry",
        "configuration",
        "reference_values",
        "features",
    }
    missing = required - set(payload)
    if missing:
        raise ValueError(f"IBSI manifest is missing keys: {sorted(missing)}.")
    resources_payload = payload["resources"]
    if not isinstance(resources_payload, list) or not resources_payload:
        raise ValueError("IBSI manifest resources must be a non-empty list.")
    resource_ids: set[str] = set()
    filenames: set[str] = set()
    for entry in resources_payload:
        if not isinstance(entry, dict):
            raise ValueError("Every IBSI resource entry must be an object.")
        resource_id = entry.get("id")
        filename = entry.get("filename")
        digest = entry.get("sha256")
        if not isinstance(resource_id, str) or not resource_id:
            raise ValueError("Every IBSI resource requires a non-empty id.")
        if not isinstance(filename, str) or Path(filename).name != filename:
            raise ValueError("IBSI resource filenames must be simple path components.")
        if (
            not isinstance(digest, str)
            or len(digest) != 64
            or any(character not in "0123456789abcdef" for character in digest.lower())
        ):
            raise ValueError("IBSI resource SHA-256 values must be 64 hexadecimal characters.")
        if resource_id in resource_ids or filename in filenames:
            raise ValueError("IBSI resource ids and filenames must be unique.")
        resource_ids.add(resource_id)
        filenames.add(filename)
    features = payload["features"]
    if not isinstance(features, list) or not features:
        raise ValueError("IBSI manifest features must be a non-empty list.")
    official_tags: set[str] = set()
    for entry in features:
        if (
            not isinstance(entry, list)
            or len(entry) != 3
            or not all(isinstance(value, str) and value for value in entry)
        ):
            raise ValueError("Every IBSI feature mapping must be [feature_key, tag, unit].")
        if entry[1] in official_tags:
            raise ValueError(f"Duplicate official IBSI tag {entry[1]!r}.")
        official_tags.add(entry[1])
    RadiomicsConfig.from_dict(payload["configuration"])


def manifest_hash(manifest: Mapping[str, Any]) -> str:
    return hashlib.sha256(_canonical_json(manifest).encode("utf-8")).hexdigest()


def _resource_records(manifest: Mapping[str, Any]) -> tuple[ReferenceResource, ...]:
    return tuple(
        ReferenceResource(
            resource_id=str(entry["id"]),
            filename=str(entry["filename"]),
            url=str(entry["url"]),
            sha256=str(entry["sha256"]).lower(),
            size_bytes=int(entry.get("size_bytes", 0)),
        )
        for entry in manifest["resources"]
    )


def verify_resource_integrity(manifest: Mapping[str, Any], cache_dir: str | Path) -> dict[str, str]:
    root = Path(cache_dir)
    hashes: dict[str, str] = {}
    for resource in _resource_records(manifest):
        path = root / resource.filename
        if not path.is_file():
            raise ReferenceResourcesMissing(f"Official IBSI resource is missing: {path.name}.")
        actual = _sha256_file(path)
        if actual != resource.sha256:
            raise ValueError(
                f"IBSI resource checksum mismatch for {path.name}: expected "
                f"{resource.sha256}, got {actual}."
            )
        hashes[resource.resource_id] = actual
    return hashes


#: Ceiling for any single reference resource whose manifest predates
#: `size_bytes`. The largest official IBSI file is a few hundred kilobytes;
#: this is generous enough never to reject a real one and small enough that a
#: hostile or misconfigured endpoint cannot fill a disk.
MAX_REFERENCE_RESOURCE_BYTES = 256 * 1024 * 1024
_DOWNLOAD_CHUNK_BYTES = 1024 * 1024
#: Headroom demanded before a download starts, as a multiple of the pinned
#: size. A file that exactly fits the free space leaves nothing behind it.
_DOWNLOAD_FREE_SPACE_MARGIN = 3.0


def _download(url: str, destination: Path, max_bytes: int = MAX_REFERENCE_RESOURCE_BYTES) -> None:
    """Fetch one pinned resource, refusing to write more than `max_bytes`.

    Two independent limits, because neither alone is enough. `Content-Length`
    is checked when the server sends one, so an oversized transfer is
    abandoned before it starts - but the header is the server's claim, and a
    server willing to send the wrong file is willing to send the wrong length
    with it. So the streamed bytes are counted too, and the transfer is cut
    off at the first byte past the limit rather than after however much the
    other end decided to send.
    """
    request = urllib.request.Request(url, headers={"User-Agent": f"QICOM/{qicom_version}"})
    with urllib.request.urlopen(request, timeout=60) as response:
        declared = response.headers.get("Content-Length")
        if declared is not None:
            try:
                declared_bytes = int(declared)
            except ValueError:
                declared_bytes = -1
            if declared_bytes > max_bytes:
                raise ValueError(
                    f"Refusing to download {url}: the server declares {declared_bytes} bytes, "
                    f"more than the {max_bytes}-byte limit for this resource."
                )
        written = 0
        with destination.open("wb") as output:
            while chunk := response.read(_DOWNLOAD_CHUNK_BYTES):
                written += len(chunk)
                if written > max_bytes:
                    raise ValueError(
                        f"Refusing to download {url}: the transfer exceeded the "
                        f"{max_bytes}-byte limit for this resource."
                    )
                output.write(chunk)


def _check_download_space(root: Path, required_bytes: int, filename: str) -> None:
    """Refuse a download the disk cannot hold, before it starts.

    The temporary file and the final copy briefly coexist, and the margin
    covers that plus ordinary headroom.
    """
    try:
        free = shutil.disk_usage(root).free
    except OSError:
        return
    needed = int(required_bytes * _DOWNLOAD_FREE_SPACE_MARGIN)
    if free < needed:
        raise ValueError(
            f"Not enough free space to download {filename}: about {needed} bytes are needed "
            f"including headroom, and {free} are available."
        )


def ensure_reference_resources(
    manifest: Mapping[str, Any],
    cache_dir: str | Path | None = None,
    *,
    consent: bool = False,
    downloader: DownloadFunction | None = None,
    progress: ProgressFunction | None = None,
) -> dict[str, Path]:
    """Return verified local resources, downloading only with explicit consent."""
    root = default_reference_cache() if cache_dir is None else Path(cache_dir)
    ensure_private_directory(root)
    records = _resource_records(manifest)
    paths: dict[str, Path] = {}
    for index, resource in enumerate(records, start=1):
        destination = root / resource.filename
        if destination.is_file() and _sha256_file(destination) == resource.sha256:
            paths[resource.resource_id] = destination
            if progress is not None:
                progress(resource.filename, index, len(records))
            continue
        if not consent:
            raise ReferenceResourcesMissing(
                "The official IBSI reference resources are not available offline. "
                "Explicit download consent is required."
            )
        budget = resource.size_bytes or MAX_REFERENCE_RESOURCE_BYTES
        _check_download_space(root, budget, resource.filename)
        with tempfile.NamedTemporaryFile(
            dir=root, prefix=f".{resource.filename}.", suffix=".download", delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        try:
            if downloader is not None:
                downloader(resource.url, temporary_path)
            else:
                _download(resource.url, temporary_path, budget)
            downloaded = temporary_path.stat().st_size
            if resource.size_bytes and downloaded != resource.size_bytes:
                raise ValueError(
                    f"Downloaded IBSI resource {resource.filename} is {downloaded} bytes; "
                    f"the manifest pins {resource.size_bytes}."
                )
            if downloaded > budget:
                raise ValueError(
                    f"Downloaded IBSI resource {resource.filename} is {downloaded} bytes, "
                    f"over the {budget}-byte limit for this resource."
                )
            actual = _sha256_file(temporary_path)
            if actual != resource.sha256:
                raise ValueError(
                    f"Downloaded IBSI resource {resource.filename} failed SHA-256 verification."
                )
            os.replace(temporary_path, destination)
            ensure_private_file(destination)
        finally:
            # Whatever happened - an overrun, a wrong digest, an interrupted
            # transfer - the partial file goes and any previously verified
            # copy at `destination` is left exactly as it was.
            temporary_path.unlink(missing_ok=True)
        paths[resource.resource_id] = destination
        if progress is not None:
            progress(resource.filename, index, len(records))
    verify_resource_integrity(manifest, root)
    return paths


def _load_reference_features(
    manifest: Mapping[str, Any], resource_paths: Mapping[str, Path]
) -> tuple[ReferenceFeature, ...]:
    settings = manifest["reference_values"]
    workbook_path = resource_paths[str(settings["resource_id"])]
    # A spreadsheet is a ZIP, and this one is downloaded rather than shipped.
    # openpyxl has no budget of its own, so the entry count, declared size and
    # compression ratio are checked against the workbook profile first.
    with zipfile.ZipFile(workbook_path) as workbook_archive:
        validate_zip_inventory(workbook_archive, limits=WORKBOOK_LIMITS)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[str(settings["worksheet"])]
        header = {
            str(cell.value): index
            for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)))
        }
        required_columns = {
            str(settings["tag_column"]),
            str(settings["value_column"]),
            str(settings["absolute_tolerance_column"]),
            "family",
            "feature",
        }
        missing_columns = required_columns - set(header)
        if missing_columns:
            raise ValueError(
                f"Official IBSI spreadsheet is missing columns: {sorted(missing_columns)}."
            )
        rows_by_tag: dict[str, tuple[Any, ...]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            tag_value = row[header[str(settings["tag_column"])]]
            if isinstance(tag_value, str) and tag_value:
                rows_by_tag[tag_value] = row
        result: list[ReferenceFeature] = []
        for feature_key, tag, unit in manifest["features"]:
            row = rows_by_tag.get(tag)
            if row is None:
                raise ValueError(f"Official IBSI spreadsheet does not contain tag {tag!r}.")
            expected = row[header[str(settings["value_column"])]]
            tolerance = row[header[str(settings["absolute_tolerance_column"])]]
            if not isinstance(expected, (int, float)) or isinstance(expected, bool):
                raise ValueError(f"IBSI tag {tag!r} has no numeric reference value.")
            if not isinstance(tolerance, (int, float)) or isinstance(tolerance, bool):
                raise ValueError(f"IBSI tag {tag!r} has no numeric absolute tolerance.")
            result.append(
                ReferenceFeature(
                    feature_key=feature_key,
                    official_tag=tag,
                    unit=unit,
                    official_family=str(row[header["family"]]),
                    official_feature_name=str(row[header["feature"]]),
                    reference_value=float(expected),
                    absolute_tolerance=float(tolerance),
                )
            )
        return tuple(result)
    finally:
        workbook.close()


def _validate_geometry_and_mask(
    manifest: Mapping[str, Any], image: RadiomicsImage, mask_image: sitk.Image
) -> np.ndarray:
    expected = manifest["expected_geometry"]
    if tuple(expected["size_xyz"]) != image.geometry.size:
        raise ValueError("Official IBSI phantom size does not match the manifest.")
    for field_name, actual in (
        ("spacing_xyz_mm", image.geometry.spacing),
        ("origin_xyz_mm", image.geometry.origin),
        ("direction", image.geometry.direction),
    ):
        if not np.allclose(expected[field_name], actual, rtol=0, atol=1e-8):
            raise ValueError(f"Official IBSI phantom {field_name} does not match the manifest.")
    if (
        tuple(mask_image.GetSize()) != image.geometry.size
        or not np.allclose(mask_image.GetSpacing(), image.geometry.spacing, rtol=0, atol=1e-8)
        or not np.allclose(mask_image.GetOrigin(), image.geometry.origin, rtol=0, atol=1e-8)
        or not np.allclose(mask_image.GetDirection(), image.geometry.direction, rtol=0, atol=1e-8)
    ):
        raise ValueError("Official IBSI image and mask geometry disagree.")
    raw_mask = sitk.GetArrayFromImage(mask_image)
    label = int(expected["mask_label"])
    labels = set(int(value) for value in np.unique(raw_mask))
    if labels - {0, label} or label not in labels:
        raise ValueError(f"Official IBSI mask labels are invalid: {sorted(labels)}.")
    mask = raw_mask == label
    if int(mask.sum()) != int(expected["mask_voxels"]):
        raise ValueError("Official IBSI mask voxel count does not match the manifest.")
    return mask


def load_ibsi_reference_set(
    reference_set_id: str = "ibsi1_digital_phantom",
    *,
    cache_dir: str | Path | None = None,
    consent: bool = False,
    downloader: DownloadFunction | None = None,
    progress: ProgressFunction | None = None,
) -> IbsiReferenceSet:
    manifest = load_reference_manifest(reference_set_id)
    resource_paths = ensure_reference_resources(
        manifest,
        cache_dir,
        consent=consent,
        downloader=downloader,
        progress=progress,
    )
    hashes = {
        record.resource_id: _sha256_file(resource_paths[record.resource_id])
        for record in _resource_records(manifest)
    }
    image = load_radiomics_image(resource_paths["image"])
    mask_image = read_spatial_image(resource_paths["mask"])
    mask = _validate_geometry_and_mask(manifest, image, mask_image)
    return IbsiReferenceSet(
        manifest=dict(manifest),
        manifest_hash=manifest_hash(manifest),
        resource_paths=dict(resource_paths),
        resource_hashes=hashes,
        image=image,
        mask=mask,
        config=RadiomicsConfig.from_dict(manifest["configuration"]),
        features=_load_reference_features(manifest, resource_paths),
    )


def _feature_parts(feature_key: str) -> tuple[str, str, str]:
    for family in MIRP_ADDITIONAL_CLASSES:
        prefix = f"original_{family}_"
        if feature_key.startswith(prefix):
            return "original", family, feature_key[len(prefix) :]
    parts = feature_key.split("_", 2)
    if len(parts) != 3:
        return "", "", feature_key
    return parts[0], parts[1], parts[2]


def _official_comparison_error(
    error: float, reference_value: float, *, zero_threshold: float
) -> float:
    """Reproduce the spreadsheet's truncated 3-significant-digit difference.

    A reference value at, or within ``zero_threshold`` of, zero has no
    significant-digit scale to truncate against. The pre-v2 rule fell back to
    ``digits = 0``, which truncated *every* difference below 1.0 to exactly
    zero: an error of 0.4 against a reference of 0 with a tolerance of 0 was
    reported as an official PASS. Near zero the raw absolute difference is
    therefore compared untruncated, which is the reproducible absolute rule
    declared as ``reference_values.comparison_rule`` in the manifest.
    """
    if not math.isfinite(error):
        return math.inf
    if abs(reference_value) <= zero_threshold:
        return abs(error)
    digits = 2 - math.floor(math.log10(abs(reference_value)))
    factor = 10.0**digits
    # Keep the arithmetic a native Python float: NumPy may expose ``eps``
    # as ``longdouble`` on some builds, which ``math.trunc`` cannot consume.
    return math.trunc((error + math.ulp(1.0)) * factor) / factor


def compare_with_ibsi_reference(
    calculated: pd.DataFrame,
    reference_set: IbsiReferenceSet,
    *,
    backend_version: str | None = None,
) -> IbsiVerificationResult:
    required_columns = {"feature_key", "value", "unit"}
    missing = required_columns - set(calculated.columns)
    if missing:
        raise ValueError(f"Calculated feature table is missing columns: {sorted(missing)}.")
    if calculated["feature_key"].duplicated().any():
        raise ValueError("Calculated IBSI feature keys must be unique.")
    calculated_rows = {str(row["feature_key"]): row for row in calculated.to_dict(orient="records")}
    resolved_backend = backend_version or _backend_version(calculated)
    zero_threshold = float(reference_set.manifest["reference_values"]["relative_zero_threshold"])
    rows: list[dict[str, Any]] = []
    for expected in reference_set.features:
        actual = calculated_rows.get(expected.feature_key)
        image_type, feature_class, feature_name = _feature_parts(expected.feature_key)
        base = {
            "feature_key": expected.feature_key,
            "official_tag": expected.official_tag,
            "image_type": image_type,
            "feature_class": feature_class,
            "feature_name": feature_name,
            "official_family": expected.official_family,
            "official_feature_name": expected.official_feature_name,
            "unit": expected.unit,
            "ibsi_reference_value": expected.reference_value,
            "absolute_tolerance": expected.absolute_tolerance,
            "tolerance_applied": (
                "Official IBSI spreadsheet absolute tolerance after truncating the raw "
                "difference to the reference value's 3-significant-digit precision; "
                "a reference value within "
                f"{zero_threshold:g} of zero is compared as an untruncated absolute "
                f"difference instead ({IBSI_COMPARISON_RULE_ID})."
            ),
            "backend_version": resolved_backend,
            "profile_hash": reference_set.profile_hash,
        }
        if actual is None:
            rows.append(
                {
                    **base,
                    "qicom_value": np.nan,
                    "absolute_error": np.nan,
                    "relative_error": np.nan,
                    "comparison_error": np.nan,
                    "status": "NOT-TESTED",
                    "reason": "Mandatory mapped feature was not returned by the backend.",
                    "validation_status": "not-covered-by-reference",
                }
            )
            continue
        value = float(actual["value"])
        absolute_error = abs(value - expected.reference_value)
        relative_is_undefined = abs(expected.reference_value) <= zero_threshold
        relative_error = (
            np.nan if relative_is_undefined else absolute_error / abs(expected.reference_value)
        )
        comparison_error = _official_comparison_error(
            absolute_error, expected.reference_value, zero_threshold=zero_threshold
        )
        unit_matches = str(actual["unit"]) == expected.unit
        passed = (
            math.isfinite(value)
            and unit_matches
            and comparison_error <= expected.absolute_tolerance
        )
        if not unit_matches:
            reason = f"Unit mismatch: calculated {actual['unit']!r}, expected {expected.unit!r}."
        else:
            reason = "Within the manifest tolerance." if passed else "Outside tolerance."
        if relative_is_undefined:
            # The relative error stays NA rather than being reported as zero, and
            # the row records why so a receipt reader is not left guessing.
            reason += (
                " Relative error is not defined: the IBSI reference value is within "
                f"{zero_threshold:g} of zero, so the untruncated absolute difference "
                "was compared against the tolerance."
            )
        rows.append(
            {
                **base,
                "qicom_value": value,
                "absolute_error": absolute_error,
                "relative_error": relative_error,
                "comparison_error": comparison_error,
                "status": "PASS" if passed else "FAIL",
                "reason": reason,
                "validation_status": "reference-verified" if passed else "reference-failed",
            }
        )
    table = pd.DataFrame(rows, columns=IBSI_VERIFICATION_COLUMNS)
    counts = table["status"].value_counts().to_dict()
    if counts.get("FAIL", 0):
        overall = "FAIL"
    elif counts.get("NOT-TESTED", 0):
        overall = "INCOMPLETE"
    else:
        overall = "PASS"
    now = datetime.now(UTC).isoformat()
    classes = sorted(set(table.loc[table["status"] != "NOT-TESTED", "feature_class"]))
    summary = {
        "passed": int(counts.get("PASS", 0)),
        "failed": int(counts.get("FAIL", 0)),
        "not_tested": int(counts.get("NOT-TESTED", 0)),
        "mandatory_features": len(reference_set.features),
        "feature_classes_covered": classes,
        "pipeline_components_covered": list(reference_set.manifest["pipeline_components"]),
        "overall_status": overall,
        "configuration": reference_set.config.to_dict(),
        "reference_set_version": reference_set.manifest["reference_set_version"],
        "qicom_version": qicom_version,
        "backend": reference_set.config.backend,
        "backend_version": resolved_backend,
        **(
            {"pyradiomics_version": resolved_backend}
            if reference_set.config.backend == "pyradiomics"
            else {"mirp_version": resolved_backend}
        ),
        "simpleitk_version": sitk.Version_VersionString(),
        "timestamp_utc": now,
        "platform": platform.platform(),
        "resource_hashes": dict(reference_set.resource_hashes),
        "profile_hash": reference_set.profile_hash,
    }
    receipt = {
        "receipt_type": "qicom-ibsi-reference-verification",
        "receipt_version": 1,
        "reference_set_id": reference_set.manifest["reference_set_id"],
        "reference_set_version": reference_set.manifest["reference_set_version"],
        "manifest_hash": reference_set.manifest_hash,
        "configuration_hash": reference_set.config.config_hash,
        "result_table_sha256": hashlib.sha256(
            table.to_csv(index=False).encode("utf-8")
        ).hexdigest(),
        "summary": summary,
        "calculation_details": calculated.attrs.get("radiomics_provenance", {}),
        "calculation_providers": {
            key: str(row.get("provider", "unrecorded")) for key, row in calculated_rows.items()
        },
        "methodology": (
            "Calculated values are matched to official IBSI spreadsheet rows selected by "
            "the hash-pinned manifest. Missing mandatory values produce INCOMPLETE."
        ),
        "limitations": reference_set.manifest["scope"],
        "non_certification_statement": (
            "This is not IBSI certification or clinical validation. It reports only whether "
            "this installation reproduced this exact reference configuration."
        ),
    }
    return IbsiVerificationResult(table, summary, reference_set.manifest, receipt)


def _backend_version(calculated: pd.DataFrame) -> str:
    provenance = calculated.attrs.get("radiomics_provenance", {})
    if provenance.get("backend_version"):
        return str(provenance["backend_version"])
    if "provider" in calculated.columns and not calculated.empty:
        provider = str(calculated.iloc[0]["provider"])
        return provider.removeprefix("pyradiomics-")
    try:
        return importlib.metadata.version("pyradiomics")
    except importlib.metadata.PackageNotFoundError:
        return "unavailable"


def run_ibsi_verification(reference_set: IbsiReferenceSet) -> IbsiVerificationResult:
    calculated = extract_image_features(
        reference_set.image,
        reference_set.mask,
        reference_set.config,
        image_role="ibsi-reference",
        roi_id="ibsi-official-mask",
        roi_name="Official IBSI mask",
    )
    return compare_with_ibsi_reference(calculated, reference_set)


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def export_ibsi_receipt_json(path: str | Path, result: IbsiVerificationResult) -> Path:
    destination = Path(path).with_suffix(".json")
    payload = {
        **result.receipt,
        "manifest": result.manifest,
        "comparisons": _json_safe(result.table.to_dict(orient="records")),
    }
    _atomic_text(destination, json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False))
    return destination


def export_ibsi_receipt_csv(path: str | Path, result: IbsiVerificationResult) -> Path:
    destination = Path(path).with_suffix(".csv")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = temporary_sibling(destination)
    try:
        result.table.to_csv(temporary, index=False)
        commit_temporary(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination


def export_ibsi_receipt_excel(path: str | Path, result: IbsiVerificationResult) -> Path:
    destination = Path(path).with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = temporary_sibling(destination)
    summary = pd.DataFrame(
        [
            {"field": key, "value": json.dumps(value) if isinstance(value, (dict, list)) else value}
            for key, value in result.summary.items()
        ]
    )
    manifest = pd.DataFrame(
        [
            {"field": key, "value": json.dumps(value) if isinstance(value, (dict, list)) else value}
            for key, value in result.manifest.items()
        ]
    )
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            safe_excel_dataframe(result.table).to_excel(
                writer, sheet_name="IBSI Comparisons", index=False
            )
            safe_excel_dataframe(summary).to_excel(writer, sheet_name="Summary", index=False)
            safe_excel_dataframe(manifest).to_excel(writer, sheet_name="Manifest", index=False)
        commit_temporary(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination


def export_ibsi_receipt_html(path: str | Path, result: IbsiVerificationResult) -> Path:
    destination = Path(path).with_suffix(".html")
    rows = []
    for record in result.table.to_dict(orient="records"):
        cells = "".join(
            f"<td>{html.escape(str(_json_safe(record[column])))}</td>"
            for column in IBSI_VERIFICATION_COLUMNS
        )
        rows.append(f"<tr>{cells}</tr>")
    headings = "".join(f"<th>{html.escape(column)}</th>" for column in IBSI_VERIFICATION_COLUMNS)
    document = "".join(
        [
            "<!doctype html><html><head><meta charset='utf-8'>"
            "<title>QICOM IBSI verification</title>",
            "<style>body{font-family:sans-serif;margin:2rem}table{border-collapse:collapse}",
            "th,td{border:1px solid #aaa;padding:.25rem;text-align:left}</style></head><body>",
            "<h1>QICOM IBSI reference verification</h1>",
            f"<p>Overall: {html.escape(result.overall_status)}</p>",
            f"<p>{html.escape(str(result.receipt['non_certification_statement']))}</p>",
            f"<p>{html.escape(str(result.receipt['limitations']))}</p>",
            f"<table><thead><tr>{headings}</tr></thead><tbody>{''.join(rows)}</tbody></table>",
            "</body></html>",
        ]
    )
    _atomic_text(destination, document)
    return destination


def _atomic_text(destination: Path, content: str) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = temporary_sibling(destination)
    try:
        temporary.write_text(content, encoding="utf-8")
        commit_temporary(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def save_ibsi_verification_project(path: str | Path, study: IbsiVerificationStudy) -> Path:
    destination = Path(path).with_suffix(".qicom-ibsi.json")
    payload: dict[str, Any] = {
        "qicom_ibsi_verification_project_version": IBSI_VERIFICATION_PROJECT_VERSION,
        "project_kind": "ibsi-reference-verification",
        "reference_set_id": study.reference_set_id,
        "manifest_hash": study.manifest_hash,
        "mode": study.mode,
        "verifiable": study.verifiable,
        "config": None if study.config is None else study.config.to_dict(),
        "derived_reason": study.derived_reason,
        "result": None,
    }
    if study.result is not None:
        payload["result"] = {
            "summary": study.result.summary,
            "receipt": study.result.receipt,
            "manifest": study.result.manifest,
            "table": _json_safe(study.result.table.to_dict(orient="records")),
        }
    _atomic_text(destination, json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False))
    ensure_private_file(destination)
    return destination


def load_ibsi_verification_project(path: str | Path) -> IbsiVerificationStudy:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if not is_supported_version(payload.get("qicom_ibsi_verification_project_version"), (1,)):
        raise ValueError("Unsupported QICOM IBSI verification project version.")
    if payload.get("project_kind") != "ibsi-reference-verification":
        raise ValueError("This file is not a QICOM IBSI verification project.")
    result_payload = payload.get("result")
    result = None
    if isinstance(result_payload, dict):
        result = IbsiVerificationResult(
            pd.DataFrame(result_payload.get("table", []), columns=IBSI_VERIFICATION_COLUMNS),
            dict(result_payload["summary"]),
            dict(result_payload["manifest"]),
            dict(result_payload["receipt"]),
        )
    config_payload = payload.get("config")
    return IbsiVerificationStudy(
        reference_set_id=str(payload["reference_set_id"]),
        manifest_hash=str(payload["manifest_hash"]),
        mode=str(payload["mode"]),  # type: ignore[arg-type]
        verifiable=bool(payload["verifiable"]),
        config=(None if config_payload is None else RadiomicsConfig.from_dict(config_payload)),
        result=result,
        derived_reason=(
            None if payload.get("derived_reason") is None else str(payload["derived_reason"])
        ),
    )


__all__ = [
    "IBSI_VERIFICATION_COLUMNS",
    "IbsiReferenceSet",
    "IbsiVerificationResult",
    "IbsiVerificationStudy",
    "ReferenceResourcesMissing",
    "available_reference_set_ids",
    "compare_with_ibsi_reference",
    "ensure_reference_resources",
    "export_ibsi_receipt_csv",
    "export_ibsi_receipt_excel",
    "export_ibsi_receipt_html",
    "export_ibsi_receipt_json",
    "load_ibsi_reference_set",
    "load_ibsi_verification_project",
    "load_reference_manifest",
    "manifest_hash",
    "run_ibsi_verification",
    "save_ibsi_verification_project",
    "verify_resource_integrity",
]
