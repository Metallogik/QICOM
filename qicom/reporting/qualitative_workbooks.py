"""The two qualitative-evaluation workbooks, writable without the GUI.

One evaluator's own answers (write_own_answers_report - answers and nothing
else, since every statistic here needs more than one rater to mean anything)
and the full aggregated multi-evaluator report (write_aggregated_report), the
openpyxl-writing counterpart to qicom.session.evaluation_aggregation's pure
pandas statistics. They lived in qicom.gui, so the aggregated report could only
be produced from a window; here a batch or a script can write it too.

They follow the native-chart conventions of
export_native_charts_to_workbook/export_line_profiles_to_workbook (a data table
plus a real, editable chart object - never a screenshot - NaN written as blank
cells, deterministic safe sheet names) and reuse export_native_charts_to_workbook
for the "Summary"/"By Evaluator" sheets.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from qicom.reporting.excel_safety import safe_excel_value, safe_sheet_name
from qicom.reporting.workbooks import (
    export_native_charts_to_workbook,
    save_workbook_atomically,
)
from qicom.session.evaluation_aggregation import (
    aggregate_by_evaluator_and_question,
    aggregate_by_question,
    category_distribution,
    descriptive_by_question,
    filter_open_text,
    inter_rater_reliability,
    inter_rater_spread,
    likert_distribution,
    paired_comparison_analysis,
    pairwise_auc_by_comparison,
    ranking_concordance,
)
from qicom.session.evaluation_answer_files import RAW_ANSWERS_SHEET
from qicom.session.qualitative_qc import (
    evaluator_quality_flags,
    intra_rater_repeat_reliability,
    precision_and_completeness,
)


def _write_dataframe(sheet: Worksheet, df: pd.DataFrame) -> None:
    sheet.append([safe_excel_value(str(name)) for name in df.columns])
    for values in df.itertuples(index=False, name=None):
        row = []
        for value in values:
            if isinstance(value, float) and value != value:  # NaN
                value = None
            row.append(safe_excel_value(value))
        sheet.append(row)


def _bar_groups_from_summary(
    summary_df: pd.DataFrame,
) -> list[tuple[str, list[str], dict[str, list[float]]]]:
    """One group per (page, comparison image) - categories are its
    questions, one "Mean Score" series - exactly the shape export_native_
    charts_to_workbook's own `groups` parameter expects. Grouped by page
    as well as comparison_role because comparison_role is only unique
    *within* one page (e.g. two different pages can each have their own
    "comparison" role pointing at two entirely different images) - see
    EvaluationPage's own docstring."""
    groups: list[tuple[str, list[str], dict[str, list[float]]]] = []
    grouped_df = summary_df.groupby(["page_title", "comparison_role"])
    for (page_title, comparison_role), group in grouped_df:
        ordered = group.sort_values("question_id")
        categories = list(ordered["question_text"])
        series = {"Mean Score": list(ordered["mean"])}
        groups.append((f"{page_title} - {comparison_role}", categories, series))
    return groups


def _bar_groups_from_by_evaluator(
    by_evaluator_df: pd.DataFrame,
) -> list[tuple[str, list[str], dict[str, list[float]]]]:
    """One group per (page, comparison image) - categories are its
    questions, one series per evaluator (their display name, or id if none
    was ever set). See _bar_groups_from_summary's own docstring for why
    grouping needs the page as well as the comparison role."""
    groups: list[tuple[str, list[str], dict[str, list[float]]]] = []
    for (page_title, comparison_role), role_group in by_evaluator_df.groupby(
        ["page_title", "comparison_role"]
    ):
        questions = (
            role_group[["question_id", "question_text"]]
            .drop_duplicates()
            .sort_values("question_id")
        )
        categories = list(questions["question_text"])
        series: dict[str, list[float]] = {}
        for evaluator_id, evaluator_group in role_group.groupby("evaluator_id"):
            display_name = evaluator_group["evaluator_display_name"].iloc[0] or evaluator_id
            by_question = dict(
                zip(evaluator_group["question_id"], evaluator_group["mean"], strict=True)
            )
            series[str(display_name)] = [
                by_question.get(question_id, float("nan"))
                for question_id in questions["question_id"]
            ]
        groups.append((f"{page_title} - {comparison_role}", categories, series))
    return groups


def _write_spread_sheet(workbook: Workbook, spread_df: pd.DataFrame) -> None:
    sheet = workbook.create_sheet("Inter-Rater Spread")
    header = ["Page", "Comparison", "Question", "Std Dev (lower = more agreement)"]
    for col, title in enumerate(header, start=1):
        sheet.cell(row=1, column=col, value=title)
    row = 2
    for _, entry in spread_df.iterrows():
        sheet.cell(row=row, column=1, value=safe_excel_value(str(entry["page_title"])))
        sheet.cell(row=row, column=2, value=safe_excel_value(str(entry["comparison_role"])))
        sheet.cell(row=row, column=3, value=safe_excel_value(str(entry["question_text"])))
        value = entry["std"]
        sheet.cell(row=row, column=4, value=None if value != value else float(value))
        row += 1
    data_end_row = row - 1

    chart = BarChart()
    chart.type = "col"
    chart.title = "Inter-Rater Spread"
    chart.y_axis.title = "Std Dev (lower = more agreement)"
    data_ref = Reference(sheet, min_col=4, max_col=4, min_row=1, max_row=data_end_row)
    cats_ref = Reference(sheet, min_col=3, min_row=2, max_row=data_end_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    sheet.add_chart(chart, "F1")


def _format_stat_and_ci(value: object, ci_low: object, ci_high: object) -> str:
    """ "0.612 (95% CI 0.481-0.743)" - or a placeholder when the statistic
    couldn't be computed (see qicom.session.evaluation_statistics.
    MIN_SUBJECTS_FOR_RELIABILITY: too little complete-case data for that
    question)."""
    if value is None or (isinstance(value, float) and value != value):
        return "n/a (not enough data)"
    if ci_low is None or ci_high is None:
        return f"{value:.3f}"
    return f"{value:.3f} (95% CI {ci_low:.3f}-{ci_high:.3f})"


def _write_reliability_sheet(workbook: Workbook, reliability_df: pd.DataFrame) -> None:
    """Cohen's kappa (2 evaluators) / Fleiss' kappa (3+) and the ICC(2,1)
    intraclass correlation coefficient, each with its own 95% bootstrap
    CI (see qicom.session.evaluation_statistics/evaluation_aggregation.
    inter_rater_reliability) - one row per (page, comparison, question).
    """
    sheet = workbook.create_sheet("Inter-Rater Reliability")
    header = [
        "Page",
        "Comparison",
        "Question",
        "Evaluators",
        "Subjects (complete cases)",
        "Kappa (type, 95% CI)",
        "ICC(2,1) (95% CI)",
    ]
    for col, title in enumerate(header, start=1):
        sheet.cell(row=1, column=col, value=title)
    row = 2
    for _, entry in reliability_df.iterrows():
        kappa_type = entry["kappa_type"]
        kappa_label = (
            _format_stat_and_ci(entry["kappa"], entry["kappa_ci_low"], entry["kappa_ci_high"])
            if kappa_type is None
            else f"{kappa_type.title()}: "
            + _format_stat_and_ci(entry["kappa"], entry["kappa_ci_low"], entry["kappa_ci_high"])
        )
        icc_label = _format_stat_and_ci(entry["icc"], entry["icc_ci_low"], entry["icc_ci_high"])
        sheet.cell(row=row, column=1, value=safe_excel_value(str(entry["page_title"])))
        sheet.cell(row=row, column=2, value=safe_excel_value(str(entry["comparison_role"])))
        sheet.cell(row=row, column=3, value=safe_excel_value(str(entry["question_text"])))
        sheet.cell(row=row, column=4, value=int(entry["n_evaluators"]))
        sheet.cell(row=row, column=5, value=int(entry["n_subjects"]))
        sheet.cell(row=row, column=6, value=kappa_label)
        sheet.cell(row=row, column=7, value=icc_label)
        row += 1


def _write_ordinal_reliability_sheet(workbook: Workbook, reliability_df: pd.DataFrame) -> None:
    sheet = workbook.create_sheet("Ordinal Reliability")
    header = [
        "Page",
        "Comparison",
        "Question",
        "Evaluators",
        "Ratings",
        "Missing",
        "PRIMARY: Krippendorff ordinal alpha (95% CI)",
        "ADVANCED: Gwet AC2 (95% CI)",
        "Method limitations",
        "AC2 formula",
    ]
    sheet.append(header)
    for _, entry in reliability_df.iterrows():
        sheet.append(
            [
                safe_excel_value(str(entry["page_title"])),
                safe_excel_value(str(entry["comparison_role"])),
                safe_excel_value(str(entry["question_text"])),
                int(entry["n_evaluators"]),
                int(entry.get("n_ratings", 0)),
                int(entry.get("n_missing", 0)),
                _format_stat_and_ci(
                    entry.get("ordinal_alpha"),
                    entry.get("alpha_ci_low"),
                    entry.get("alpha_ci_high"),
                ),
                _format_stat_and_ci(
                    entry.get("ac2"), entry.get("ac2_ci_low"), entry.get("ac2_ci_high")
                ),
                safe_excel_value(entry.get("method_limitations")),
                safe_excel_value(entry.get("ac2_method")),
            ]
        )


def _write_auc_sheet(workbook: Workbook, auc_df: pd.DataFrame) -> None:
    """Paired probability of superiority between comparison images,
    per question (see evaluation_aggregation.pairwise_auc_by_comparison) -
    only written when the evaluation actually has 2+ comparison roles to
    pair up (a single-comparison evaluation has nothing to compare)."""
    sheet = workbook.create_sheet("AUC Comparisons")
    header = [
        "Page",
        "Question",
        "Comparison A",
        "Comparison B",
        "Paired superiority (95% CI)",
        "Wilcoxon paired p-value",
    ]
    for col, title in enumerate(header, start=1):
        sheet.cell(row=1, column=col, value=title)
    row = 2
    for _, entry in auc_df.iterrows():
        p_value = entry["p_value"]
        sheet.cell(row=row, column=1, value=safe_excel_value(str(entry["page_title"])))
        sheet.cell(row=row, column=2, value=safe_excel_value(str(entry["question_text"])))
        sheet.cell(row=row, column=3, value=safe_excel_value(str(entry["comparison_role_a"])))
        sheet.cell(row=row, column=4, value=safe_excel_value(str(entry["comparison_role_b"])))
        sheet.cell(
            row=row,
            column=5,
            value=_format_stat_and_ci(entry["auc"], entry["auc_ci_low"], entry["auc_ci_high"]),
        )
        sheet.cell(row=row, column=6, value=None if p_value is None else float(p_value))
        row += 1


def _write_methodology_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Methodology")
    notes = [
        (
            "Reliability",
            "The primary ordinal agreement coefficient is Krippendorff's alpha, with Gwet AC2 "
            "labelled advanced. AC2 v2 uses mean unit category proportions (including singleton "
            "units) and Gwet's weighted chance agreement; v1 is not numerically comparable. "
            "Effective ratings and missing values are shown. Legacy weighted "
            "kappa and ICC(A,1) absolute agreement are retained for continuity; repeated copies "
            "are excluded from inter-rater estimates.",
        ),
        (
            "Paired comparisons",
            "Scores are averaged across slices within evaluator-by-case cells. Confidence "
            "intervals independently resample evaluator and case clusters. Reports include "
            "median paired difference, probability of superiority, paired Wilcoxon raw p-values, "
            "and Holm family-wise corrected p-values.",
        ),
        (
            "Quality control",
            "Timing, straightlining, missingness, and leave-one-out influence are descriptive "
            "review flags only. QICOM never automatically excludes or recentres an evaluator.",
        ),
        (
            "Interpretation",
            "These statistics support research and quality-control workflows and are not a "
            "substitute for a study design with independent scans and a prespecified analysis.",
        ),
    ]
    sheet.cell(row=1, column=1, value="Analysis")
    sheet.cell(row=1, column=2, value="Method and limitations")
    for row, (topic, note) in enumerate(notes, start=2):
        sheet.cell(row=row, column=1, value=topic)
        sheet.cell(row=row, column=2, value=note)


def _write_per_question_sheets(workbook: Workbook, answers_df: pd.DataFrame) -> None:
    """One sheet per question - mean score per slice, one series per
    comparison image, mirroring export_line_profiles_to_workbook's own
    "one sheet per item, deterministic safe name, real title never the
    sheet name" convention. A LineChart (x=slice) for a multi-slice
    evaluation, a bar chart otherwise - matching how DualImageView's own
    2D-vs-3D distinction is handled everywhere else in this app. Sheet
    names are `P<page_id>_Q<n>` (n resets per page) - question_id is
    globally unique (see EvaluationPage's own docstring) so this is never
    strictly needed for uniqueness, but it keeps a multi-page evaluation's
    sheets grouped and scannable rather than one long undifferentiated
    Q1..Qn run."""
    if answers_df.empty:
        return

    questions = (
        answers_df[["page_id", "page_title", "question_id", "question_text"]]
        .drop_duplicates()
        .sort_values(["page_id", "question_id"])
    )
    question_index_within_page: dict[object, int] = {}
    for _, question in questions.iterrows():
        page_id = question["page_id"]
        question_index_within_page[page_id] = question_index_within_page.get(page_id, 0) + 1
        question_index = question_index_within_page[page_id]

        question_df = answers_df[answers_df["question_id"] == question["question_id"]]
        sheet_name = safe_sheet_name(f"P{page_id}_Q{question_index}", workbook.sheetnames)
        display_title = f"{question['page_title']} - {question['question_text']}"
        if not question_df["score"].notna().any():
            # Categorical, multiple-choice, ranking and open-text questions
            # must never be forced into a meaningless numeric mean chart.
            # Give them a clean per-response detail sheet instead.
            sheet = workbook.create_sheet(sheet_name)
            sheet.cell(row=1, column=1, value=safe_excel_value(display_title))
            detail_columns = [
                "evaluator_id",
                "evaluator_display_name",
                "comparison_role",
                "slice_index",
                "response",
                "free_text",
            ]
            detail = question_df[detail_columns]
            for column, heading in enumerate(detail_columns, start=1):
                sheet.cell(row=2, column=column, value=heading)
            for row_index, values in enumerate(detail.itertuples(index=False, name=None), start=3):
                for column, value in enumerate(values, start=1):
                    sheet.cell(row=row_index, column=column, value=safe_excel_value(value))
            continue
        per_slice = (
            question_df.groupby(["comparison_role", "slice_index"])["score"].mean().reset_index()
        )
        if per_slice.empty:
            continue

        sheet = workbook.create_sheet(sheet_name)
        sheet.cell(row=1, column=1, value=safe_excel_value(display_title))

        roles = sorted(str(role) for role in per_slice["comparison_role"].unique())
        slices = sorted(int(s) for s in per_slice["slice_index"].unique())
        by_role_slice = {
            (str(row["comparison_role"]), int(row["slice_index"])): row["score"]
            for _, row in per_slice.iterrows()
        }

        header_row = 2
        sheet.cell(row=header_row, column=1, value="Slice")
        for col, role in enumerate(roles, start=2):
            sheet.cell(row=header_row, column=col, value=safe_excel_value(role))
        for row_offset, slice_index in enumerate(slices):
            data_row = header_row + 1 + row_offset
            sheet.cell(row=data_row, column=1, value=slice_index)
            for col, role in enumerate(roles, start=2):
                value = by_role_slice.get((role, slice_index))
                cell_value = None if value is None or value != value else float(value)
                sheet.cell(row=data_row, column=col, value=cell_value)
        data_start_row = header_row
        data_end_row = header_row + len(slices)

        is_multi_slice = len(slices) > 1
        chart: LineChart | BarChart = LineChart() if is_multi_slice else BarChart()
        if isinstance(chart, BarChart):
            chart.type = "col"
        chart.title = safe_excel_value(display_title)
        chart.x_axis.title = "Slice"
        chart.y_axis.title = "Mean Score"
        n_series = len(roles)
        data_ref = Reference(
            sheet, min_col=2, max_col=1 + n_series, min_row=data_start_row, max_row=data_end_row
        )
        cats_ref = Reference(sheet, min_col=1, min_row=data_start_row + 1, max_row=data_end_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        anchor_column = get_column_letter(1 + n_series + 2)
        sheet.add_chart(chart, f"{anchor_column}{data_start_row}")


def write_own_answers_report(xlsx_path: Path, answers_df: pd.DataFrame) -> None:
    """Write one evaluator's answers workbook (see qicom.gui.evaluation_export)."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = RAW_ANSWERS_SHEET
    _write_dataframe(sheet, answers_df)
    save_workbook_atomically(workbook, xlsx_path)


def write_aggregated_report(xlsx_path: Path, answers_df: pd.DataFrame) -> None:
    """Write the aggregated multi-evaluator workbook (see qicom.gui.evaluation_export)."""
    workbook = Workbook()
    raw_sheet = workbook.active
    assert raw_sheet is not None
    raw_sheet.title = RAW_ANSWERS_SHEET
    _write_dataframe(raw_sheet, answers_df)

    summary_groups = _bar_groups_from_summary(aggregate_by_question(answers_df))
    if summary_groups:
        export_native_charts_to_workbook(workbook, [], summary_groups, "Summary")

    by_evaluator_df = aggregate_by_evaluator_and_question(answers_df)
    by_evaluator_groups = _bar_groups_from_by_evaluator(by_evaluator_df)
    if by_evaluator_groups:
        export_native_charts_to_workbook(workbook, [], by_evaluator_groups, "By Evaluator")

    spread_df = inter_rater_spread(answers_df)
    if not spread_df.empty:
        _write_spread_sheet(workbook, spread_df)

    reliability_df = inter_rater_reliability(answers_df)
    if not reliability_df.empty:
        _write_reliability_sheet(workbook, reliability_df)
        _write_ordinal_reliability_sheet(workbook, reliability_df)

    auc_df = pairwise_auc_by_comparison(answers_df)
    if not auc_df.empty:
        _write_auc_sheet(workbook, auc_df)

    supplementary = [
        ("Ordinal Descriptives", descriptive_by_question(answers_df)),
        ("Category Distributions", category_distribution(answers_df)),
        ("Likert Distributions", likert_distribution(answers_df)),
        ("Ranking Concordance", ranking_concordance(answers_df)),
        ("Paired Study Analysis", paired_comparison_analysis(answers_df)),
        ("Intra-Rater Repeats", intra_rater_repeat_reliability(answers_df)),
        ("Response QC Flags", evaluator_quality_flags(answers_df)),
        ("Precision Completeness", precision_and_completeness(answers_df)),
    ]
    open_text = filter_open_text(answers_df)
    if "question_type" in open_text.columns:
        open_text = open_text[
            open_text["question_type"].eq("text")
            | open_text["free_text"].fillna("").astype(str).str.len().gt(0)
        ]
    supplementary.append(("Open Text", open_text))
    for sheet_name, frame in supplementary:
        if frame.empty:
            continue
        sheet = workbook.create_sheet(sheet_name)
        _write_dataframe(sheet, frame)

    _write_per_question_sheets(workbook, answers_df)
    if not answers_df.empty:
        _write_methodology_sheet(workbook)
    save_workbook_atomically(workbook, xlsx_path)


__all__ = ["write_aggregated_report", "write_own_answers_report"]
