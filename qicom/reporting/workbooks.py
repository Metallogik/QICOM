"""Reusable native-Excel exporters, independent of GUI controllers."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

from qicom.atomic_files import commit_temporary, temporary_sibling
from qicom.i18n import translator
from qicom.metrics.types import ImageArray
from qicom.reporting.excel_safety import safe_excel_value, safe_sheet_name

_STAT_TILE_DIRECTION_KEYS = {
    "higher": "metric.direction_higher",
    "lower": "metric.direction_lower",
    "target": "metric.direction_target",
    "none": "metric.direction_none",
}


def _resolve_workbook(target: Path | Workbook) -> tuple[Workbook, Path | None]:
    if isinstance(target, Workbook):
        return target, None
    return load_workbook(target), target


def _save_owned_workbook(workbook: Workbook, path: Path | None) -> None:
    if path is not None:
        save_workbook_atomically(workbook, path)


def save_workbook_atomically(workbook: Workbook, path: Path) -> None:
    """Save once, then atomically replace the requested workbook path."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = temporary_sibling(path)
    try:
        workbook.save(temporary)
        commit_temporary(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def export_native_charts_to_workbook(
    target: Path | Workbook,
    stat_tiles: list[tuple[str, float, str, float | None, str]],
    groups: list[tuple[str, list[str], dict[str, list[float]]]],
    sheet_name: str,
) -> None:
    """Add editable charts to a workbook object or an existing file."""
    workbook, owned_path = _resolve_workbook(target)
    sheet = workbook.create_sheet(safe_sheet_name(sheet_name, workbook.sheetnames))

    row = 1
    if stat_tiles:
        sheet.cell(row=row, column=1, value=translator.tr("results_table.column_metric_name"))
        sheet.cell(row=row, column=2, value=translator.tr("results_table.column_value"))
        row += 1
        for metric_name, value, better_direction, target_value, description in stat_tiles:
            hint_key = _STAT_TILE_DIRECTION_KEYS.get(better_direction, "metric.direction_none")
            sheet.cell(row=row, column=1, value=safe_excel_value(metric_name))
            sheet.cell(row=row, column=2, value=value)
            sheet.cell(
                row=row,
                column=3,
                value=safe_excel_value(translator.tr(hint_key, target_value=target_value)),
            )
            # A tile's longer description (only the QICOM Index has one so
            # far, carrying its "experimental / not clinically validated"
            # caveat) is hover-only on screen, so a sheet has to spell it
            # out: a column of its own, left empty for every other tile.
            if description:
                sheet.cell(row=row, column=4, value=safe_excel_value(description))
            row += 1
        row += 1

    for title, categories, series in groups:
        data_start_row = row
        series_names = list(series.keys())
        sheet.cell(row=row, column=1, value=translator.tr("results_table.column_category"))
        for col_offset, name in enumerate(series_names, start=2):
            sheet.cell(row=row, column=col_offset, value=safe_excel_value(name))
        row += 1
        for row_offset, category in enumerate(categories):
            sheet.cell(row=row, column=1, value=safe_excel_value(category))
            for col_offset, name in enumerate(series_names, start=2):
                values = series[name]
                cell_value = values[row_offset] if row_offset < len(values) else None
                sheet.cell(
                    row=row,
                    column=col_offset,
                    value=None if cell_value != cell_value else cell_value,
                )
            row += 1
        data_end_row = row - 1

        chart = BarChart()
        chart.type = "col"
        chart.title = safe_excel_value(title)
        n_series = len(series_names)
        data_ref = Reference(
            sheet, min_col=2, max_col=1 + n_series, min_row=data_start_row, max_row=data_end_row
        )
        cats_ref = Reference(sheet, min_col=1, min_row=data_start_row + 1, max_row=data_end_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        sheet.add_chart(chart, f"{get_column_letter(1 + n_series + 2)}{data_start_row}")
        row = data_end_row + 16

    _save_owned_workbook(workbook, owned_path)


def export_distributions_data_to_workbook(
    target: Path | Workbook,
    sheet_name: str,
    groups: list[tuple[str, list[str], dict[str, list[float]]]],
) -> None:
    """Add raw per-slice distribution tables to an existing workbook."""
    if not groups:
        return

    workbook, owned_path = _resolve_workbook(target)
    sheet = workbook.create_sheet(safe_sheet_name(sheet_name, workbook.sheetnames))
    row = 1
    for title, categories, distributions in groups:
        sheet.cell(row=row, column=1, value=safe_excel_value(title))
        row += 1
        for col_offset, label in enumerate(categories, start=1):
            sheet.cell(row=row, column=col_offset, value=safe_excel_value(label))
        row += 1
        data_start_row = row
        max_len = max((len(distributions.get(label, [])) for label in categories), default=0)
        for col_offset, label in enumerate(categories, start=1):
            for offset, value in enumerate(distributions.get(label, [])):
                sheet.cell(
                    row=data_start_row + offset,
                    column=col_offset,
                    value=None if value != value else value,
                )
        row = data_start_row + max_len + 2
    _save_owned_workbook(workbook, owned_path)


def export_line_profiles_to_workbook(
    xlsx_path: Path,
    profiles: list[tuple[str, dict[str, ImageArray]]],
    sheet_name_prefix: str = "Line",
) -> None:
    """Write line-profile values and editable charts to a new workbook."""
    workbook = Workbook()
    created_any = False

    for index, (line_name, series) in enumerate(profiles, start=1):
        labels = list(series.keys())
        if not labels:
            continue
        sheet = workbook.create_sheet(
            safe_sheet_name(f"{sheet_name_prefix} {index}", workbook.sheetnames)
        )
        created_any = True
        reference_label = labels[0]
        reference_values = series[reference_label]
        comparison_labels = labels[1:]

        sheet.cell(row=1, column=1, value=safe_excel_value(line_name))
        header_row = 2
        header = [
            "Distance (px)",
            *labels,
            *[f"|{label} - {reference_label}|" for label in comparison_labels],
        ]
        for col, title in enumerate(header, start=1):
            sheet.cell(row=header_row, column=col, value=safe_excel_value(title))

        n_points = max((len(values) for values in series.values()), default=0)
        for offset in range(n_points):
            data_row = header_row + 1 + offset
            sheet.cell(row=data_row, column=1, value=offset)
            for col_offset, label in enumerate(labels, start=2):
                values = series[label]
                if offset < len(values):
                    value = float(values[offset])
                    sheet.cell(
                        row=data_row,
                        column=col_offset,
                        value=None if value != value else value,
                    )
            diff_start_col = 2 + len(labels)
            for col_offset, label in enumerate(comparison_labels, start=diff_start_col):
                values = series[label]
                if offset < len(values) and offset < len(reference_values):
                    value = float(abs(values[offset] - reference_values[offset]))
                    sheet.cell(
                        row=data_row,
                        column=col_offset,
                        value=None if value != value else value,
                    )

        if n_points == 0:
            continue
        data_end_row = header_row + n_points
        chart = LineChart()
        chart.title = safe_excel_value(line_name)
        chart.x_axis.title = "Distance (px)"
        chart.y_axis.title = "Pixel value"
        n_series = len(header) - 1
        data_ref = Reference(
            sheet, min_col=2, max_col=1 + n_series, min_row=header_row, max_row=data_end_row
        )
        cats_ref = Reference(sheet, min_col=1, min_row=header_row + 1, max_row=data_end_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        sheet.add_chart(chart, f"{get_column_letter(1 + n_series + 2)}{header_row}")

    if created_any:
        workbook.remove(workbook.active)
    save_workbook_atomically(workbook, xlsx_path)
