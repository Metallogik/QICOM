"""Tests for qicom.gui.evaluation_export (export_aggregated_report)."""

from __future__ import annotations

from pathlib import Path
from typing import Any
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from qicom.gui.evaluation_export import (
    export_aggregated_report,
    export_own_answers_report,
)
from qicom.session.evaluation_aggregation import build_answers_dataframe
from qicom.session.evaluation_folder import EvaluationAnswer, EvaluationPage, EvaluationQuestion
from qicom.session.evaluation_package import EvaluationPackage, EvaluatorAnswers
from qicom.session.project_folder import ProjectPanelRecord


def _comparison_panels() -> list[ProjectPanelRecord]:
    return [
        ProjectPanelRecord(role, f"images/{role}.png", False, None)
        for role in ("comparison", "comparison_2")
    ]


def _chart_title_text(chart: object) -> str:
    return chart.title.tx.rich.p[0].r[0].t  # type: ignore[attr-defined]


def _make_package() -> EvaluationPackage:
    return EvaluationPackage(
        folder=Path("."),
        version=2,
        test_id="t",
        created_at="",
        creator_note=None,
        pages=[
            EvaluationPage(
                page_id=1,
                title="Page 1",
                questions=[
                    EvaluationQuestion(1, "Overall difference?"),
                    EvaluationQuestion(2, "Noise?"),
                ],
                panels=_comparison_panels(),
            )
        ],
        warnings=[],
    )


def _make_two_evaluators_two_slices() -> list[EvaluatorAnswers]:
    return [
        EvaluatorAnswers(
            test_id="t",
            evaluator_id="A",
            evaluator_display_name="Alice",
            answers=[
                EvaluationAnswer("comparison", 1, 0, 4),
                EvaluationAnswer("comparison", 1, 1, 3),
                EvaluationAnswer("comparison", 2, 0, 5),
                EvaluationAnswer("comparison", 2, 1, 4),
            ],
            completed=True,
            exported_at="",
        ),
        EvaluatorAnswers(
            test_id="t",
            evaluator_id="B",
            evaluator_display_name="Bob",
            answers=[
                EvaluationAnswer("comparison", 1, 0, 2),
                EvaluationAnswer("comparison", 1, 1, None),
                EvaluationAnswer("comparison", 2, 0, 3),
                EvaluationAnswer("comparison", 2, 1, 3),
            ],
            completed=True,
            exported_at="",
        ),
    ]


class TestExportAggregatedReport:
    def test_complete_report_saves_the_workbook_once(self, tmp_path: Path) -> None:
        df = build_answers_dataframe(_make_package(), _make_two_evaluators_two_slices())
        original_save = Workbook.save
        saves = 0

        def counted_save(workbook: Workbook, filename: Any) -> None:
            nonlocal saves
            saves += 1
            original_save(workbook, filename)

        with patch.object(Workbook, "save", counted_save):
            export_aggregated_report(tmp_path / "one-save.xlsx", df)

        assert saves == 1

    def test_writes_every_expected_sheet_with_real_charts_not_images(self, tmp_path: Path) -> None:
        df = build_answers_dataframe(_make_package(), _make_two_evaluators_two_slices())
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        assert set(workbook.sheetnames) >= {
            "Raw Answers",
            "Summary",
            "By Evaluator",
            "Inter-Rater Spread",
            "P1_Q1",
            "P1_Q2",
        }
        for name in ("Summary", "By Evaluator", "Inter-Rater Spread", "P1_Q1", "P1_Q2"):
            sheet = workbook[name]
            assert len(sheet._charts) >= 1
            assert len(sheet._images) == 0

    def test_raw_answers_sheet_has_the_expected_header_and_a_nan_as_a_blank_cell(
        self, tmp_path: Path
    ) -> None:
        df = build_answers_dataframe(_make_package(), _make_two_evaluators_two_slices())
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        raw = workbook["Raw Answers"]
        assert [cell.value for cell in raw[1]] == list(df.columns)
        # Bob's skipped slice-1 answer to Q1 must be a blank cell, not the
        # literal string "nan" or any other misleading marker.
        skip_row = next(
            row
            for row in raw.iter_rows(min_row=2, values_only=True)
            if row[0] == "B" and row[5] == 1 and row[7] == 1
        )
        score_column = list(df.columns).index("score")
        assert skip_row[score_column] is None

    def test_structured_answers_are_exported_as_readable_results_not_blank_charts(
        self, tmp_path: Path
    ) -> None:
        package = EvaluationPackage(
            folder=Path("."),
            version=3,
            test_id="structured",
            created_at="",
            creator_note=None,
            pages=[
                EvaluationPage(
                    1,
                    "Image quality",
                    [
                        EvaluationQuestion(
                            1,
                            "Visible artefacts?",
                            response_type="multiple_choice",
                            options=("Motion", "Noise", "Aliasing"),
                            allow_comment=True,
                        )
                    ],
                    [],
                )
            ],
            warnings=[],
        )
        evaluators = [
            EvaluatorAnswers(
                test_id="structured",
                evaluator_id="A",
                evaluator_display_name="Alice",
                answers=[
                    EvaluationAnswer(
                        "reference", 1, 0, None, "Mostly peripheral", 1, ("Motion", "Noise")
                    )
                ],
                completed=True,
                exported_at="",
            )
        ]
        dataframe = build_answers_dataframe(package, evaluators)
        xlsx_path = tmp_path / "structured.xlsx"

        export_aggregated_report(xlsx_path, dataframe)

        workbook = load_workbook(xlsx_path)
        assert "Summary" not in workbook.sheetnames
        details = workbook["P1_Q1"]
        assert [cell.value for cell in details[2]] == [
            "evaluator_id",
            "evaluator_display_name",
            "comparison_role",
            "slice_index",
            "response",
            "free_text",
        ]
        assert details.cell(3, 5).value == "Motion → Noise"
        assert details.cell(3, 6).value == "Mostly peripheral"
        assert details._charts == []

    def test_q1_sheet_uses_a_line_chart_for_a_multi_slice_evaluation(self, tmp_path: Path) -> None:
        from openpyxl.chart import LineChart

        df = build_answers_dataframe(_make_package(), _make_two_evaluators_two_slices())
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        chart = workbook["P1_Q1"]._charts[0]
        assert isinstance(chart, LineChart)

    def test_single_slice_evaluation_uses_a_bar_chart_per_question(self, tmp_path: Path) -> None:
        from openpyxl.chart import BarChart

        single_slice_evaluators = [
            EvaluatorAnswers(
                test_id="t",
                evaluator_id="A",
                evaluator_display_name="Alice",
                answers=[EvaluationAnswer("comparison", 1, 0, 4)],
                completed=True,
                exported_at="",
            )
        ]
        df = build_answers_dataframe(_make_package(), single_slice_evaluators)
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        chart = workbook["P1_Q1"]._charts[0]
        assert isinstance(chart, BarChart)

    def test_an_empty_dataframe_still_produces_a_readable_workbook(self, tmp_path: Path) -> None:
        df = build_answers_dataframe(_make_package(), [])
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        assert workbook.sheetnames == ["Raw Answers"]


class TestMultiPageReport:
    def _make_two_page_package(self) -> EvaluationPackage:
        return EvaluationPackage(
            folder=Path("."),
            version=2,
            test_id="t",
            created_at="",
            creator_note=None,
            pages=[
                EvaluationPage(
                    page_id=1,
                    title="Baseline",
                    questions=[EvaluationQuestion(1, "Q1")],
                    panels=_comparison_panels(),
                ),
                EvaluationPage(
                    page_id=2,
                    title="Follow-up",
                    questions=[EvaluationQuestion(2, "Q2")],
                    panels=_comparison_panels(),
                ),
            ],
            warnings=[],
        )

    def test_each_page_gets_its_own_sheet_and_the_titles_stay_distinct(
        self, tmp_path: Path
    ) -> None:
        answers = [
            EvaluatorAnswers(
                test_id="t",
                evaluator_id="A",
                evaluator_display_name="Alice",
                answers=[
                    EvaluationAnswer("comparison", 1, 0, 4, page_id=1),
                    EvaluationAnswer("comparison", 2, 0, 2, page_id=2),
                ],
                completed=True,
                exported_at="",
            )
        ]
        df = build_answers_dataframe(self._make_two_page_package(), answers)
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        assert "P1_Q1" in workbook.sheetnames
        assert "P2_Q1" in workbook.sheetnames
        assert workbook["P1_Q1"]["A1"].value == "Baseline - Q1"
        assert workbook["P2_Q1"]["A1"].value == "Follow-up - Q2"

    def test_summary_keeps_same_named_roles_on_different_pages_in_separate_chart_groups(
        self, tmp_path: Path
    ) -> None:
        # Both pages use the "comparison" role for a completely different
        # image - the Summary chart grouping must not merge them (see
        # _bar_groups_from_summary's own docstring).
        answers = [
            EvaluatorAnswers(
                test_id="t",
                evaluator_id="A",
                evaluator_display_name="Alice",
                answers=[
                    EvaluationAnswer("comparison", 1, 0, 4, page_id=1),
                    EvaluationAnswer("comparison", 2, 0, 2, page_id=2),
                ],
                completed=True,
                exported_at="",
            )
        ]
        df = build_answers_dataframe(self._make_two_page_package(), answers)
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        summary_sheet = workbook["Summary"]
        chart_titles = {_chart_title_text(chart) for chart in summary_sheet._charts}
        assert chart_titles == {"Baseline - comparison", "Follow-up - comparison"}


def _make_three_evaluators_five_slices_two_roles() -> list[EvaluatorAnswers]:
    """Enough complete-case data (3 evaluators x 5 slices) for Inter-Rater
    Reliability to actually compute something, across two comparison
    roles so AUC Comparisons has a pair to report on too."""
    comparison_scores = {"A": [1, 1, 1, 2, 2], "B": [1, 1, 2, 2, 2], "C": [1, 1, 1, 1, 2]}
    comparison_2_scores = {"A": [4, 5, 4, 5, 5], "B": [5, 5, 4, 5, 4], "C": [4, 4, 5, 5, 5]}
    evaluators = []
    for evaluator_id in ("A", "B", "C"):
        answers = [
            EvaluationAnswer("comparison", 1, i, score)
            for i, score in enumerate(comparison_scores[evaluator_id])
        ] + [
            EvaluationAnswer("comparison_2", 1, i, score)
            for i, score in enumerate(comparison_2_scores[evaluator_id])
        ]
        evaluators.append(
            EvaluatorAnswers(
                test_id="t",
                evaluator_id=evaluator_id,
                evaluator_display_name=evaluator_id,
                answers=answers,
                completed=True,
                exported_at="",
            )
        )
    return evaluators


class TestInterRaterReliabilitySheet:
    def test_sheet_is_written_with_expected_header_and_kappa_type(self, tmp_path: Path) -> None:
        df = build_answers_dataframe(
            _make_package(), _make_three_evaluators_five_slices_two_roles()
        )
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        assert "Inter-Rater Reliability" in workbook.sheetnames
        sheet = workbook["Inter-Rater Reliability"]
        header = [cell.value for cell in sheet[1]]
        assert header == [
            "Page",
            "Comparison",
            "Question",
            "Evaluators",
            "Subjects (complete cases)",
            "Kappa (type, 95% CI)",
            "ICC(2,1) (95% CI)",
        ]
        kappa_cells = [sheet.cell(row=r, column=6).value for r in range(2, sheet.max_row + 1)]
        assert any(cell.startswith("Fleiss:") for cell in kappa_cells)

    def test_reports_not_enough_data_rather_than_omitting_the_question(
        self, tmp_path: Path
    ) -> None:
        # Only 1-2 complete-case slices per question here (see
        # MIN_SUBJECTS_FOR_RELIABILITY) - too little to compute kappa/ICC
        # from, but the question still gets a row, saying so explicitly,
        # rather than silently vanishing from the report.
        df = build_answers_dataframe(_make_package(), _make_two_evaluators_two_slices())
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        sheet = workbook["Inter-Rater Reliability"]
        kappa_cells = [sheet.cell(row=r, column=6).value for r in range(2, sheet.max_row + 1)]
        assert all("not enough data" in cell for cell in kappa_cells)

    def test_sheet_absent_when_the_report_has_no_answers_at_all(self, tmp_path: Path) -> None:
        df = build_answers_dataframe(_make_package(), [])
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        assert "Inter-Rater Reliability" not in workbook.sheetnames


class TestAucComparisonsSheet:
    def test_sheet_is_written_with_expected_header(self, tmp_path: Path) -> None:
        df = build_answers_dataframe(
            _make_package(), _make_three_evaluators_five_slices_two_roles()
        )
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        assert "AUC Comparisons" in workbook.sheetnames
        sheet = workbook["AUC Comparisons"]
        header = [cell.value for cell in sheet[1]]
        assert header == [
            "Page",
            "Question",
            "Comparison A",
            "Comparison B",
            "Paired superiority (95% CI)",
            "Wilcoxon paired p-value",
        ]
        # Three evaluators are not enough for a two-sided exact paired
        # Wilcoxon result below 0.05. The former pooled-slice test reported a
        # spuriously significant result by treating correlated slices as
        # independent observations.
        row2 = [sheet.cell(row=2, column=c).value for c in range(1, 7)]
        assert row2[2] == "comparison"
        assert row2[3] == "comparison_2"
        assert row2[5] == 0.25

    def test_not_written_with_only_one_comparison_role(self, tmp_path: Path) -> None:
        df = build_answers_dataframe(_make_package(), _make_two_evaluators_two_slices())
        xlsx_path = tmp_path / "report.xlsx"

        export_aggregated_report(xlsx_path, df)

        workbook = load_workbook(xlsx_path)
        assert "AUC Comparisons" not in workbook.sheetnames


class TestExportOwnAnswersReport:
    """One evaluator's own file - the thing they send to whoever
    aggregates. Everything the aggregation step computes across evaluators
    is deliberately absent from it."""

    def test_writes_only_the_answers_sheet(self, tmp_path: Path) -> None:
        package = _make_package()
        alice = _make_two_evaluators_two_slices()[0]
        xlsx_path = tmp_path / "alice.xlsx"

        export_own_answers_report(xlsx_path, build_answers_dataframe(package, [alice]))

        workbook = load_workbook(xlsx_path)
        # Not Summary/By Evaluator (a "mean" over one rater is that
        # rater's own score), not Inter-Rater Spread (blank) or
        # Inter-Rater Reliability ("n/a (not enough data)" on every row),
        # and not Methodology, which described analyses this file never
        # contained.
        assert workbook.sheetnames == ["Raw Answers"]

    def test_keeps_every_answer_row_intact(self, tmp_path: Path) -> None:
        package = _make_package()
        alice = _make_two_evaluators_two_slices()[0]
        xlsx_path = tmp_path / "alice.xlsx"

        export_own_answers_report(xlsx_path, build_answers_dataframe(package, [alice]))

        rows = list(load_workbook(xlsx_path)["Raw Answers"].iter_rows(values_only=True))
        assert rows[0][0] == "evaluator_id"
        assert len(rows) == 1 + len(alice.answers)
        assert [row[-2] for row in rows[1:]] == [4, 3, 5, 4]

    def test_the_sheet_name_is_the_one_the_aggregation_step_reads_back(
        self, tmp_path: Path
    ) -> None:
        """This file is an accepted input to the aggregation workspace (see
        qicom.session.evaluation_answer_files.read_answer_file), which
        looks the answers up by sheet name."""
        from qicom.session.evaluation_answer_files import RAW_ANSWERS_SHEET

        package = _make_package()
        alice = _make_two_evaluators_two_slices()[0]
        xlsx_path = tmp_path / "alice.xlsx"

        export_own_answers_report(xlsx_path, build_answers_dataframe(package, [alice]))

        assert load_workbook(xlsx_path).sheetnames == [RAW_ANSWERS_SHEET]

    def test_untrusted_text_is_still_neutralised(self, tmp_path: Path) -> None:
        """The formula-injection guard is not a property of the aggregated
        report - it has to hold for this leaner one too."""
        package = EvaluationPackage(
            folder=Path("."),
            version=2,
            test_id="t",
            created_at="",
            creator_note=None,
            pages=[
                EvaluationPage(
                    page_id=1,
                    title="=2+2",
                    questions=[EvaluationQuestion(1, "+question")],
                    panels=[ProjectPanelRecord("-comparison", "images/c.png", False, None)],
                )
            ],
            warnings=[],
        )
        evaluator = EvaluatorAnswers(
            test_id="t",
            evaluator_id="e",
            evaluator_display_name="@evaluator",
            answers=[EvaluationAnswer("-comparison", 1, 0, 4)],
            completed=True,
            exported_at="",
        )
        xlsx_path = tmp_path / "safe.xlsx"

        export_own_answers_report(xlsx_path, build_answers_dataframe(package, [evaluator]))

        raw = load_workbook(xlsx_path)["Raw Answers"]
        values = [raw.cell(row=2, column=column).value for column in (2, 4, 5, 7)]
        assert values == ["'@evaluator", "'=2+2", "'-comparison", "'+question"]


def test_untrusted_report_text_is_not_written_as_an_excel_formula(tmp_path: Path) -> None:
    package = EvaluationPackage(
        folder=Path("."),
        version=2,
        test_id="t",
        created_at="",
        creator_note=None,
        pages=[
            EvaluationPage(
                page_id=1,
                title="=2+2",
                questions=[EvaluationQuestion(1, "+question")],
                panels=[ProjectPanelRecord("-comparison", "images/c.png", False, None)],
            )
        ],
        warnings=[],
    )
    evaluator = EvaluatorAnswers(
        test_id="t",
        evaluator_id="e",
        evaluator_display_name="@evaluator",
        answers=[EvaluationAnswer("-comparison", 1, 0, 4)],
        completed=True,
        exported_at="",
    )
    xlsx_path = tmp_path / "safe.xlsx"

    export_aggregated_report(xlsx_path, build_answers_dataframe(package, [evaluator]))

    raw = load_workbook(xlsx_path)["Raw Answers"]
    values = [raw.cell(row=2, column=column).value for column in (2, 4, 5, 7)]
    assert values == ["'@evaluator", "'=2+2", "'-comparison", "'+question"]
    assert all(raw.cell(row=2, column=column).data_type == "s" for column in (2, 4, 5, 7))
