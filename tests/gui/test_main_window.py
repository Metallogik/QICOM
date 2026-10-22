"""Tests for qicom.gui.main_window (MainWindow) - requires a Qt platform.

Runs headlessly via QT_QPA_PLATFORM=offscreen, set in tests/gui/conftest.py.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import pytest
from message_box_stubs import stub_message_box
from PIL import Image
from PySide6.QtCore import Qt
from PySide6.QtGui import QCloseEvent
from PySide6.QtWidgets import QDialog, QDockWidget, QLabel, QMenuBar, QMessageBox
from pytestqt.qtbot import QtBot

from qicom.gui import main_window as main_window_module
from qicom.gui.error_reporting import explained_error
from qicom.gui.i18n import translator
from qicom.gui.main_window import MainWindow
from qicom.gui.theming import theme_manager
from qicom.gui.widgets.dual_image_view import ImagePanel
from qicom.gui.widgets.heatmap_panel import WHOLE_IMAGE_TARGET
from qicom.gui.widgets.toolbar import MODE_ROI_ELLIPSE
from qicom.io.dicom_anonymize import DicomDeidentificationError, UnverifiedDicomFile
from qicom.metrics.background import METHOD_MANUAL, BackgroundEstimate
from qicom.metrics.types import RESULT_COLUMNS
from qicom.session import recent_projects as recent_projects_module
from qicom.session.project_folder import ProjectFolderData
from qicom.session.recent_projects import load_recent_projects
from qicom.session.storage import SessionStorage


def _make_results() -> pd.DataFrame:
    rows = [
        {
            "metric_id": "mse",
            "metric_name": "MSE",
            "category": "reference_error",
            "image": "pair",
            "scope": "per-slice",
            "value": 12.5,
            "std": float("nan"),
            "n_slices": None,
            "better_direction": "lower",
            "target_value": 0.0,
        }
    ]
    return pd.DataFrame(rows, columns=RESULT_COLUMNS)


def _make_chartable_results() -> pd.DataFrame:
    """_make_results()'s single "pair" row has nothing to bar-chart against
    (see ChartsPanel's stat-tile design) - these two rows (a no_reference
    metric with both a reference and a comparison value) actually render a
    chart, for the export tests that need one to exist."""
    rows = [
        {
            "metric_id": "niqe",
            "metric_name": "NIQE",
            "category": "no_reference",
            "image": "reference",
            "scope": "per-slice",
            "value": 4.0,
            "std": float("nan"),
            "n_slices": None,
            "better_direction": "lower",
            "target_value": None,
        },
        {
            "metric_id": "niqe",
            "metric_name": "NIQE",
            "category": "no_reference",
            "image": "comparison",
            "scope": "per-slice",
            "value": 5.0,
            "std": float("nan"),
            "n_slices": None,
            "better_direction": "lower",
            "target_value": None,
        },
    ]
    return pd.DataFrame(rows, columns=RESULT_COLUMNS)


def _save_png(path: Path, seed: int) -> None:
    data = np.random.default_rng(seed).integers(0, 256, size=(16, 16), dtype=np.uint8)
    Image.fromarray(data).save(path)


@pytest.fixture(autouse=True)
def _fake_user_data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Recent-projects entries (qicom.session.recent_projects) live under
    platformdirs' real user-data directory by default - every test in this
    module must stay off it."""
    fake_dir = tmp_path / "platformdirs_home"
    monkeypatch.setattr(
        recent_projects_module.platformdirs, "user_data_dir", lambda *a, **k: str(fake_dir)
    )
    return fake_dir


@pytest.fixture
def window(qtbot: QtBot) -> MainWindow:
    translator.set_language("en")
    # In-memory storage: tests must never touch the real on-disk session DB.
    win = MainWindow(storage=SessionStorage(":memory:"))
    qtbot.addWidget(win)
    return win


class TestMainWindow:
    def test_constructs_without_error(self, window: MainWindow) -> None:
        assert window.windowTitle() == "QICOM - Quality Image Comparator"

    def test_view_toolbar_is_a_top_level_toolbar_above_the_images(self, window: MainWindow) -> None:
        assert window.toolBarArea(window.view_toolbar) == Qt.ToolBarArea.TopToolBarArea

    def test_roi_shape_toolbar_is_embedded_in_the_roi_panel_not_floating(
        self, window: MainWindow
    ) -> None:
        assert window.roi_manager_panel.isAncestorOf(window.roi_shape_toolbar)
        assert window.toolBarArea(window.roi_shape_toolbar) == Qt.ToolBarArea.NoToolBarArea

    def test_status_bar_shows_current_workflow_step(self, window: MainWindow) -> None:
        assert window.statusBar().currentMessage() == "Current step: Import"

    def test_freehand_add_button_arms_drawing_and_finished_path_creates_an_roi(
        self, window: MainWindow
    ) -> None:
        """End-to-end wiring check: RoiManagerPanel's Add button (Freehand
        mode selected) -> DualImageView.start_freehand_draw, and
        DualImageView.freehand_roi_ready -> AppController.add_freehand_roi
        - both connected directly by MainWindow.__init__, never exercised
        by a controller/dual_image_view-level test alone."""
        from qicom.gui.widgets.toolbar import MODE_ROI_FREEHAND

        window.dual_image_view.reference_panel.set_image(np.zeros((50, 50)))
        window.roi_manager_panel.set_active_shape_mode(MODE_ROI_FREEHAND)
        window.roi_manager_panel.set_images_loaded(True)

        window.roi_manager_panel._add_button.click()

        assert window.dual_image_view._freehand_active is True
        assert window.dual_image_view.reference_panel._freehand_capture is True

        window.dual_image_view.freehand_roi_ready.emit([(0.0, 0.0), (10.0, 0.0), (5.0, 10.0)])

        pairs = window.dual_image_view.list_roi_pairs()
        assert len(pairs) == 1
        assert pairs[0].shape_type == "freehand"

    def test_language_switch_retranslates_menus_live(self, window: MainWindow) -> None:
        translator.set_language("it")
        try:
            assert window._file_menu.title() == "File"
            assert window.statusBar().currentMessage() == "Fase corrente: Importa"
        finally:
            translator.set_language("en")

    def test_theme_action_switches_application_stylesheet(self, window: MainWindow) -> None:
        window._theme_actions["dark"].trigger()
        assert theme_manager.theme == "dark"

        window._theme_actions["light"].trigger()  # restore default for test isolation
        assert theme_manager.theme == "light"

    def test_number_format_action_switches_the_shared_preference(self, window: MainWindow) -> None:
        from qicom.gui.number_format import FULL, SCIENTIFIC, number_format_settings

        try:
            window._number_format_actions["scientific"].trigger()
            assert number_format_settings.notation == SCIENTIFIC

            window._number_format_actions["full"].trigger()
            assert number_format_settings.notation == FULL
        finally:
            number_format_settings.set_notation(FULL)  # restore default for test isolation


class TestInPageCalculateProgress:
    """Calculate reports into the Analyze page itself, not into a modal
    dialog floating over it - see AnalysisSetupPanel's own docstring."""

    def test_calculate_requested_opens_the_pages_own_progress_panel(
        self, window: MainWindow
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))

        window._on_calculate_requested(1.0)

        assert window.analysis_setup_panel.is_calculating()

    def test_no_modal_dialog_is_opened_for_a_calculation(self, window: MainWindow) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))

        window._on_calculate_requested(1.0)

        assert window._busy_dialog is None

    def test_calculating_from_the_results_table_brings_the_analyze_page_up(
        self, window: MainWindow
    ) -> None:
        """The results table has a Calculate button of its own; the progress
        now lives on the Analyze page, so that page has to be the one on
        screen or the run would report to nobody."""
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        window.workflow_rail.set_current_stage("results")

        window.results_panel._calculate_button.click()

        assert window.workflow_rail.current_stage == "analyze"
        assert window.analysis_setup_panel.is_calculating()

    def test_calculate_success_closes_the_progress_panel(self, window: MainWindow) -> None:
        window.analysis_setup_panel.set_calculating(True)

        window._on_calculate_success(_make_results())

        assert not window.analysis_setup_panel.is_calculating()

    def test_calculate_error_closes_the_progress_panel(self, window: MainWindow) -> None:
        window.analysis_setup_panel.set_calculating(True)

        window._on_calculate_error("boom")

        assert not window.analysis_setup_panel.is_calculating()

    def test_no_images_error_does_not_leave_the_page_stuck_calculating(
        self, window: MainWindow
    ) -> None:
        # Missing images make AppController.calculate() call on_error
        # synchronously (no background thread involved) - the panel must
        # not linger open in that case.
        window._on_calculate_requested(1.0)

        assert not window.analysis_setup_panel.is_calculating()

    def test_a_progress_update_switches_the_bar_to_determinate_mode(
        self, window: MainWindow
    ) -> None:
        window.analysis_setup_panel.set_calculating(True)

        window._on_calculate_progress(5, 20)

        bar = window.analysis_setup_panel._progress_bar
        assert bar.maximum() == 20
        assert bar.value() == 5

    def test_a_zero_total_is_ignored_rather_than_dividing_by_zero(self, window: MainWindow) -> None:
        window.analysis_setup_panel.set_calculating(True)

        window._on_calculate_progress(0, 0)  # must not raise

        assert window.analysis_setup_panel._progress_bar.maximum() == 0

    def test_a_step_update_names_what_is_running(self, window: MainWindow) -> None:
        window.analysis_setup_panel.set_calculating(True)

        window._on_calculate_step("Computing SSIM...")

        assert window.analysis_setup_panel._progress_step.text() == "Computing SSIM..."

    def test_a_stray_update_after_the_run_ended_is_a_noop(self, window: MainWindow) -> None:
        window.analysis_setup_panel.set_calculating(True)
        window.analysis_setup_panel.set_calculating(False)

        window._on_calculate_progress(1, 4)  # must not raise
        window._on_calculate_step("Computing SSIM...")  # must not raise

        assert not window.analysis_setup_panel.is_calculating()

    def test_cancelling_from_the_page_reports_it_and_ends_the_run(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.analysis_setup_panel.set_calculating(True)
        monkeypatch.setattr(window.controller, "cancel_calculation", lambda: True)

        window.analysis_setup_panel._progress_cancel.click()

        assert not window.analysis_setup_panel._progress_cancel.isEnabled()

        window._on_calculate_cancelled()

        assert not window.analysis_setup_panel.is_calculating()

    def test_a_refused_cancel_leaves_the_button_alone(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """cancel_calculation() returns False when there is nothing running
        to cancel - the panel must not then claim it is stopping."""
        window.analysis_setup_panel.set_calculating(True)
        monkeypatch.setattr(window.controller, "cancel_calculation", lambda: False)

        window.analysis_setup_panel._progress_cancel.click()

        assert window.analysis_setup_panel._progress_cancel.isEnabled()


class TestCalculationLock:
    """What the modal dialog used to enforce for free: while a run is in
    flight, nothing else in the window can be reached - see
    MainWindow._set_calculation_lock."""

    def test_the_other_workflow_steps_refuse_navigation_while_running(
        self, window: MainWindow
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))

        window._on_calculate_requested(1.0)

        assert window.workflow_rail.is_stage_enabled("analyze")
        assert not window.workflow_rail.is_stage_enabled("results")
        assert not window.workflow_rail.is_stage_enabled("measure")

    def test_the_menus_toolbars_and_canvas_are_disabled_while_running(
        self, window: MainWindow
    ) -> None:
        branded_container = window.menuWidget()
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))

        window._on_calculate_requested(1.0)

        assert window.menuWidget() is branded_container
        assert not window._menu_bar.isEnabled()
        assert not window.view_toolbar.isEnabled()
        assert not window.dual_image_view.isEnabled()

    def test_the_analyze_dock_stays_live_so_cancel_is_reachable(self, window: MainWindow) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))

        window._on_calculate_requested(1.0)

        assert window._analysis_setup_dock.isEnabled()
        assert not window._results_dock.isEnabled()

    def test_finishing_gives_the_whole_window_back(self, window: MainWindow) -> None:
        branded_container = window.menuWidget()
        window._set_calculation_lock(True)

        window._on_calculate_success(_make_results())

        assert window.menuWidget() is branded_container
        assert window._menu_bar.isEnabled()
        assert window.dual_image_view.isEnabled()
        assert all(
            window.workflow_rail.is_stage_enabled(key) for key in window.workflow_rail.stage_keys
        )

    def test_a_failed_run_gives_the_whole_window_back_too(self, window: MainWindow) -> None:
        window._set_calculation_lock(True)

        window._on_calculate_error("boom")

        assert window.dual_image_view.isEnabled()
        assert window.workflow_rail.is_stage_enabled("results")


class TestExportBusyDialog:
    """Export still runs synchronously on the GUI thread behind a modal
    busy dialog - only Calculate moved into the page."""

    def test_a_progress_update_switches_the_dialog_to_determinate_mode(
        self, window: MainWindow
    ) -> None:
        window._show_busy_dialog("Exporting...")

        window._on_export_progress(5, 20)

        assert window._busy_dialog is not None
        assert window._busy_dialog.maximum() == 20
        assert window._busy_dialog.value() == 5

    def test_the_label_text_includes_the_percentage(self, window: MainWindow) -> None:
        window._show_busy_dialog("Exporting...")

        window._on_export_progress(1, 4)  # 25%

        assert "25" in window._busy_dialog.labelText()
        assert "Exporting..." in window._busy_dialog.labelText()

    def test_reaching_the_total_shows_100_percent(self, window: MainWindow) -> None:
        window._show_busy_dialog("Exporting...")

        window._on_export_progress(20, 20)

        assert "100" in window._busy_dialog.labelText()

    def test_a_stray_update_after_the_dialog_closed_is_a_noop(self, window: MainWindow) -> None:
        window._show_busy_dialog("Exporting...")
        window._hide_busy_dialog()

        window._on_export_progress(1, 4)  # must not raise (no dialog to update)

        assert window._busy_dialog is None

    def test_a_zero_total_is_ignored_rather_than_dividing_by_zero(self, window: MainWindow) -> None:
        window._show_busy_dialog("Exporting...")

        window._on_export_progress(0, 0)  # must not raise

        assert window._busy_dialog is not None

    def test_a_step_update_appears_in_the_label_text(self, window: MainWindow) -> None:
        """Export reuses _on_calculate_step for its own step_callback."""
        window._show_busy_dialog("Exporting...")

        window._on_calculate_step("Writing charts...")

        assert "Writing charts..." in window._busy_dialog.labelText()

    def test_step_and_progress_updates_both_appear_together(self, window: MainWindow) -> None:
        window._show_busy_dialog("Exporting...")

        window._on_export_progress(1, 4)  # 25%
        window._on_calculate_step("Writing charts...")

        label = window._busy_dialog.labelText()
        assert "25" in label
        assert "Writing charts..." in label

    def test_a_stray_step_update_after_the_dialog_closed_is_a_noop(
        self, window: MainWindow
    ) -> None:
        window._show_busy_dialog("Exporting...")
        window._hide_busy_dialog()

        window._on_calculate_step("Writing charts...")  # must not raise

        assert window._busy_dialog is None


class TestCalculateAllPairsWiring:
    def test_two_images_uses_the_single_pair_calculate_path(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        calls: list[str] = []
        monkeypatch.setattr(
            window.controller, "calculate", lambda *a, **k: calls.append("calculate")
        )
        monkeypatch.setattr(
            window.controller,
            "calculate_all_pairs",
            lambda *a, **k: calls.append("calculate_all_pairs"),
        )

        window._on_calculate_requested(1.0)

        assert calls == ["calculate"]

    def test_three_images_uses_the_all_pairs_calculate_path(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        panel_c = window.dual_image_view.add_comparison_panel()
        panel_c.set_image(np.zeros((10, 10)))
        calls: list[str] = []
        monkeypatch.setattr(
            window.controller, "calculate", lambda *a, **k: calls.append("calculate")
        )
        monkeypatch.setattr(
            window.controller,
            "calculate_all_pairs",
            lambda *a, **k: calls.append("calculate_all_pairs"),
        )

        window._on_calculate_requested(1.0)

        assert calls == ["calculate_all_pairs"]

    def test_all_pairs_success_populates_results_and_charts_with_the_first_pair(
        self, window: MainWindow
    ) -> None:
        results = {"AvsB": _make_chartable_results(), "AvsC": _make_results()}

        window._on_calculate_all_pairs_success(results)

        assert window.results_panel.model.dataframe is results["AvsB"]
        assert not window.results_panel._pair_combo.isHidden()
        assert window.charts_panel._results_df is results["AvsB"]
        assert not window.charts_panel._pair_combo.isHidden()
        assert window.charts_panel.current_pair_label() == "AvsB"

    def test_pair_changed_updates_the_charts_panel(self, window: MainWindow) -> None:
        chartable = _make_chartable_results()
        results = {"AvsB": chartable, "AvsC": _make_results()}
        window._on_calculate_all_pairs_success(results)

        window.results_panel._pair_combo.setCurrentIndex(1)

        assert window.charts_panel._results_df is results["AvsC"]
        # Charts' own selector must stay truthful about what's now shown,
        # even though the pair was picked from Results' combo instead.
        assert window.charts_panel.current_pair_label() == "AvsC"

    def test_charts_own_pair_combo_can_select_a_pair_independently(
        self, window: MainWindow
    ) -> None:
        # Regression test: charts used to only ever show whichever pair
        # Calculate computed first, with no way to view any other pair's
        # chart at all from the Charts dock itself.
        results = {
            "AvsB": _make_chartable_results(),
            "AvsC": _make_results(),
            "BvsC": _make_chartable_results(),
        }
        window._on_calculate_all_pairs_success(results)

        window.charts_panel._pair_combo.setCurrentText("BvsC")

        assert window.charts_panel._results_df is results["BvsC"]
        # Picking a pair from Charts' own combo does not disturb Results.
        assert window.results_panel.model.dataframe is results["AvsB"]


class TestChartsPairsWithAllHelper:
    """_charts_pairs_with_all appends the special "R vs All" pair entry
    (mirrors HeatmapPanel's own _refresh_heatmap_pairs convention) once
    there's more than one reference-based pair to combine."""

    def test_appends_r_vs_all_when_multiple_reference_pairs_exist(self) -> None:
        pairs = main_window_module._charts_pairs_with_all(
            {"R vs C1": pd.DataFrame(), "R vs C2": pd.DataFrame(), "C1 vs C2": pd.DataFrame()}
        )

        assert pairs == ["R vs C1", "R vs C2", "C1 vs C2", "R vs All"]

    def test_no_all_entry_with_only_one_reference_pair(self) -> None:
        assert main_window_module._charts_pairs_with_all({"R vs C1": pd.DataFrame()}) == ["R vs C1"]

    def test_no_all_entry_when_empty(self) -> None:
        assert main_window_module._charts_pairs_with_all({}) == []


class TestChartsAllComparisonsWiring:
    def test_three_images_appends_r_vs_all_to_the_charts_pair_combo(
        self, window: MainWindow
    ) -> None:
        results = {
            "R vs C1": _make_chartable_results(),
            "R vs C2": _make_chartable_results(),
            "C1 vs C2": _make_chartable_results(),
        }

        window._on_calculate_all_pairs_success(results)

        pair_items = [
            window.charts_panel._pair_combo.itemText(i)
            for i in range(window.charts_panel._pair_combo.count())
        ]
        assert "R vs All" in pair_items

    def test_charts_gets_every_pairs_own_dataframe(self, window: MainWindow) -> None:
        results = {
            "R vs C1": _make_chartable_results(),
            "R vs C2": _make_chartable_results(),
        }

        window._on_calculate_all_pairs_success(results)

        assert window.charts_panel._results_by_pair == results


class TestPerSliceDistributionsWiring:
    """Both Calculate success paths kick off AppController.
    compute_per_slice_distributions afterward (see MainWindow.
    _refresh_per_slice_distributions) - feeding the Charts panel's
    boxplot/violin "Distributions" view without the user having to do
    anything beyond the Calculate they already pressed."""

    def test_calculate_success_passes_through_the_same_data_range_and_resolution(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        monkeypatch.setattr(window.controller, "calculate", lambda *a, **k: None)
        calls: list[tuple[float, str | None]] = []
        monkeypatch.setattr(
            window.controller,
            "compute_per_slice_distributions",
            lambda data_range, on_success, on_error, resolution=None, background=None: calls.append(
                (data_range, resolution)
            ),
        )

        window._start_calculate(2.5, "truncate")
        window._on_calculate_success(_make_results())

        assert calls == [(2.5, "truncate")]

    def test_all_pairs_success_also_triggers_distributions(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(window.controller, "calculate_all_pairs", lambda *a, **k: None)
        calls: list[float] = []
        monkeypatch.setattr(
            window.controller,
            "compute_per_slice_distributions",
            lambda data_range, on_success, on_error, resolution=None, background=None: calls.append(
                data_range
            ),
        )

        window._start_calculate(3.0, None)
        window._on_calculate_all_pairs_success({"AvsB": _make_results()})

        assert calls == [3.0]

    def test_success_routes_into_the_charts_panel(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        distributions = {"mse": {"": [1.0, 2.0, 3.0]}}
        monkeypatch.setattr(
            window.controller,
            "compute_per_slice_distributions",
            lambda data_range, on_success, on_error, resolution=None, background=None: on_success(
                distributions
            ),
        )

        window._on_calculate_success(_make_results())

        assert window.charts_panel._per_slice_distributions == distributions

    def test_failure_is_reported_via_the_errors_panel(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(
            window.controller,
            "compute_per_slice_distributions",
            lambda data_range, on_success, on_error, resolution=None, background=None: on_error(
                "boom"
            ),
        )

        window._on_calculate_success(_make_results())

        assert "boom" in window.errors_panel._table.item(0, 1).text()


class _FakeSliceMismatchDialog:
    """Stands in for the real SliceMismatchDialog: records the `counts`
    dict it was constructed with, and `exec()`/`resolution` are driven by
    the two class attributes below - set them before triggering the
    action under test."""

    accepted = True
    resolution: str | None = None
    received_counts: list[dict[str, int]] = []

    def __init__(self, counts: dict[str, int], parent: object = None) -> None:
        type(self).received_counts.append(counts)

    def exec(self) -> int:
        return (
            QDialog.DialogCode.Accepted.value
            if self.accepted
            else QDialog.DialogCode.Rejected.value
        )


class TestSliceMismatchWiring:
    def test_no_mismatch_calculates_immediately_without_any_dialog(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        fake = _FakeSliceMismatchDialog
        fake.received_counts = []
        monkeypatch.setattr(main_window_module, "SliceMismatchDialog", fake)
        calls: list[object] = []
        monkeypatch.setattr(
            window.controller, "calculate", lambda *a, **k: calls.append(k.get("resolution"))
        )

        window._on_calculate_requested(1.0)

        assert fake.received_counts == []
        assert calls == [None]

    def test_mismatch_prompts_with_each_panels_display_name(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((500, 4, 4)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((460, 4, 4)))
        fake = _FakeSliceMismatchDialog
        fake.received_counts = []
        fake.accepted = False
        fake.resolution = None
        monkeypatch.setattr(main_window_module, "SliceMismatchDialog", fake)

        window._on_calculate_requested(1.0)

        assert fake.received_counts == [
            {
                window.dual_image_view.reference_panel.display_name(): 500,
                window.dual_image_view.comparison_panel.display_name(): 460,
            }
        ]

    def test_cancelling_the_mismatch_dialog_skips_calculate(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((500, 4, 4)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((460, 4, 4)))
        fake = _FakeSliceMismatchDialog
        fake.received_counts = []
        fake.accepted = False
        fake.resolution = None
        monkeypatch.setattr(main_window_module, "SliceMismatchDialog", fake)
        calls: list[object] = []
        monkeypatch.setattr(window.controller, "calculate", lambda *a, **k: calls.append(True))

        window._on_calculate_requested(1.0)

        assert calls == []

    def test_chosen_resolution_is_passed_through_to_calculate(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((500, 4, 4)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((460, 4, 4)))
        fake = _FakeSliceMismatchDialog
        fake.received_counts = []
        fake.accepted = True
        fake.resolution = "truncate"
        monkeypatch.setattr(main_window_module, "SliceMismatchDialog", fake)
        calls: list[object] = []
        monkeypatch.setattr(
            window.controller, "calculate", lambda *a, **k: calls.append(k.get("resolution"))
        )

        window._on_calculate_requested(1.0)

        assert calls == ["truncate"]


class TestSessionPersistenceWiring:
    def test_save_analysis_without_results_shows_error_in_status_bar(
        self, window: MainWindow
    ) -> None:
        window._on_save_analysis()
        assert "Calculate metrics" in window.statusBar().currentMessage()

    def test_save_and_reopen_updates_history_panel_and_results_table(
        self, window: MainWindow, qtbot: QtBot, tmp_path: Path
    ) -> None:
        reference_path = tmp_path / "ref.png"
        comparison_path = tmp_path / "cmp.png"
        _save_png(reference_path, seed=0)
        _save_png(comparison_path, seed=1)
        with qtbot.waitSignal(window.dual_image_view.reference_loaded, timeout=5000):
            window.controller.open_reference(reference_path)
        with qtbot.waitSignal(window.dual_image_view.comparison_loaded, timeout=5000):
            window.controller.open_comparison(comparison_path)
        window.controller._last_results = _make_results()
        window.controller._last_data_range = 255.0

        window._on_save_analysis()
        assert window.history_panel._list.count() == 1

        window.history_panel._list.setCurrentRow(0)
        window.history_panel._reopen_button.click()

        assert list(window.results_panel.model.dataframe["metric_id"]) == ["mse"]


class TestConsistencyWarningReporting:
    def test_an_analysis_warning_reaches_the_errors_panel_with_a_prefix(
        self, window: MainWindow
    ) -> None:
        """A pre-flight finding (see AppController.analysis_warning) must
        outlive the status bar's timeout: it competes with the results
        arriving right behind it."""
        window.controller.analysis_warning.emit("Reference and C1 disagree about spacing")

        recorded = [
            window.errors_panel._table.item(row, 1).text()
            for row in range(window.errors_panel._table.rowCount())
        ]
        assert any("disagree about spacing" in message for message in recorded)
        assert any(message.startswith("Warning:") for message in recorded)


class TestExportWiring:
    """_on_save_export_requested is the single shared handler for both
    the Results panel's Export button and File > Save/Export... - these
    tests drive it through the "results only" choice (matching the old,
    single-purpose Export Results behavior); TestProjectFolderMenuWiring
    covers the "project"/"both" choices."""

    @pytest.fixture(autouse=True)
    def _choose_results_only(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(MainWindow, "_prompt_save_export_choice", lambda self: "results")

    @pytest.fixture(autouse=True)
    def _stub_success_dialog(self, monkeypatch: pytest.MonkeyPatch) -> None:
        # A successful save/export now always ends in a real
        # QMessageBox.information popup (see _on_save_export_requested) -
        # its .exec() would otherwise block forever in these headless
        # tests, waiting for a click that never comes.
        stub_message_box(monkeypatch, "qicom.gui.main_window", information=lambda *a, **k: None)

    def test_export_without_results_shows_error_and_never_opens_a_dialog(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        opened = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: opened.append(True) or "",
        )

        window._on_save_export_requested()

        assert opened == []
        assert "Calculate metrics" in window.statusBar().currentMessage()

    def test_export_cancelled_folder_dialog_is_a_noop(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # The choice dialog used to be left unanswered, so the flow stopped
        # before the folder dialog this test is about was ever shown.
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(window, "_prompt_save_export_choice", lambda: "both")
        asked: list[str] = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: asked.append("folder") or "",
        )

        def no_name(*_args: object, **_kwargs: object) -> tuple[str, bool]:
            raise AssertionError("a name was asked for after the folder was cancelled")

        monkeypatch.setattr("qicom.gui.main_window.QInputDialog.getText", no_name)
        saved: list[object] = []
        monkeypatch.setattr(window, "_save_project_zip", lambda *a, **k: saved.append(a))

        window._on_save_export_requested()

        assert asked == ["folder"]
        assert saved == []
        assert window._busy_dialog is None

    def test_cancelling_the_choice_dialog_is_a_noop(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(MainWindow, "_prompt_save_export_choice", lambda self: None)
        opened = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: opened.append(True) or "",
        )

        window._on_save_export_requested()

        assert opened == []

    def test_export_writes_files_and_reports_success(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText",
            lambda *a, **k: ("my_export", True),
        )

        window._on_save_export_requested()

        # Its own subfolder (named after the base name, suffixed
        # "_Results"), not loose directly into the folder the user picked.
        assert (tmp_path / "my_export_Results" / "my_export.xlsx").exists()
        assert window.statusBar().currentMessage() == "Ready"

    def test_export_name_cannot_escape_the_chosen_directory(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        # Whatever the user types is a filename stem, not a path: "../"
        # segments and separators used to build paths outside the folder
        # they picked (and, on Windows, names the filesystem rejects
        # outright, which surfaced as an export that silently wrote
        # nothing).
        destination = tmp_path / "chosen"
        destination.mkdir()
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(destination),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText",
            lambda *a, **k: ("../escaped/report:1", True),
        )

        window._on_save_export_requested()

        written = sorted(path.name for path in destination.iterdir())
        assert written == ["___escaped_report_1_Results"]
        assert not (tmp_path / "escaped").exists()
        assert list((destination / "___escaped_report_1_Results").glob("*.xlsx"))

    def test_export_shows_a_success_dialog_with_the_save_path(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText",
            lambda *a, **k: ("my_export", True),
        )
        info_calls: list[str] = []
        stub_message_box(
            monkeypatch,
            "qicom.gui.main_window",
            information=lambda *a, **k: info_calls.append(a[2]),
        )

        window._on_save_export_requested()

        assert len(info_calls) == 1
        assert str(tmp_path / "my_export_Results") in info_calls[0]

    def test_export_writes_native_charts_not_an_image(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from openpyxl import load_workbook

        results = _make_chartable_results()
        window.results_panel.set_dataframe(results)
        window.charts_panel.set_results(results)
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText",
            lambda *a, **k: ("my_export", True),
        )

        window._on_save_export_requested()

        workbook = load_workbook(tmp_path / "my_export_Results" / "my_export.xlsx")
        # One sheet per Charts panel view (see ChartsPanel.
        # export_data_for_every_view) - "Metrics Overview" is the only one
        # with data for this plain 2-image, no-ROI fixture.
        assert "Metrics Overview" in workbook.sheetnames
        # A real, editable openpyxl chart object - never a screenshot/
        # embedded picture (see export_native_charts_to_workbook).
        assert len(workbook["Metrics Overview"]._charts) >= 1
        assert len(workbook["Metrics Overview"]._images) == 0
        assert not (tmp_path / "my_export_Results" / "my_export_chart.jpg").exists()

    def test_export_shows_a_busy_dialog_and_hides_it_when_done(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText",
            lambda *a, **k: ("my_export", True),
        )
        seen_dialog_during_export: list[bool] = []
        real_export_results = window.controller.export_results

        def _wrapped_export_results(*args: object, **kwargs: object) -> object:
            seen_dialog_during_export.append(window._busy_dialog is not None)
            return real_export_results(*args, **kwargs)  # type: ignore[arg-type]

        monkeypatch.setattr(window.controller, "export_results", _wrapped_export_results)

        window._on_save_export_requested()

        assert seen_dialog_during_export == [True]
        assert window._busy_dialog is None

    def test_export_hides_the_busy_dialog_even_on_failure(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText",
            lambda *a, **k: ("my_export", True),
        )
        monkeypatch.setattr(
            window.controller,
            "export_results",
            lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")),
        )

        window._on_save_export_requested()

        assert window._busy_dialog is None

    def test_export_failure_is_reported_via_errors_panel(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText",
            lambda *a, **k: ("my_export", True),
        )
        monkeypatch.setattr(
            window.controller,
            "export_results",
            lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")),
        )

        window._on_save_export_requested()

        assert "disk full" in window.errors_panel._table.item(0, 1).text()


class TestHeatmapWiring:
    def _load_images(self, window: MainWindow) -> None:
        rng = np.random.default_rng(0)
        window.dual_image_view.reference_panel.set_image(rng.random((32, 32)))
        window.dual_image_view.comparison_panel.set_image(np.random.default_rng(1).random((32, 32)))

    def test_show_button_overlays_the_heatmap_and_updates_the_legend(
        self, window: MainWindow
    ) -> None:
        self._load_images(window)

        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )
        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert window.dual_image_view.reference_panel.heatmap_levels() is not None
        assert window.heatmap_panel._colorbar_widget.isEnabled()

    def test_hide_button_removes_the_heatmap_and_clears_the_legend(
        self, window: MainWindow
    ) -> None:
        self._load_images(window)
        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert window.dual_image_view.reference_panel.heatmap_levels() is None
        assert not window.heatmap_panel._colorbar_widget.isEnabled()

    def test_show_without_images_reports_an_error(self, window: MainWindow) -> None:
        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert window.errors_panel._table.rowCount() == 1

    def test_heatmap_failed_signal_clears_legend_and_reports_error(
        self, window: MainWindow
    ) -> None:
        self._load_images(window)
        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()
        assert window.heatmap_panel._colorbar_widget.isEnabled()

        # DualImageView always calls hide_heatmap() before emitting
        # heatmap_failed (see _safe_refresh_heatmap) - heatmap_hidden's own
        # handler is what actually clears the legend/active row.
        window.dual_image_view.hide_heatmap()
        window.dual_image_view.heatmap_failed.emit("boom")

        assert not window.heatmap_panel._colorbar_widget.isEnabled()
        assert "boom" in window.errors_panel._table.item(0, 1).text()

    def test_adding_a_panel_populates_the_heatmap_pair_combo(self, window: MainWindow) -> None:
        window.dual_image_view.add_comparison_panel()

        assert not window.heatmap_panel._pair_combo.isHidden()
        pairs = [
            window.heatmap_panel._pair_combo.itemText(i)
            for i in range(window.heatmap_panel._pair_combo.count())
        ]
        # "R vs All" only once there's more than one comparison to
        # meaningfully tell apart from a single "R vs C1"; "C1 vs C2" is
        # offered too, same as Results/Charts (every unordered pair).
        assert pairs == ["R vs C1", "R vs C2", "C1 vs C2", "R vs All"]

    def test_pair_combo_also_offers_comparison_vs_comparison(self, window: MainWindow) -> None:
        window.dual_image_view.add_comparison_panel()
        window.dual_image_view.add_comparison_panel()

        pairs = [
            window.heatmap_panel._pair_combo.itemText(i)
            for i in range(window.heatmap_panel._pair_combo.count())
        ]

        assert pairs == [
            "R vs C1",
            "R vs C2",
            "R vs C3",
            "C1 vs C2",
            "C1 vs C3",
            "C2 vs C3",
            "R vs All",
        ]

    def test_show_button_applies_to_a_comparison_vs_comparison_pair(
        self, window: MainWindow
    ) -> None:
        # comparison_panel is already "C1" by construction - add_comparison_
        # panel() adds C2 (see TestCalculateAllPairsWiring's own use of the
        # same default-panel layout elsewhere in this module).
        panel_c1 = window.dual_image_view.comparison_panel
        panel_c1.set_image(np.random.default_rng(1).random((32, 32)))
        panel_c2 = window.dual_image_view.add_comparison_panel()
        panel_c2.set_image(np.random.default_rng(2).random((32, 32)))
        window.heatmap_panel._pair_combo.setCurrentText("C1 vs C2")
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )

        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert panel_c1.heatmap_levels() is not None
        assert panel_c2.heatmap_levels() is not None
        assert window.dual_image_view.reference_panel.heatmap_levels() is None

    def test_show_button_applies_to_a_comparison_beyond_the_first(self, window: MainWindow) -> None:
        window.dual_image_view.reference_panel.set_image(np.random.default_rng(0).random((32, 32)))
        window.dual_image_view.add_comparison_panel()
        panel_d = window.dual_image_view.add_comparison_panel()
        panel_d.set_image(np.random.default_rng(1).random((32, 32)))
        window.heatmap_panel._pair_combo.setCurrentText("R vs C3")
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )

        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert panel_d.heatmap_levels() is not None

    def test_r_vs_all_shows_a_heatmap_on_every_comparison_panel(self, window: MainWindow) -> None:
        window.dual_image_view.reference_panel.set_image(np.random.default_rng(0).random((32, 32)))
        panel_c1 = window.dual_image_view.comparison_panel
        panel_c1.set_image(np.random.default_rng(1).random((32, 32)))
        panel_c2 = window.dual_image_view.add_comparison_panel()
        panel_c2.set_image(np.random.default_rng(2).random((32, 32)))
        window.heatmap_panel._pair_combo.setCurrentText("R vs All")
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )

        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert panel_c1.heatmap_levels() is not None
        assert panel_c2.heatmap_levels() is not None
        assert window.dual_image_view.reference_panel.heatmap_levels() is not None
        assert window.errors_panel._table.rowCount() == 0

    def test_removing_a_panel_shrinks_the_heatmap_pair_combo(self, window: MainWindow) -> None:
        panel_c = window.dual_image_view.add_comparison_panel()

        window.dual_image_view.remove_comparison_panel(panel_c.role)

        assert window.heatmap_panel._pair_combo.isHidden()

    def test_show_button_uses_the_heatmap_panels_selected_pair(self, window: MainWindow) -> None:
        window.dual_image_view.reference_panel.set_image(np.random.default_rng(0).random((32, 32)))
        panel_c = window.dual_image_view.add_comparison_panel()
        panel_c.set_image(np.random.default_rng(1).random((32, 32)))
        window.heatmap_panel._pair_combo.setCurrentText("R vs C2")
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )

        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert panel_c.heatmap_levels() is not None
        assert window.dual_image_view.comparison_panel.heatmap_levels() is None

    def test_show_button_passes_the_clicked_rois_id(self, window: MainWindow) -> None:
        self._load_images(window)
        pair = window.dual_image_view.add_roi("rectangle")
        window.heatmap_panel.set_available_rois([(pair.roi_id, pair.name)])
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )

        window.heatmap_panel._row_buttons[pair.roi_id].click()

        data = window.dual_image_view.reference_panel._heatmap_item.image
        mask = window.dual_image_view.get_roi_mask(pair.roi_id, "reference")
        assert np.isnan(data[~mask]).all()
        assert not np.isnan(data[mask]).any()

    def test_two_roi_rows_can_be_active_at_once(self, window: MainWindow) -> None:
        # Regression test: clicking a second ROI's Show button used to
        # replace the first one's heatmap entirely instead of the two
        # coexisting.
        self._load_images(window)
        pair_a = window.dual_image_view.add_roi("rectangle")
        pair_b = window.dual_image_view.add_roi("rectangle")
        pair_b.rois["reference"].setPos([0, 0])
        window.heatmap_panel.set_available_rois(
            [(pair_a.roi_id, pair_a.name), (pair_b.roi_id, pair_b.name)]
        )
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )

        window.heatmap_panel._row_buttons[pair_a.roi_id].click()
        window.heatmap_panel._row_buttons[pair_b.roi_id].click()

        hide_text = translator.tr("heatmap.hide")
        assert window.heatmap_panel._row_buttons[pair_a.roi_id].text() == hide_text
        assert window.heatmap_panel._row_buttons[pair_b.roi_id].text() == hide_text
        data = window.dual_image_view.reference_panel._heatmap_item.image
        union = window.dual_image_view.get_roi_mask(
            pair_a.roi_id, "reference"
        ) | window.dual_image_view.get_roi_mask(pair_b.roi_id, "reference")
        assert not np.isnan(data[union]).any()
        assert np.isnan(data[~union]).all()

    def test_clicking_whole_image_deactivates_active_roi_rows(self, window: MainWindow) -> None:
        self._load_images(window)
        pair = window.dual_image_view.add_roi("rectangle")
        window.heatmap_panel.set_available_rois([(pair.roi_id, pair.name)])
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )
        window.heatmap_panel._row_buttons[pair.roi_id].click()

        window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].click()

        assert window.heatmap_panel._row_buttons[pair.roi_id].text() == translator.tr(
            "heatmap.show"
        )
        assert window.heatmap_panel._row_buttons[WHOLE_IMAGE_TARGET].text() == translator.tr(
            "heatmap.hide"
        )
        data = window.dual_image_view.reference_panel._heatmap_item.image
        assert not np.isnan(data).any()  # whole image, nothing masked out

    def test_roi_heatmap_activates_for_a_comparison_panel_added_after_the_fact(
        self, window: MainWindow
    ) -> None:
        # Regression test: a comparison panel added after the reference/
        # first comparison were already loaded must still support a
        # ROI-restricted heatmap, the same as the original two panels.
        self._load_images(window)
        panel_c2 = window.dual_image_view.add_comparison_panel()
        panel_c2.set_image(np.random.default_rng(2).random((32, 32)))
        pair = window.dual_image_view.add_roi("rectangle")
        window.heatmap_panel._pair_combo.setCurrentText("R vs C2")
        window.heatmap_panel._metric_combo.setCurrentIndex(
            window.heatmap_panel._metric_combo.findData("mse")
        )

        window.heatmap_panel._row_buttons[pair.roi_id].click()

        assert panel_c2.heatmap_levels() is not None
        assert window.errors_panel._table.rowCount() == 0

    def test_adding_an_roi_populates_the_targets_list(self, window: MainWindow) -> None:
        window.dual_image_view.add_roi("rectangle")

        # index 0 is always the permanent "Whole Image" row.
        assert window.heatmap_panel._targets_list.count() == 2

    def test_removing_an_roi_shrinks_the_targets_list(self, window: MainWindow) -> None:
        pair = window.dual_image_view.add_roi("rectangle")

        window.dual_image_view.remove_roi(pair.roi_id)

        assert window.heatmap_panel._targets_list.count() == 1

    def test_renaming_an_roi_updates_the_targets_list_text(self, window: MainWindow) -> None:
        pair = window.dual_image_view.add_roi("rectangle")

        window.dual_image_view.rename_roi(pair.roi_id, "Signal")

        row_widget = window.heatmap_panel._targets_list.itemWidget(
            window.heatmap_panel._targets_list.item(1)
        )
        assert row_widget is not None
        label = row_widget.findChild(QLabel)
        assert label is not None
        assert label.text() == "Signal"

    def test_targets_list_narrows_to_the_current_slices_rois_on_scroll(
        self, window: MainWindow
    ) -> None:
        """Regression test for the reported bug: a ROI drawn without
        persist-across-slices only exists on the slice it was drawn on
        (see DualImageView._is_roi_visible_on_slice) - the heatmap's own
        per-ROI target list must track that, not keep listing every ROI
        ever drawn regardless of which slice is showing."""
        window.dual_image_view.reference_panel.set_image(
            np.random.default_rng(0).random((3, 32, 32))
        )
        window.dual_image_view.comparison_panel.set_image(
            np.random.default_rng(1).random((3, 32, 32))
        )
        window.dual_image_view.add_roi("rectangle")
        assert window.heatmap_panel._targets_list.count() == 2  # Whole Image + the new ROI

        window.dual_image_view.reference_panel.step_slice(1)
        assert window.heatmap_panel._targets_list.count() == 1  # Whole Image only

        window.dual_image_view.reference_panel.step_slice(-1)
        assert window.heatmap_panel._targets_list.count() == 2  # back on its creation slice


class TestExportHeatmapsWiring:
    def _load_images(self, window: MainWindow) -> None:
        rng = np.random.default_rng(0)
        window.dual_image_view.reference_panel.set_image(rng.random((32, 32)))
        window.dual_image_view.comparison_panel.set_image(np.random.default_rng(1).random((32, 32)))

    def test_no_rois_reports_an_error_and_never_opens_a_dialog(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        opened = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: opened.append(True) or "",
        )

        window._on_export_heatmaps_requested()

        assert opened == []
        assert window.errors_panel._table.rowCount() == 1

    def test_r_vs_all_reports_an_error_and_never_opens_a_dialog(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # export_heatmap_for_roi needs exactly one specific pair - "R vs
        # All" has no single pair to resolve to.
        self._load_images(window)
        window.dual_image_view.add_roi("rectangle")
        window.dual_image_view.add_comparison_panel()
        window.heatmap_panel._pair_combo.setCurrentText("R vs All")
        opened = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: opened.append(True) or "",
        )

        window._on_export_heatmaps_requested()

        assert opened == []
        assert window.errors_panel._table.rowCount() == 1

    def test_works_without_ever_having_clicked_show(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        # Regression test: exporting used to require an active on-screen
        # heatmap - it's now computed fresh per ROI, so it must succeed even
        # when Show was never clicked.
        self._load_images(window)
        window.dual_image_view.add_roi("rectangle")
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )

        window._on_export_heatmaps_requested()

        assert window.errors_panel._table.rowCount() == 0
        assert "1" in window.statusBar().currentMessage()

    def test_cancelled_folder_dialog_is_a_noop(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._load_images(window)
        window.dual_image_view.add_roi("rectangle")
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory", lambda *a, **k: ""
        )

        window._on_export_heatmaps_requested()  # must not raise

        assert window.errors_panel._table.rowCount() == 0

    def test_writes_one_file_per_roi_and_reports_success(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        self._load_images(window)
        pair_a = window.dual_image_view.add_roi("rectangle")
        pair_b = window.dual_image_view.add_roi("rectangle")
        window.dual_image_view.rename_roi(pair_a.roi_id, "Signal")
        window.dual_image_view.rename_roi(pair_b.roi_id, "Noise")
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )

        window._on_export_heatmaps_requested()

        # Its own "heatmaps" subfolder, not loose directly into the folder
        # the user picked.
        assert (tmp_path / "heatmaps" / "heatmap_Signal.png").exists()
        assert (tmp_path / "heatmaps" / "heatmap_Noise.png").exists()
        assert "2" in window.statusBar().currentMessage()


class TestResetAllWiring:
    def test_declining_confirmation_leaves_everything_untouched(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        stub_message_box(
            monkeypatch,
            "qicom.gui.main_window",
            question=lambda *a, **k: QMessageBox.StandardButton.No,
        )

        window.dual_image_view.add_roi("rectangle")
        window.results_panel.set_dataframe(_make_results())

        window._on_reset_all()

        assert window.dual_image_view.reference_panel.image is not None
        assert len(window.dual_image_view.list_roi_pairs()) == 1
        assert not window.results_panel.model.dataframe.empty

    def test_confirming_clears_images_rois_and_results(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.add_roi("rectangle")
        window.results_panel.set_dataframe(_make_results())
        window.charts_panel.set_results(_make_results())
        window.charts_panel.set_per_slice_distributions({"mse": {"": [1.0, 2.0, 3.0]}})
        window.errors_panel.add_error("some earlier error")
        stub_message_box(
            monkeypatch,
            "qicom.gui.main_window",
            question=lambda *a, **k: QMessageBox.StandardButton.Yes,
        )

        window._on_reset_all()

        assert window.dual_image_view.reference_panel.image is None
        assert window.dual_image_view.list_roi_pairs() == []
        assert window.results_panel.model.dataframe.empty
        assert window.charts_panel._per_slice_distributions == {}
        assert window.errors_panel._table.rowCount() == 0
        assert window.statusBar().currentMessage() == "Ready"

    def test_confirming_disables_export_and_clears_any_per_pair_results(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.results_panel.set_results_by_pair({"AvsB": _make_results(), "AvsC": _make_results()})
        stub_message_box(
            monkeypatch,
            "qicom.gui.main_window",
            question=lambda *a, **k: QMessageBox.StandardButton.Yes,
        )

        window._on_reset_all()

        assert window.results_panel.dataframe_for_pair("AvsB") is None
        assert window.results_panel._pair_combo.isHidden()
        assert not window.results_panel._export_button.isEnabled()


def _click_message_box_role(monkeypatch: pytest.MonkeyPatch, role: QMessageBox.ButtonRole) -> None:
    """QMessageBox.buttons() reorders per the platform's own button-order
    convention (observed here as [Save, Cancel, Discard], not insertion
    order) - matching by ButtonRole is the only reliable way to simulate
    "the user clicked Save/Discard/Cancel" regardless of that ordering."""
    monkeypatch.setattr(QMessageBox, "exec", lambda self: 0)
    monkeypatch.setattr(
        QMessageBox,
        "clickedButton",
        lambda self: next(b for b in self.buttons() if self.buttonRole(b) == role),
    )


class TestCloseEventWiring:
    """closeEvent's prompt box is a real QMessageBox(self) (not a static
    QMessageBox.question(...) call), so it's driven here by monkeypatching
    exec()/clickedButton() directly rather than answering a static dialog
    with stub_message_box."""

    def test_no_unsaved_results_closes_without_any_prompt(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        exec_calls: list[bool] = []
        monkeypatch.setattr(QMessageBox, "exec", lambda self: exec_calls.append(True) or 0)

        event = QCloseEvent()
        window.closeEvent(event)

        assert exec_calls == []
        assert event.isAccepted()

    def test_cancel_ignores_the_close(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.controller._has_unsaved_results = True
        _click_message_box_role(monkeypatch, QMessageBox.ButtonRole.RejectRole)

        event = QCloseEvent()
        window.closeEvent(event)

        assert not event.isAccepted()

    def test_discard_accepts_the_close_without_saving(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.controller._has_unsaved_results = True
        _click_message_box_role(monkeypatch, QMessageBox.ButtonRole.DestructiveRole)
        save_calls: list[bool] = []
        monkeypatch.setattr(
            window.controller, "save_project_zip", lambda *a, **k: save_calls.append(True)
        )

        event = QCloseEvent()
        window.closeEvent(event)

        assert save_calls == []
        assert event.isAccepted()

    def test_save_saves_then_accepts_the_close(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.controller._has_unsaved_results = True
        window.controller._current_project_zip = Path("/tmp/existing-project.zip")
        _click_message_box_role(monkeypatch, QMessageBox.ButtonRole.AcceptRole)
        save_calls: list[Path] = []
        monkeypatch.setattr(
            window.controller, "save_project_zip", lambda path, **_kwargs: save_calls.append(path)
        )
        monkeypatch.setattr(main_window_module, "record_recent_project", lambda *a, **k: None)

        event = QCloseEvent()
        window.closeEvent(event)

        assert save_calls == [Path("/tmp/existing-project.zip")]
        assert event.isAccepted()

    def test_save_failure_reports_error_and_ignores_the_close(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.controller._has_unsaved_results = True
        window.controller._current_project_zip = Path("/tmp/existing-project.zip")
        _click_message_box_role(monkeypatch, QMessageBox.ButtonRole.AcceptRole)
        monkeypatch.setattr(
            window.controller,
            "save_project_zip",
            lambda *a, **k: (_ for _ in ()).throw(ValueError("nothing to save")),
        )

        event = QCloseEvent()
        window.closeEvent(event)

        assert not event.isAccepted()
        assert "nothing to save" in window.errors_panel._table.item(0, 1).text()

    def test_save_without_an_existing_project_uses_a_portable_project_picker(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        window.controller._has_unsaved_results = True
        _click_message_box_role(monkeypatch, QMessageBox.ButtonRole.AcceptRole)
        selected = tmp_path / "new-project"
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getSaveFileName",
            lambda *a, **k: (str(selected), "QICOM Project (*.zip)"),
        )
        save_calls: list[Path] = []
        monkeypatch.setattr(
            window.controller, "save_project_zip", lambda path, **_kwargs: save_calls.append(path)
        )
        monkeypatch.setattr(main_window_module, "record_recent_project", lambda *a, **k: None)

        event = QCloseEvent()
        window.closeEvent(event)

        assert save_calls == [selected.with_suffix(".zip")]
        assert event.isAccepted()

    def test_ignores_the_close_while_a_calculation_is_still_running(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # Regression: a calculation still in flight has produced no
        # results yet, so has_unsaved_results is False and the ordinary
        # unsaved-results prompt would never even fire - closing here must
        # not let Qt tear down the still-running background QThread (see
        # workers.ThreadJob's own docstring on why that's a fatal abort).
        window.analysis_setup_panel.set_calculating(True)
        exec_calls: list[bool] = []
        monkeypatch.setattr(QMessageBox, "exec", lambda self: exec_calls.append(True) or 0)

        event = QCloseEvent()
        window.closeEvent(event)

        assert exec_calls == []
        assert not event.isAccepted()

    def test_calculation_finishing_lets_the_window_close_again(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.analysis_setup_panel.set_calculating(True)
        window.analysis_setup_panel.set_calculating(False)
        monkeypatch.setattr(QMessageBox, "exec", lambda self: 0)

        event = QCloseEvent()
        window.closeEvent(event)

        assert event.isAccepted()


class TestCloseProjectWiring:
    """File > Close Project routes through the exact same closeEvent/
    unsaved-results confirmation as Exit - see MainWindow._on_close_project.
    The only new behavior is return_to_launcher, which qicom.gui.app.main()
    reads after this window (the last one open) closes to decide whether to
    show the launcher again."""

    def test_closes_and_sets_return_to_launcher_when_nothing_unsaved(
        self, window: MainWindow
    ) -> None:
        window.show()

        window._on_close_project()

        assert window.return_to_launcher is True
        assert not window.isVisible()

    def test_cancelling_the_unsaved_results_prompt_keeps_the_window_open(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.show()
        window.controller._has_unsaved_results = True
        _click_message_box_role(monkeypatch, QMessageBox.ButtonRole.RejectRole)

        window._on_close_project()

        assert window.return_to_launcher is False
        assert window.isVisible()


class TestProjectFolderMenuWiring:
    @pytest.fixture(autouse=True)
    def _choose_project_only(self, monkeypatch: pytest.MonkeyPatch) -> None:
        # Overridden per-test (e.g. to "both") where that test cares.
        monkeypatch.setattr(MainWindow, "_prompt_save_export_choice", lambda self: "project")

    @pytest.fixture(autouse=True)
    def _stub_success_dialog(self, monkeypatch: pytest.MonkeyPatch) -> None:
        # A successful save/export now always ends in a real
        # QMessageBox.information popup (see _on_save_export_requested) -
        # its .exec() would otherwise block forever in these headless
        # tests, waiting for a click that never comes. Stubbed by default
        # for every test in this class; test_save_project_zip_shows_a_
        # success_dialog overrides this to actually inspect the call.
        stub_message_box(monkeypatch, "qicom.gui.main_window", information=lambda *a, **k: None)

    def _stub_name_dialog(self, monkeypatch: pytest.MonkeyPatch, name: str = "my_project") -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText", lambda *a, **k: (name, True)
        )

    def test_save_project_zip_prompts_for_a_folder_and_a_name_then_saves(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        self._stub_name_dialog(monkeypatch)
        calls: list[Path] = []

        def _fake_save(zip_path: Path, **_k: object) -> None:
            calls.append(zip_path)
            zip_path.touch()  # load_recent_projects() drops entries whose file is gone

        monkeypatch.setattr(window.controller, "save_project_zip", _fake_save)

        window._on_save_export_requested()

        assert calls == [tmp_path / "my_project_Project.zip"]
        assert window.errors_panel._table.rowCount() == 0
        recent = load_recent_projects()
        assert len(recent) == 1
        assert recent[0].path == tmp_path / "my_project_Project.zip"
        assert recent[0].mode == "quantitative"

    def test_busy_dialog_is_already_visible_when_project_compression_starts(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        self._stub_name_dialog(monkeypatch)
        seen: list[bool] = []

        def _fake_save(_path: Path, **_kwargs: object) -> None:
            seen.append(window._busy_dialog is not None and window._busy_dialog.isVisible())

        monkeypatch.setattr(window.controller, "save_project_zip", _fake_save)

        window._on_save_export_requested()

        assert seen == [True]
        assert window._busy_dialog is None

    def test_save_project_zip_failure_does_not_record_it_as_recent(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        self._stub_name_dialog(monkeypatch)
        monkeypatch.setattr(
            window.controller,
            "save_project_zip",
            lambda zip_path, **k: (_ for _ in ()).throw(ValueError("disk full")),
        )

        window._on_save_export_requested()

        assert load_recent_projects() == []

    def test_save_project_zip_cancelled_folder_dialog_is_a_noop(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory", lambda *a, **k: ""
        )
        calls: list[Path] = []
        monkeypatch.setattr(
            window.controller, "save_project_zip", lambda zip_path, **k: calls.append(zip_path)
        )

        window._on_save_export_requested()  # must not raise

        assert calls == []

    def test_save_project_zip_cancelled_name_dialog_is_a_noop(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        monkeypatch.setattr(
            "qicom.gui.main_window.QInputDialog.getText", lambda *a, **k: ("", False)
        )
        calls: list[Path] = []
        monkeypatch.setattr(
            window.controller, "save_project_zip", lambda zip_path, **k: calls.append(zip_path)
        )

        window._on_save_export_requested()  # must not raise

        assert calls == []

    def test_cancelling_the_choice_dialog_never_opens_the_folder_picker(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(MainWindow, "_prompt_save_export_choice", lambda self: None)
        opened: list[bool] = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: opened.append(True) or "",
        )

        window._on_save_export_requested()  # must not raise

        assert opened == []

    def test_save_project_zip_failure_reports_error(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        self._stub_name_dialog(monkeypatch)
        monkeypatch.setattr(
            window.controller,
            "save_project_zip",
            lambda *a, **k: (_ for _ in ()).throw(ValueError("no source file")),
        )

        window._on_save_export_requested()

        assert "no source file" in window.errors_panel._table.item(0, 1).text()

    def test_save_project_zip_shows_a_success_dialog_with_the_path(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        self._stub_name_dialog(monkeypatch)
        monkeypatch.setattr(window.controller, "save_project_zip", lambda zip_path, **k: None)
        info_calls: list[str] = []
        stub_message_box(
            monkeypatch,
            "qicom.gui.main_window",
            information=lambda *a, **k: info_calls.append(a[2]),
        )

        window._on_save_export_requested()

        assert len(info_calls) == 1
        assert str(tmp_path / "my_project_Project.zip") in info_calls[0]

    def test_project_only_choice_never_calls_export_results(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        self._stub_name_dialog(monkeypatch)
        monkeypatch.setattr(window.controller, "save_project_zip", lambda zip_path, **k: None)
        export_calls: list[object] = []
        monkeypatch.setattr(
            window.controller, "export_results", lambda *a, **k: export_calls.append(True)
        )

        window._on_save_export_requested()

        assert export_calls == []

    def test_both_choice_saves_the_project_zip_and_exports_results_as_sibling_outputs(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(MainWindow, "_prompt_save_export_choice", lambda self: "both")
        window.results_panel.set_dataframe(_make_results())
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: str(tmp_path),
        )
        self._stub_name_dialog(monkeypatch)
        save_calls: list[Path] = []
        monkeypatch.setattr(
            window.controller,
            "save_project_zip",
            lambda zip_path, **k: save_calls.append(zip_path),
        )
        export_calls: list[tuple[Path, str]] = []
        monkeypatch.setattr(
            window.controller,
            "export_results",
            lambda directory, base_name, results, **k: (
                export_calls.append((directory, base_name)) or [directory / f"{base_name}.xlsx"]
            ),
        )

        window._on_save_export_requested()

        # Sibling outputs, same name, same parent directory - not one
        # bundled inside the other (see _on_save_export_requested's own
        # docstring).
        assert save_calls == [tmp_path / "my_project_Project.zip"]
        assert export_calls == [(tmp_path / "my_project_Results", "my_project")]

    def test_both_choice_without_any_results_reports_an_error_and_saves_nothing(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(MainWindow, "_prompt_save_export_choice", lambda self: "both")
        opened: list[bool] = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getExistingDirectory",
            lambda *a, **k: opened.append(True) or str(tmp_path),
        )

        window._on_save_export_requested()

        assert opened == []  # blocked before even prompting for a folder
        assert "Calculate metrics" in window.statusBar().currentMessage()

    def test_open_project_zip_prompts_and_loads(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        zip_path = tmp_path / "my_project.zip"
        zip_path.touch()  # load_recent_projects() drops entries whose file is gone
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(zip_path), "QICOM Project (*.zip)"),
        )
        fake_data = ProjectFolderData(
            folder=tmp_path,
            version=1,
            created_at="now",
            saved_at="now",
            notes=None,
            data_range=123.0,
            panels=[],
            rois=[],
            annotations=[],
            lines=[],
            results={},
            warnings=[],
        )
        calls: list[Path] = []
        monkeypatch.setattr(
            window.controller,
            "load_project_zip",
            lambda zp, **k: calls.append(zp) or fake_data,
        )

        window._on_open_project_folder()

        assert calls == [zip_path]
        assert window.results_panel.data_range() == 123.0
        assert window.errors_panel._table.rowCount() == 0
        recent = load_recent_projects()
        assert len(recent) == 1
        assert recent[0].path == zip_path
        assert recent[0].mode == "quantitative"
        assert recent[0].name == "my_project"

    def test_open_project_zip_failure_does_not_record_it_as_recent(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(tmp_path / "my_project.zip"), "QICOM Project (*.zip)"),
        )
        monkeypatch.setattr(
            window.controller,
            "load_project_zip",
            lambda zp, **k: (_ for _ in ()).throw(ValueError("bad manifest")),
        )

        window._on_open_project_folder()

        assert load_recent_projects() == []

    def test_open_project_zip_surfaces_warnings(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(tmp_path / "my_project.zip"), "QICOM Project (*.zip)"),
        )
        fake_data = ProjectFolderData(
            folder=tmp_path,
            version=1,
            created_at="now",
            saved_at="now",
            notes=None,
            data_range=None,
            panels=[],
            rois=[],
            annotations=[],
            lines=[],
            results={},
            warnings=["Missing file for panel 'reference': /gone.png"],
        )
        monkeypatch.setattr(window.controller, "load_project_zip", lambda zp, **k: fake_data)

        window._on_open_project_folder()

        assert "Missing file" in window.errors_panel._table.item(0, 1).text()

    def test_open_project_zip_respects_unsaved_results_confirmation(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.controller._has_unsaved_results = True
        _click_message_box_role(monkeypatch, QMessageBox.ButtonRole.RejectRole)
        opened: list[bool] = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getOpenFileName",
            lambda *a, **k: (opened.append(True) or "", ""),
        )

        window._on_open_project_folder()

        assert opened == []  # Cancel must stop before the file picker ever opens

    def test_open_project_zip_failure_reports_error(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getOpenFileName",
            lambda *a, **k: (str(tmp_path / "my_project.zip"), "QICOM Project (*.zip)"),
        )
        monkeypatch.setattr(
            window.controller,
            "load_project_zip",
            lambda zp, **k: (_ for _ in ()).throw(ValueError("bad manifest")),
        )

        window._on_open_project_folder()

        assert "bad manifest" in window.errors_panel._table.item(0, 1).text()


class TestProjectZipRealRoundTrip:
    """Exercises AppController.load_project_zip's real async load path
    (not mocked) through MainWindow.open_project_zip directly - covers
    the Charts panel's ROI Stats view getting refreshed once every panel
    has actually finished (re)loading (see _refresh_roi_stats_for_charts's
    own docstring for the reported "charts don't come back after
    reopening a project" gap this closes)."""

    def test_roi_stats_chart_is_populated_after_reload(
        self, window: MainWindow, qtbot: QtBot, tmp_path: Path
    ) -> None:
        reference_path = tmp_path / "ref.png"
        comparison_path = tmp_path / "cmp.png"
        _save_png(reference_path, seed=0)
        _save_png(comparison_path, seed=1)
        with qtbot.waitSignal(window.dual_image_view.reference_loaded, timeout=5000):
            window.controller.open_reference(reference_path)
        with qtbot.waitSignal(window.dual_image_view.comparison_loaded, timeout=5000):
            window.controller.open_comparison(comparison_path)
        pair = window.dual_image_view.add_roi("rectangle")
        window.dual_image_view.set_roi_role(pair.roi_id, "signal_roi")
        zip_path = tmp_path / "project.zip"
        window.controller.save_project_zip(zip_path)

        # A fresh window, as if reopening the project from scratch.
        fresh_window = MainWindow(storage=SessionStorage(":memory:"))
        qtbot.addWidget(fresh_window)
        assert fresh_window.charts_panel._roi_stats_df.empty

        fresh_window.open_project_zip(zip_path)
        qtbot.waitUntil(
            lambda: all(p.image is not None for p in fresh_window.dual_image_view.list_panels()),
            timeout=5000,
        )
        qtbot.waitUntil(lambda: not fresh_window.charts_panel._roi_stats_df.empty, timeout=5000)

        assert not fresh_window.charts_panel._roi_stats_df.empty
        fresh_window.controller.cleanup()


class _FakeOpenImagesDialog:
    """Stands in for the real OpenImagesDialog: records which `panels` list
    it was constructed with, and `exec()`/`result_info` are driven entirely
    by the two class attributes below - set them before triggering the
    action under test."""

    accepted = True
    result_info: object = None
    received_panels: list[list[ImagePanel]] = []

    def __init__(self, panels: list[ImagePanel], parent: object = None) -> None:
        type(self).received_panels.append(panels)

    def exec(self) -> int:
        return (
            QDialog.DialogCode.Accepted.value
            if self.accepted
            else QDialog.DialogCode.Rejected.value
        )


class TestOpenImagesWiring:
    def test_menu_action_prompts_with_every_panel_as_a_target(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        fake = _FakeOpenImagesDialog
        fake.received_panels = []
        fake.accepted = False
        fake.result_info = None
        monkeypatch.setattr(main_window_module, "OpenImagesDialog", fake)

        window._on_open_images()

        assert fake.received_panels == [window.dual_image_view.list_panels()]

    def test_panel_add_images_button_prompts_with_only_that_panel_as_target(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        fake = _FakeOpenImagesDialog
        fake.received_panels = []
        fake.accepted = False
        fake.result_info = None
        monkeypatch.setattr(main_window_module, "OpenImagesDialog", fake)
        panel = window.dual_image_view.reference_panel

        window.dual_image_view.image_add_requested.emit(panel)

        assert fake.received_panels == [[panel]]

    def test_single_mode_result_loads_via_open_panel(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        path = str(tmp_path / "c.png")
        panel = window.dual_image_view.reference_panel
        fake = _FakeOpenImagesDialog
        fake.received_panels = []
        fake.accepted = True
        fake.result_info = (panel, path, "single")
        monkeypatch.setattr(main_window_module, "OpenImagesDialog", fake)
        received: list[object] = []
        monkeypatch.setattr(
            window.controller, "open_panel", lambda p, path: received.append((p, path))
        )

        window._on_open_images()

        assert received == [(panel, path)]

    def test_stack_mode_result_loads_via_open_panel_files(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        paths = [str(tmp_path / "a.png"), str(tmp_path / "b.png")]
        panel = window.dual_image_view.comparison_panel
        fake = _FakeOpenImagesDialog
        fake.received_panels = []
        fake.accepted = True
        fake.result_info = (panel, paths, "stack")
        monkeypatch.setattr(main_window_module, "OpenImagesDialog", fake)
        received: list[object] = []
        monkeypatch.setattr(
            window.controller, "open_panel_files", lambda p, paths: received.append((p, paths))
        )

        window._on_open_images()

        assert received == [(panel, paths)]

    def test_cancelled_dialog_is_a_noop(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        fake = _FakeOpenImagesDialog
        fake.received_panels = []
        fake.accepted = False
        fake.result_info = None
        monkeypatch.setattr(main_window_module, "OpenImagesDialog", fake)
        received: list[object] = []
        monkeypatch.setattr(
            window.controller, "open_panel", lambda p, path: received.append((p, path))
        )

        window._on_open_images()

        assert received == []


class TestPanelAddedWiring:
    def test_adding_a_comparison_panel_does_not_prompt_any_dialog(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # Loading is now opt-in via the menu's "Open Images..." action or a
        # panel's own "Load Images" header button - simply adding a new
        # comparison panel must not pop any dialog on its own.
        def _fail(*_args: object, **_kwargs: object) -> None:
            raise AssertionError("adding a panel must not open a dialog")

        monkeypatch.setattr(main_window_module, "OpenImagesDialog", _fail)

        panel = window.dual_image_view.add_comparison_panel()

        assert panel.image is None


class TestAnnotationManagerWiring:
    """The annotation list used to show the raw internal role
    ("reference"/"comparison"/"comparison_2") - now the same short display
    letter used everywhere else in the app (pair labels, panel headers)."""

    def test_reference_annotation_shows_r(self, window: MainWindow) -> None:
        window.dual_image_view.add_annotation("Note", (1.0, 2.0), "reference")

        item = window.annotation_manager_panel._list.item(0)
        assert item is not None
        assert item.text() == "[R] Note"

    def test_first_comparison_annotation_shows_c1_not_the_raw_role(
        self, window: MainWindow
    ) -> None:
        window.dual_image_view.add_annotation("Note", (1.0, 2.0), "comparison")

        item = window.annotation_manager_panel._list.item(0)
        assert item is not None
        assert item.text() == "[C1] Note"

    def test_third_panel_annotation_shows_c2(self, window: MainWindow) -> None:
        window.dual_image_view.add_comparison_panel()

        window.dual_image_view.add_annotation("Note", (1.0, 2.0), "comparison_2")

        item = window.annotation_manager_panel._list.item(0)
        assert item is not None
        assert item.text() == "[C2] Note"


class TestRoiPanelUnlockWiring:
    def test_roi_manager_panel_starts_with_the_original_two_panel_columns(
        self, window: MainWindow
    ) -> None:
        from qicom.gui.widgets.roi_manager import _COLUMNS

        assert window.roi_manager_panel._table.columnCount() == len(_COLUMNS) + 2

    def test_adding_a_panel_adds_its_own_unlock_column(self, window: MainWindow) -> None:
        from qicom.gui.widgets.roi_manager import _COLUMNS

        window.dual_image_view.add_comparison_panel()

        assert window.roi_manager_panel._table.columnCount() == len(_COLUMNS) + 3
        assert (
            window.roi_manager_panel._table.horizontalHeaderItem(len(_COLUMNS) + 2).text()
            == "Unlock C2"
        )

    def test_checking_the_unlock_checkbox_lets_that_panels_roi_move_independently(
        self, window: MainWindow
    ) -> None:
        from qicom.gui.widgets.roi_manager import _COLUMNS

        window.dual_image_view.reference_panel.set_image(np.random.default_rng(0).random((32, 32)))
        window.dual_image_view.comparison_panel.set_image(np.random.default_rng(1).random((32, 32)))
        pair = window.dual_image_view.add_roi("rectangle")

        item = window.roi_manager_panel._table.item(0, len(_COLUMNS) + 1)
        item.setCheckState(Qt.CheckState.Checked)
        pair.reference_roi.setPos([5, 5])

        assert pair.comparison_roi.pos() != pair.reference_roi.pos()


class TestRoiValuesPopulateOnCreation:
    """A new ROI's values used to stay blank in the table until the ROI was
    dragged: DualImageView.add_roi emitted roi_added (which creates the
    table row) only after computing/emitting the initial values, so
    RoiManagerPanel.update_row_value silently dropped them (no row to
    write into yet) - see add_roi's own comment."""

    def test_row_is_populated_immediately_without_moving_the_roi(self, window: MainWindow) -> None:
        window.dual_image_view.reference_panel.set_image(np.full((20, 20), 5.0))
        window.dual_image_view.comparison_panel.set_image(np.full((20, 20), 9.0))

        pair = window.dual_image_view.add_roi("rectangle")

        row = window.roi_manager_panel._rows[pair.roi_id]
        assert row.reference_values["roi_mean"] == pytest.approx(5.0)
        assert row.comparison_values["roi_mean"] == pytest.approx(9.0)


class TestReportBugWiring:
    def test_report_bug_action_opens_the_dialog(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        opened: list[object] = []
        records_seen: list[list[str]] = []

        class _FakeDialog:
            def __init__(self, parent: object = None, **kwargs: Any) -> None:
                opened.append(parent)
                records_seen.append([error.message for error in kwargs["error_records"]()])

            def exec(self) -> int:
                return 0

        monkeypatch.setattr(main_window_module, "BugReportDialog", _FakeDialog)
        window._report_error("boom")

        window._report_bug_action.trigger()

        # The dialog can attach the errors this window recorded to the report.
        assert opened == [window]
        assert records_seen == [["boom"]]

    def test_help_menu_has_a_translated_title_and_action_text(self, window: MainWindow) -> None:
        assert window._help_menu.title() == translator.tr("menu.help")
        assert window._report_bug_action.text() == translator.tr("menu.help.report_bug")


class TestConcatWiring:
    def test_shows_error_when_an_image_is_missing(self, window: MainWindow) -> None:
        window._on_concat_stacks()
        assert "Load both" in window.statusBar().currentMessage()

    def test_accepted_dialog_replaces_the_target_panel_image(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((2, 5, 5)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((3, 5, 5)))
        concatenated = np.ones((5, 5, 5))

        class _FakeDialog:
            def __init__(self, *args: object, **kwargs: object) -> None:
                pass

            def exec(self) -> int:
                return 1  # QDialog.DialogCode.Accepted

            result_stack = concatenated
            target_role = "comparison"

        monkeypatch.setattr(main_window_module, "ConcatDialog", _FakeDialog)

        window._on_concat_stacks()

        assert np.array_equal(window.dual_image_view.comparison_panel.image, concatenated)

    def test_dialog_is_offered_every_panel_with_a_loaded_image(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((2, 5, 5)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((3, 5, 5)))
        panel_c = window.dual_image_view.add_comparison_panel()
        panel_c.set_image(np.zeros((4, 5, 5)))
        received_panels: list[object] = []

        class _FakeDialog:
            def __init__(self, panels: object, parent: object = None) -> None:
                received_panels.append(panels)

            def exec(self) -> int:
                return 0  # QDialog.DialogCode.Rejected

            result_stack = None
            target_role = None

        monkeypatch.setattr(main_window_module, "ConcatDialog", _FakeDialog)

        window._on_concat_stacks()

        assert received_panels == [window.dual_image_view.list_panels()]


class TestErrorsPanelWiring:
    def test_load_failure_is_recorded_in_the_errors_panel(self, window: MainWindow) -> None:
        window._on_load_failed("boom")

        assert window.errors_panel._table.rowCount() == 1
        assert window.errors_panel._table.item(0, 1).text() == explained_error(
            "errors.summary.load_image", "boom"
        )
        assert window.statusBar().currentMessage() == explained_error(
            "errors.summary.load_image", "boom"
        )

    def test_concat_missing_images_error_is_recorded(self, window: MainWindow) -> None:
        window._on_concat_stacks()

        assert window.errors_panel._table.rowCount() == 1

    def test_save_without_results_error_is_recorded(self, window: MainWindow) -> None:
        window._on_save_analysis()

        assert window.errors_panel._table.rowCount() == 1

    def test_calculate_error_is_recorded(self, window: MainWindow) -> None:
        window._on_calculate_error("ValueError: something failed")

        assert window.errors_panel._table.rowCount() == 1
        assert "something failed" in window.errors_panel._table.item(0, 1).text()


class TestErrorTabBadge:
    def test_reporting_an_error_shows_a_count_of_one(self, window: MainWindow) -> None:
        window._report_error("boom")

        assert window._error_badge.text() == "1"
        assert window._error_badge.width() > 0

    def test_reporting_an_error_shows_the_status_bar_indicator(self, window: MainWindow) -> None:
        # The always-visible half of the count: no workflow step opens the
        # Errors dock any more (the Activity step is gone), so the tab
        # badge alone would only exist while the user already had the
        # panel open.
        assert window._errors_status_button.isHidden()

        window._report_error("boom")

        assert not window._errors_status_button.isHidden()
        assert "1" in window._errors_status_button.text()

    def test_status_bar_indicator_opens_the_errors_window(self, window: MainWindow) -> None:
        window._report_error("boom")

        window._errors_status_button.click()

        assert window._errors_dialog is not None
        assert window._errors_dialog.isVisible()
        assert any(
            "boom" in record.message for record in window._errors_dialog.records_for_action()
        )

    def test_closing_the_errors_window_leaves_the_indicator_in_the_status_bar(
        self, window: MainWindow
    ) -> None:
        # The whole point of counting *recorded* errors rather than unread
        # ones: reading them must not make them disappear from the
        # interface, only clearing the log does.
        window._report_error("boom")
        window._errors_status_button.click()
        assert window._errors_dialog is not None

        window._errors_dialog.close()

        assert not window._errors_status_button.isHidden()
        assert "1" in window._errors_status_button.text()

    def test_errors_reported_while_the_window_is_open_are_appended_to_it(
        self, window: MainWindow
    ) -> None:
        window._report_error("first")
        window._errors_status_button.click()

        window._report_error("second")

        assert window._errors_dialog is not None
        messages = [record.message for record in window._errors_dialog.records_for_action()]
        assert messages == ["first", "second"]

    def test_clearing_from_the_errors_window_empties_log_window_and_indicator(
        self, window: MainWindow
    ) -> None:
        window._report_error("boom")
        window._errors_status_button.click()
        assert window._errors_dialog is not None

        window._errors_dialog._clear_button.click()

        assert window.errors_panel._table.rowCount() == 0
        assert window._errors_dialog.records_for_action() == []
        assert window._errors_status_button.isHidden()

    def test_badge_is_attached_as_soon_as_the_errors_dock_is_shown(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        # Attached before the first error, not lazily on it - the whole
        # point is that the tab's size never jumps around, including at
        # the moment of the very first error. Qt only rebuilds the tab bar
        # of a laid-out window, and only creates one for a tabified group
        # with more than one visible dock - hence the show() and the
        # surrounding step (any step showing bottom panels will do).
        window.show()
        qtbot.waitExposed(window)
        window.workflow_rail.set_current_stage("measure")
        window._show_errors_panel()
        window._refresh_visible_dock_chrome()
        index = window._errors_tab_index()
        tab_bar = window._find_tab_bar()
        assert index is not None and tab_bar is not None

        assert tab_bar.tabButton(index, tab_bar.ButtonPosition.RightSide) is window._error_badge

    def test_no_unread_errors_collapses_the_badge_to_zero_width(self, window: MainWindow) -> None:
        # setVisible(False) doesn't stick once embedded via setTabButton
        # (verified empirically) - "nothing to show" is zero width instead.
        assert window._error_badge.width() == 0
        assert window._error_badge.text() == ""

    def test_badge_never_overflows_past_its_tabs_right_edge(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        # Regression test: widening the badge (e.g. "1" -> "99+") used to
        # leave the tab's own reserved width stuck at whatever it was when
        # the badge was first attached (zero-width, at construction time) -
        # QTabBar only consults tabSizeHint() when a tab button is first
        # set, not on every later resize of that same button - so the
        # badge visibly spilled past the tab's right edge. See
        # MainWindow._resize_error_badge_tab.
        window.show()
        qtbot.waitExposed(window)
        for i in range(150):
            window._report_error(f"error {i}")
        window.workflow_rail.set_current_stage("measure")
        window._show_errors_panel()
        window._refresh_visible_dock_chrome()

        tab_bar = window._find_tab_bar()
        index = window._errors_tab_index()
        assert tab_bar is not None and index is not None
        tab_right_edge = tab_bar.tabRect(index).right()
        badge_right_edge = window._error_badge.geometry().right()
        assert badge_right_edge <= tab_right_edge

    def test_repeated_errors_increment_the_count(self, window: MainWindow) -> None:
        window._report_error("first")
        window._report_error("second")
        window._report_error("third")

        assert window._error_badge.text() == "3"

    def test_errors_dock_becoming_visible_clears_the_badge(self, window: MainWindow) -> None:
        window._report_error("boom")

        window._on_errors_dock_visibility_changed(True)

        assert window._error_badge.text() == ""
        assert window._error_badge.width() == 0

    def test_errors_dock_becoming_hidden_does_not_clear_the_badge(self, window: MainWindow) -> None:
        window._report_error("boom")

        window._on_errors_dock_visibility_changed(False)

        assert window._error_badge.text() == "1"

    def test_reset_all_clears_the_badge(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window._report_error("boom")
        stub_message_box(
            monkeypatch,
            "qicom.gui.main_window",
            question=lambda *a, **k: QMessageBox.StandardButton.Yes,
        )

        window._on_reset_all()

        assert window._error_badge.text() == ""

    def test_already_visible_errors_dock_never_increments_the_badge(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        window.show()
        qtbot.waitExposed(window)
        window._show_errors_panel()
        qtbot.waitUntil(lambda: window._errors_tab_currently_visible)

        window._report_error("boom")

        assert window._error_badge.text() == ""


class TestDockIcons:
    def test_every_tabified_dock_gets_a_non_null_tab_icon(self, window: MainWindow) -> None:
        tab_bar = window._find_tab_bar()
        assert tab_bar is not None
        assert tab_bar.count() > 0
        for index in range(tab_bar.count()):
            assert not tab_bar.tabIcon(index).isNull()

    def test_icons_repaint_on_theme_switch(self, window: MainWindow) -> None:
        tab_bar = window._find_tab_bar()
        assert tab_bar is not None
        window._theme_actions["dark"].trigger()
        dark_pixmap = tab_bar.tabIcon(0).pixmap(20, 20).toImage()
        window._theme_actions["light"].trigger()  # restore default for test isolation
        light_pixmap = tab_bar.tabIcon(0).pixmap(20, 20).toImage()

        assert dark_pixmap != light_pixmap

    def test_concat_action_gets_a_non_null_icon(self, window: MainWindow) -> None:
        assert not window._concat_action.icon().isNull()


class TestBrandBar:
    def test_menu_widget_carries_the_logo_and_the_real_menu_bar(self, window: MainWindow) -> None:
        container = window.menuWidget()
        assert container is not None
        menu_bars = container.findChildren(QMenuBar)
        assert len(menu_bars) == 1
        logos = container.findChildren(QLabel)
        assert len(logos) == 1
        assert logos[0].pixmap() is not None and not logos[0].pixmap().isNull()

    def test_file_menu_is_still_reachable_through_the_embedded_menu_bar(
        self, window: MainWindow
    ) -> None:
        assert window._file_menu.title() != ""


class TestLineProfileWiring:
    def _load_images(self, window: MainWindow) -> None:
        # image_loaded, not just set_image(): this drives the ROI/Line Add
        # buttons' own "at least one image loaded" gating (see
        # MainWindow._refresh_images_loaded_gating), same as a real file
        # load already does via ImagePanel._on_load_success.
        reference_image = np.random.default_rng(0).random((32, 32))
        comparison_image = np.random.default_rng(1).random((32, 32))
        window.dual_image_view.reference_panel.set_image(reference_image)
        window.dual_image_view.reference_panel.image_loaded.emit(reference_image)
        window.dual_image_view.comparison_panel.set_image(comparison_image)
        window.dual_image_view.comparison_panel.image_loaded.emit(comparison_image)

    def test_add_button_creates_a_line_and_populates_the_table(self, window: MainWindow) -> None:
        self._load_images(window)

        window.line_profile_panel._add_button.click()

        assert len(window.dual_image_view.list_line_pairs()) == 1
        assert window.line_profile_panel._table.rowCount() == 1

    def test_add_button_disabled_until_an_image_is_loaded(self, window: MainWindow) -> None:
        assert not window.line_profile_panel._add_button.isEnabled()

        self._load_images(window)

        assert window.line_profile_panel._add_button.isEnabled()

    def test_renaming_updates_both_the_table_and_dual_image_view(self, window: MainWindow) -> None:
        pair = window.dual_image_view.add_line()

        window.line_profile_panel.rename_line_requested.emit(pair.line_id, "My Line")

        assert window.dual_image_view.get_line_pair(pair.line_id).name == "My Line"
        assert window.line_profile_panel._table.item(0, 0).text() == "My Line"

    def test_deleting_removes_the_row(self, window: MainWindow) -> None:
        pair = window.dual_image_view.add_line()

        window.line_profile_panel.delete_line_requested.emit(pair.line_id)

        assert window.dual_image_view.list_line_pairs() == []
        assert window.line_profile_panel._table.rowCount() == 0

    def test_duplicating_creates_a_second_line(self, window: MainWindow) -> None:
        pair = window.dual_image_view.add_line()

        window.line_profile_panel.duplicate_line_requested.emit(pair.line_id)

        assert len(window.dual_image_view.list_line_pairs()) == 2
        assert window.line_profile_panel._table.rowCount() == 2

    def test_selecting_a_row_locks_other_lines_and_populates_the_chart(
        self, window: MainWindow
    ) -> None:
        self._load_images(window)
        pair1 = window.dual_image_view.add_line()
        pair2 = window.dual_image_view.add_line()

        window.line_profile_panel._table.setCurrentCell(0, 0)

        assert pair1.lines["reference"].translatable
        assert not pair2.lines["reference"].translatable
        assert len(window.line_profile_panel._plot_widget.getPlotItem().listDataItems()) == 2

    def test_dragging_the_selected_line_refreshes_the_chart(self, window: MainWindow) -> None:
        self._load_images(window)
        pair = window.dual_image_view.add_line()
        window.line_profile_panel._table.setCurrentCell(0, 0)

        line = pair.lines["reference"]
        line.movePoint(line.getHandles()[1], (10.0, 10.0))

        assert len(window.line_profile_panel._plot_widget.getPlotItem().listDataItems()) == 2

    def test_adding_a_comparison_panel_extends_the_available_panel_labels(
        self, window: MainWindow
    ) -> None:
        window.dual_image_view.add_comparison_panel()

        assert window.line_profile_panel._panel_labels == [
            ("reference", "R"),
            ("comparison", "C1"),
            ("comparison_2", "C2"),
        ]

    def test_full_flow_add_drag_select_see_chart(self, window: MainWindow) -> None:
        self._load_images(window)
        window.dual_image_view.add_comparison_panel()
        window.dual_image_view.list_panels()[2].set_image(np.random.default_rng(2).random((32, 32)))

        window.line_profile_panel._add_button.click()
        pair = window.dual_image_view.list_line_pairs()[0]
        line = pair.lines["reference"]
        line.movePoint(line.getHandles()[1], (20.0, 25.0))
        window.line_profile_panel._table.setCurrentCell(0, 0)

        data_items = window.line_profile_panel._plot_widget.getPlotItem().listDataItems()
        assert len(data_items) == 3  # R, C1, C2

    def test_a_profile_recompute_failure_reaches_the_errors_panel(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._load_images(window)
        pair = window.dual_image_view.add_line()
        window.dual_image_view.set_active_line(pair.line_id)

        def _boom(*_args: object, **_kwargs: object) -> None:
            raise RuntimeError("synthetic failure")

        monkeypatch.setattr("qicom.gui.widgets.dual_image_view._sample_line_profile", _boom)

        window.dual_image_view._recompute_active_line_profile()  # must not raise

        assert window.errors_panel._table.rowCount() == 1


class TestExportLineProfilesWiring:
    def _load_images(self, window: MainWindow) -> None:
        rng = np.random.default_rng(0)
        window.dual_image_view.reference_panel.set_image(rng.random((32, 32)))
        window.dual_image_view.comparison_panel.set_image(np.random.default_rng(1).random((32, 32)))

    def test_no_lines_reports_an_error_and_never_opens_a_dialog(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        opened = []
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getSaveFileName",
            lambda *a, **k: opened.append(True) or ("", ""),
        )

        window._on_export_line_profiles_requested()

        assert opened == []
        assert window.errors_panel._table.rowCount() == 1

    def test_cancelled_save_dialog_is_a_noop(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._load_images(window)
        window.dual_image_view.add_line()
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getSaveFileName", lambda *a, **k: ("", "")
        )

        window._on_export_line_profiles_requested()  # must not raise

        assert window.errors_panel._table.rowCount() == 0

    def test_writes_a_workbook_with_one_sheet_per_line_and_reports_success(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        self._load_images(window)
        pair_a = window.dual_image_view.add_line()
        pair_b = window.dual_image_view.add_line()
        window.dual_image_view.rename_line(pair_a.line_id, "First")
        window.dual_image_view.rename_line(pair_b.line_id, "Second")
        out_path = tmp_path / "lines.xlsx"
        monkeypatch.setattr(
            "qicom.gui.main_window.QFileDialog.getSaveFileName",
            lambda *a, **k: (str(out_path), "Excel (*.xlsx)"),
        )

        window._on_export_line_profiles_requested()

        assert window.errors_panel._table.rowCount() == 0
        assert out_path.exists()
        from openpyxl import load_workbook

        workbook = load_workbook(out_path)
        assert workbook.sheetnames == ["Line 1", "Line 2"]
        assert workbook["Line 1"].cell(row=1, column=1).value == "First"
        assert workbook["Line 2"].cell(row=1, column=1).value == "Second"


class TestImportStep:
    def test_import_step_shows_its_own_panel(self, window: MainWindow) -> None:
        window.show()
        window.workflow_rail.set_current_stage("import")

        assert window.import_panel.isVisible()

    def test_import_project_button_opens_the_project_picker(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Reported: step 1 offered only "load images", while reopening a
        saved project was buried in the File menu."""
        opened: list[bool] = []
        monkeypatch.setattr(
            main_window_module.QFileDialog,
            "getOpenFileName",
            lambda *a, **k: opened.append(True) or ("", ""),
        )

        window.import_panel._import_project_button.click()

        assert opened == [True]


class TestBackgroundExclusionWiring:
    """The Remove-air step's exclusion flow, end to end through
    MainWindow: the green sample circle, the Detect button, the painted
    result and the paint-style eraser."""

    @staticmethod
    def _load_phantom(window: MainWindow) -> np.ndarray:
        rows, columns = np.mgrid[0:64, 0:64]
        body = (rows - 32) ** 2 + (columns - 32) ** 2 < 20**2
        image = np.full((64, 64), 5.0)
        image[body] = 300.0
        window.dual_image_view.reference_panel.set_image(image)
        window.dual_image_view.comparison_panel.set_image(image.copy())
        return body

    @staticmethod
    def _enable_and_detect(window: MainWindow, qtbot: QtBot) -> None:
        window.background_panel._enabled_check.setChecked(True)
        window.background_panel._air_roi_button.click()
        roi = window.dual_image_view._air_roi
        assert roi is not None
        roi.setPos(2, 2)
        roi.setSize((10, 10))
        window.background_panel._detect_button.click()
        qtbot.waitUntil(lambda: window.background_panel._estimate is not None, timeout=5000)

    def test_the_panel_belongs_to_the_remove_air_step(self, window: MainWindow) -> None:
        window.show()

        window.workflow_rail.set_current_stage("background")
        assert window.background_panel.isVisible()

        window.workflow_rail.set_current_stage("measure")
        assert not window.background_panel.isVisible()

    def test_the_create_button_puts_the_green_circle_on_the_canvas(
        self, window: MainWindow
    ) -> None:
        self._load_phantom(window)
        window.background_panel._enabled_check.setChecked(True)

        window.background_panel._air_roi_button.click()

        assert window.dual_image_view.has_air_roi
        assert window.dual_image_view.air_roi_count == 1
        assert window.background_panel._air_roi_button.text() == translator.tr(
            "background.air_roi_placed"
        )

    def test_creating_the_sample_stays_on_the_background_step(self, window: MainWindow) -> None:
        """Placing the circle is the first of this step's three actions
        (place it, detect, correct by hand) - moving the window on to
        Analyze threw the user out of the step they had just started."""
        self._load_phantom(window)
        window.workflow_rail.set_current_stage("background")
        window.background_panel._enabled_check.setChecked(True)

        window.background_panel._air_roi_button.click()

        assert window.workflow_rail.current_stage == "background"

    def test_a_second_click_cannot_add_another_air_sample(self, window: MainWindow) -> None:
        """The air level is read from one circle only, so a second sample
        would silently widen what counts as air."""
        self._load_phantom(window)
        window.background_panel._enabled_check.setChecked(True)

        window.background_panel._air_roi_button.click()
        window.background_panel._air_roi_button.click()

        assert window.dual_image_view.air_roi_count == 1
        assert not window.background_panel._air_roi_button.isEnabled()
        assert window.background_panel._air_roi_count_label.text() == translator.tr(
            "background.air_roi_count", count=1
        )

    def test_the_air_circle_is_not_a_measurement_roi(self, window: MainWindow) -> None:
        """It must not reach the ROI panel, the results, the exports or the
        heatmap targets - which is exactly what staying out of the ROI
        registry guarantees."""
        self._load_phantom(window)
        window.background_panel._enabled_check.setChecked(True)

        window.background_panel._air_roi_button.click()

        assert window.dual_image_view.list_roi_pairs() == []

    def test_removing_the_circle_clears_the_detection(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        self._load_phantom(window)
        self._enable_and_detect(window, qtbot)

        window.background_panel._clear_air_rois_button.click()

        assert not window.dual_image_view.has_air_roi
        assert window.background_panel._estimate is None
        for panel in window.dual_image_view.list_panels():
            assert panel._heatmap_item is None

    def test_detecting_reports_the_estimate_and_paints_the_result(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        self._load_phantom(window)

        self._enable_and_detect(window, qtbot)

        estimate = window.background_panel._estimate
        assert estimate is not None
        assert 0.0 < estimate.excluded_fraction < 1.0
        for panel in window.dual_image_view.list_panels():
            assert panel._heatmap_item is not None

    def test_a_finished_older_detection_cannot_replace_the_newer_one(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._load_phantom(window)
        window.background_panel._enabled_check.setChecked(True)
        window.background_panel._air_roi_button.click()
        roi = window.dual_image_view._air_roi
        assert roi is not None
        roi.setPos(2, 2)
        roi.setSize((10, 10))
        window._background_detection_requested = True
        pending: list[tuple[object, tuple[object, ...], object]] = []

        def defer(fn, *args, on_success, on_error, **_kwargs):
            pending.append((fn, args, on_success))
            return object()

        monkeypatch.setattr(main_window_module, "run_in_thread", defer)
        window._refresh_background_estimate()
        window.background_panel._tolerance_spin.setValue(4.0)
        window._background_refresh_timer.stop()
        window._refresh_background_estimate()
        assert len(pending) == 2

        old_fn, old_args, old_success = pending[0]
        old_success(old_fn(*old_args))
        assert window.background_panel._estimate is None

        new_fn, new_args, new_success = pending[1]
        new_success(new_fn(*new_args))
        estimate = window.background_panel._estimate
        assert estimate is not None
        assert new_args[0].settings.tolerance_factor == pytest.approx(4.0)

    def test_nothing_is_detected_before_the_button_is_pressed(self, window: MainWindow) -> None:
        """Enabling the checkbox and moving the circle are steps 1 and 2 -
        the answer appears when the user asks for it, not while they are
        still placing the sample."""
        self._load_phantom(window)
        window.background_panel._enabled_check.setChecked(True)
        window.background_panel._air_roi_button.click()

        window._refresh_background_estimate()

        assert window.background_panel._estimate is None
        for panel in window.dual_image_view.list_panels():
            assert panel._heatmap_item is None

    def test_a_failed_detection_is_shown_in_the_panel_not_raised(self, window: MainWindow) -> None:
        self._load_phantom(window)
        window.background_panel._enabled_check.setChecked(True)

        window.background_panel._detect_button.click()  # no circle placed

        assert window.background_panel._estimate is None
        assert window.background_panel._summary.text()

    def test_the_eraser_puts_painted_pixels_back_into_the_measurement(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        self._load_phantom(window)
        self._enable_and_detect(window, qtbot)
        before = window.controller.build_background_exclusion(window.background_panel.settings())
        assert before is not None

        window.background_panel._brush_spin.setValue(6)
        window.background_panel._eraser_button.setChecked(True)
        window.dual_image_view.reference_panel.air_eraser_point.emit((50.0, 50.0))

        assert window.dual_image_view.air_eraser_active
        after = window.controller.build_background_exclusion(window.background_panel.settings())
        assert after is not None
        assert after[0][50, 50]
        assert after[1].excluded_fraction < before[1].excluded_fraction
        # and the readout followed the correction, without re-detecting
        assert f"{100.0 * after[1].excluded_fraction:.1f}" in (
            window.background_panel._summary.text()
        )

    def test_the_air_brush_excludes_background_the_threshold_missed(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        self._load_phantom(window)
        self._enable_and_detect(window, qtbot)
        before = window.controller.build_background_exclusion(window.background_panel.settings())
        assert before is not None
        # The bright body is measured by automatic detection. Painting it
        # as air is artificial but gives an unambiguous end-to-end check of
        # the manual exclusion path.
        assert before[0][32, 32]

        window.background_panel._brush_spin.setValue(3)
        window.background_panel._paint_button.setChecked(True)
        window.dual_image_view.reference_panel.air_eraser_point.emit((32.0, 32.0))

        after = window.controller.build_background_exclusion(window.background_panel.settings())
        assert after is not None
        assert window.dual_image_view.air_brush_active
        assert not after[0][32, 32]
        assert after[1].excluded_fraction > before[1].excluded_fraction

    def test_undoing_the_corrections_goes_back_to_the_detected_result(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        self._load_phantom(window)
        self._enable_and_detect(window, qtbot)
        detected = window.controller.build_background_exclusion(window.background_panel.settings())
        assert detected is not None
        window.background_panel._eraser_button.setChecked(True)
        window.dual_image_view.reference_panel.air_eraser_point.emit((50.0, 50.0))

        window.background_panel._clear_erasures_button.click()

        restored = window.controller.build_background_exclusion(window.background_panel.settings())
        assert restored is not None
        assert np.array_equal(restored[0], detected[0])

    def test_picking_another_tool_disarms_the_eraser(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        """Otherwise the button would keep claiming the next click erases
        something, while the canvas has moved on to another tool."""
        self._load_phantom(window)
        self._enable_and_detect(window, qtbot)
        window.background_panel._eraser_button.setChecked(True)

        window.dual_image_view.set_interaction_mode(MODE_ROI_ELLIPSE)

        assert not window.dual_image_view.air_eraser_active
        assert not window.background_panel.is_eraser_active()

    def test_leaving_the_remove_air_step_takes_the_preview_off_the_canvas(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        self._load_phantom(window)
        window.workflow_rail.set_current_stage("background")
        self._enable_and_detect(window, qtbot)

        window.workflow_rail.set_current_stage("results")

        for panel in window.dual_image_view.list_panels():
            assert panel._heatmap_item is None
        assert not window.background_panel.is_eraser_active()

    def test_returning_to_the_remove_air_step_paints_it_again(
        self, window: MainWindow, qtbot: QtBot
    ) -> None:
        self._load_phantom(window)
        window.workflow_rail.set_current_stage("background")
        self._enable_and_detect(window, qtbot)
        window.workflow_rail.set_current_stage("results")

        window.workflow_rail.set_current_stage("background")

        for panel in window.dual_image_view.list_panels():
            assert panel._heatmap_item is not None

    def test_calculate_is_given_the_panels_settings(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._load_phantom(window)
        window.background_panel._enabled_check.setChecked(True)
        captured: list[object] = []
        monkeypatch.setattr(
            window.controller,
            "calculate",
            lambda *args, background=None, **kwargs: captured.append(background),
        )

        window._start_calculate(1.0, None)

        assert len(captured) == 1
        assert getattr(captured[0], "enabled", None) is True

    def test_the_results_table_says_what_was_excluded(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._load_phantom(window)
        monkeypatch.setattr(
            window.controller,
            "compute_per_slice_distributions",
            lambda *args, **kwargs: None,
        )
        window.controller._last_background_estimate = BackgroundEstimate(
            method=METHOD_MANUAL, air_level=5.0, tolerance=3.0, excluded_fraction=0.7
        )

        window._on_calculate_success(_make_results())

        note = window.results_panel._background_note_label
        assert note.isVisibleTo(window.results_panel)
        assert "70.0" in note.text()

    def test_a_plain_run_leaves_no_note_behind(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._load_phantom(window)
        monkeypatch.setattr(
            window.controller,
            "compute_per_slice_distributions",
            lambda *args, **kwargs: None,
        )

        window._on_calculate_success(_make_results())

        assert not window.results_panel._background_note_label.isVisibleTo(window.results_panel)


class TestResultsAndHeatmapSteps:
    def test_results_step_hides_the_image_canvas(self, window: MainWindow) -> None:
        """Reported: "in the results section I only want the numeric
        results and the charts, I don't want to see the images here taking
        up space for nothing"."""
        window.show()

        window.workflow_rail.set_current_stage("results")

        assert not window.dual_image_view.isVisible()
        assert window.results_panel.isVisible()
        assert window.charts_panel.isVisible()
        assert window._heatmap_dock.isHidden()

    def test_heatmaps_step_is_its_own_step_with_the_images_back(self, window: MainWindow) -> None:
        window.show()

        window.workflow_rail.set_current_stage("heatmaps")

        assert window.dual_image_view.isVisible()
        assert window.heatmap_panel.isVisible()
        assert window._results_dock.isHidden()
        assert window._charts_dock.isHidden()

    def test_there_is_no_activity_step_any_more(self, window: MainWindow) -> None:
        assert "activity" not in window.workflow_rail.stage_keys
        assert window._history_dock.isHidden()
        assert window._log_dock.isHidden()
        assert window._errors_dock.isHidden()

    def test_history_log_and_errors_are_reachable_from_the_view_menu(
        self, window: MainWindow
    ) -> None:
        actions = window._panels_menu.actions()

        assert [action.text() for action in actions] == ["History", "Log", "Errors"]
        actions[2].trigger()
        assert not window._errors_dock.isHidden()


class TestExportStep:
    def test_export_step_shows_the_written_summary_instead_of_the_results_table(
        self, window: MainWindow
    ) -> None:
        window.show()
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.results_panel.set_dataframe(_make_results())

        window.workflow_rail.set_current_stage("export")

        assert window.export_summary_panel.isVisible()
        assert window._results_dock.isHidden()
        assert not window.dual_image_view.isVisible()

    def test_summary_lists_what_was_analysed_and_what_will_be_written(
        self, window: MainWindow
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        window.controller.add_roi("rectangle")
        window.results_panel.set_dataframe(_make_results())

        window.workflow_rail.set_current_stage("export")

        report = window.export_summary_panel._report.toPlainText()
        assert "ROI 1" in report
        assert "Results workbook" in report
        assert "key image" in report
        # Regression: display_name()/slice_count() are methods, not
        # properties - the summary once rendered "<bound method ...>" for
        # every loaded image, which no assertion on ROI names caught.
        assert "bound method" not in report
        assert "Image R: REFERENCE (R) (1 slice(s))" in report

    def test_both_export_options_start_checked(self, window: MainWindow) -> None:
        assert window.export_summary_panel.include_heatmaps
        assert window.export_summary_panel.include_key_images

    def test_unchecking_an_option_drops_it_from_the_promised_file_list(
        self, window: MainWindow
    ) -> None:
        window.dual_image_view.reference_panel.set_image(np.zeros((10, 10)))
        window.dual_image_view.comparison_panel.set_image(np.zeros((10, 10)))
        window.workflow_rail.set_current_stage("export")
        assert "key image" in window.export_summary_panel._report.toPlainText()

        window.export_summary_panel._include_key_images_check.setChecked(False)

        assert "key image" not in window.export_summary_panel._report.toPlainText()

    def test_unchecked_options_are_passed_to_the_export(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        captured: dict[str, object] = {}

        def _fake_export(*args: object, **kwargs: object) -> list[Path]:
            captured.update(kwargs)
            return [tmp_path / "results.xlsx"]

        monkeypatch.setattr(window.controller, "export_results", _fake_export)
        monkeypatch.setattr(window, "_prompt_save_export_choice", lambda: "results")
        monkeypatch.setattr(
            main_window_module.QFileDialog, "getExistingDirectory", lambda *a, **k: str(tmp_path)
        )
        monkeypatch.setattr(
            main_window_module.QInputDialog, "getText", lambda *a, **k: ("run", True)
        )
        stub_message_box(monkeypatch, main_window_module, information=lambda *a, **k: None)
        window.results_panel.set_dataframe(_make_results())
        window.export_summary_panel._include_heatmaps_check.setChecked(False)

        window._on_save_export_requested()

        assert captured["include_heatmaps"] is False
        assert captured["include_key_images"] is True


class TestDockLayout:
    def test_results_dock_is_not_tabified_with_log_or_errors(self, window: MainWindow) -> None:
        """Regression test: Results used to be tabbed with Log, so clicking
        the Log tab hid the Calculate button - reported as confusing by a
        user. Results must now be the sole occupant of its own dock area.
        Results/ROIs were later swapped per user request (ROIs+History+
        Annotations+Log+Errors tabbed at the bottom), then Results was
        moved to the right so images stay on the left, calculations on
        the right."""
        assert window.dockWidgetArea(window._results_dock) == Qt.DockWidgetArea.RightDockWidgetArea
        assert window.dockWidgetArea(window._roi_dock) == Qt.DockWidgetArea.BottomDockWidgetArea
        assert window.dockWidgetArea(window._log_dock) == Qt.DockWidgetArea.BottomDockWidgetArea
        assert window.dockWidgetArea(window._errors_dock) == Qt.DockWidgetArea.BottomDockWidgetArea

    def test_results_and_charts_are_side_by_side(self, window: MainWindow, qtbot: QtBot) -> None:
        window.resize(1280, 800)
        window.show()
        window.workflow_rail.set_current_stage("results")
        qtbot.waitUntil(
            lambda: window._charts_dock.width() > 0 and window._results_dock.width() > 0
        )

        charts = window._charts_dock.geometry()
        results = window._results_dock.geometry()
        assert charts.right() < results.left() or results.right() < charts.left()
        assert charts.top() <= results.bottom() and results.top() <= charts.bottom()

    def test_heatmap_controls_are_a_compact_right_inspector(self, window: MainWindow) -> None:
        assert window.dockWidgetArea(window._heatmap_dock) == Qt.DockWidgetArea.RightDockWidgetArea
        window.show()
        window.workflow_rail.set_current_stage("heatmaps")
        window._resize_workflow_docks("heatmaps")

        assert window._heatmap_dock.width() < window.dual_image_view.width()
        assert window.dual_image_view.geometry().right() < window._heatmap_dock.geometry().left()

    def test_every_dock_is_locked_against_dragging_or_floating(self, window: MainWindow) -> None:
        for dock in window.findChildren(QDockWidget):
            assert not dock.features() & QDockWidget.DockWidgetFeature.DockWidgetMovable
            assert not dock.features() & QDockWidget.DockWidgetFeature.DockWidgetFloatable

    def test_analyze_has_the_metric_setup_instead_of_results(self, window: MainWindow) -> None:
        window.show()
        assert not window.results_panel.isVisible()

        window.workflow_rail.set_current_stage("analyze")
        assert window.analysis_setup_panel.isVisible()
        assert not window.results_panel.isVisible()
        assert not window.dual_image_view.isVisible()

        window.workflow_rail.set_current_stage("results")
        assert window.results_panel.isVisible()
        assert not window.analysis_setup_panel.isVisible()

    def test_save_and_export_actions_exist_only_in_export_stage(self, window: MainWindow) -> None:
        window.workflow_rail.set_current_stage("analyze")
        assert not window._save_analysis_action.isVisible()
        assert not window._save_project_folder_action.isVisible()
        assert window.results_panel._export_button.isHidden()

        window.workflow_rail.set_current_stage("export")
        assert window._save_analysis_action.isVisible()
        assert window._save_project_folder_action.isVisible()
        assert not window.results_panel._export_button.isHidden()

    def test_measurement_bottom_docks_are_grouped(self, window: MainWindow, qtbot: QtBot) -> None:
        from PySide6.QtWidgets import QTabBar

        window.show()
        tab_bar = window.findChild(QTabBar)
        assert tab_bar is not None

        window.workflow_rail.set_current_stage("measure")
        qtbot.waitUntil(lambda: tab_bar.count() == 3)
        assert [tab_bar.tabText(i) for i in range(tab_bar.count())] == [
            "ROIs",
            "Line",
            "Annotations",
        ]


class TestDefaultDockSizing:
    """The central DualImageView used to end up squeezed by the bottom-
    docked panels' own sizeHint-driven default share of the window - see
    MainWindow._apply_default_dock_sizes."""

    def test_caps_the_bottom_dock_group_to_reserve_55_percent_for_images(
        self, window: MainWindow, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        window.resize(1200, 1000)
        calls: list[tuple[list[object], list[int]]] = []
        monkeypatch.setattr(
            window,
            "resizeDocks",
            lambda docks, sizes, orientation: calls.append((docks, sizes)),
        )

        window._apply_default_dock_sizes()

        assert calls == [([window._roi_dock], [450])]

    def test_first_show_schedules_it_exactly_once(
        self, window: MainWindow, qtbot: QtBot, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        calls: list[bool] = []
        monkeypatch.setattr(window, "_apply_default_dock_sizes", lambda: calls.append(True))

        window.show()
        qtbot.waitUntil(lambda: len(calls) == 1, timeout=2000)
        window.hide()
        window.show()

        # Still exactly one call - the flag guards every subsequent show.
        assert calls == [True]


class TestProjectSaveDeidentificationOverride:
    @staticmethod
    def _refusing_controller(window: MainWindow, monkeypatch: pytest.MonkeyPatch) -> list[object]:
        calls: list[object] = []

        def save(_path: Path, **kwargs: Any) -> None:
            unverified = kwargs.get("unverified_dicom")
            calls.append(unverified)
            if unverified is None:
                raise DicomDeidentificationError("DICOM export blocked: BurnedInAnnotation")
            unverified.append(UnverifiedDicomFile("scan.dcm", "BurnedInAnnotation", True))

        monkeypatch.setattr(window.controller, "save_project_zip", save)
        return calls

    def test_declining_stops_the_save(
        self, window: MainWindow, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        calls = self._refusing_controller(window, monkeypatch)
        monkeypatch.setattr(
            main_window_module, "confirm_unverified_dicom_export", lambda *_args: False
        )

        assert window._save_project_zip(tmp_path / "p.zip") is False

        assert calls == [None]
        assert "BurnedInAnnotation" in window.errors_panel.records()[0].message

    def test_accepting_saves_again_with_the_override_and_lists_the_files(
        self, window: MainWindow, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        calls = self._refusing_controller(window, monkeypatch)
        monkeypatch.setattr(
            main_window_module, "confirm_unverified_dicom_export", lambda *_args: True
        )

        assert window._save_project_zip(tmp_path / "p.zip") is True

        assert calls[0] is None and isinstance(calls[1], list)
        messages = [record.message for record in window.errors_panel.records()]
        assert any("scan.dcm" in message for message in messages)
